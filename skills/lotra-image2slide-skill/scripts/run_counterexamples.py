#!/usr/bin/env python3
"""Persist four mandatory M0 counterexamples using the real saved-deck preview."""

import argparse
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from core.common import digest, read_json, write_json
from core.inspect import check, NS


def run(scene_path,pptx,preview,out):
    scene_path,pptx,preview,out=map(Path,(scene_path,pptx,preview,out))
    out.mkdir(parents=True,exist_ok=False)
    scene=read_json(scene_path)
    cases=[]
    positive=check(scene,scene_path.parent,pptx)
    for fault in ("png_icon","rich_text_formula","deleted_connector","hidden_native_screenshot"):
        deck=Presentation(pptx)
        slide=deck.slides[0]
        named={s.name:s for s in slide.shapes}
        if fault=="png_icon":
            group=named["Device"]
            image=Image.open(preview)
            rx,ry=image.width/deck.slide_width,image.height/deck.slide_height
            crop=image.crop((round(group.left*rx),round(group.top*ry),round((group.left+group.width)*rx),round((group.top+group.height)*ry)))
            image_path=out/"icon_lookalike.png"
            crop.save(image_path)
            pic=slide.shapes.add_picture(str(image_path),group.left,group.top,group.width,group.height)
            pic.name="Device"
            slide.shapes._spTree.remove(group._element)
        elif fault=="rich_text_formula":
            shape=named["Equation_scripts"]
            paragraph=shape.text_frame.paragraphs[0]
            paragraph.clear()
            for text,baseline in [("f",0),("k",-25000),("max",35000)]:
                run=paragraph.add_run()
                run.text=text
                run.font.name="Cambria Math"
                run.font.size=Pt(19 if baseline==0 else 13)
                run._r.get_or_add_rPr().set("baseline",str(baseline))
        elif fault=="deleted_connector":
            slide.shapes._spTree.remove(named["Link"]._element)
        else:
            for shape in slide.shapes:
                shape._element.find(".//p:cNvPr",NS).set("hidden","1")
            slide.shapes.add_picture(str(preview),0,0,deck.slide_width,deck.slide_height).name="ScreenshotCover"
        file=out/(fault+".pptx")
        deck.save(file)
        result=check(scene,scene_path.parent,file)
        write_json(out/(fault+".json"),result)
        cases.append({"case":fault,"pptx_hash":digest(file),"expected":"FAIL","actual":result["status"],
                      "rejected":result["status"]=="FAIL","failed_checks":[c["check_id"] for c in result["checks"] if c["status"]=="FAIL"]})
    report={"scope":"four_core_M0_structural_counterexamples","positive_status":positive["status"],"cases":cases,
            "false_rejects":int(positive["status"]!="PASS"),"missed_counterexamples":sum(not x["rejected"] for x in cases),
            "source_pptx_hash":digest(pptx),"actual_preview_hash":digest(preview),"status":"PASS" if positive["status"]=="PASS" and all(c["rejected"] for c in cases) else "FAIL"}
    write_json(out/"results.json",report)
    return report


if __name__=="__main__":
    p=argparse.ArgumentParser()
    p.add_argument("scene");p.add_argument("pptx");p.add_argument("preview");p.add_argument("--out",required=True)
    a=p.parse_args()
    result=run(a.scene,a.pptx,a.preview,a.out)
    print(result["status"])
    raise SystemExit(0 if result["status"]=="PASS" else 2)
