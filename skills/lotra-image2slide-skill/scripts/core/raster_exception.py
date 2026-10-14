"""Measured, scoped raster-exception trial against an actual native candidate.

This gate can make a raster candidate *eligible for review*, never turn a
picture into native content or certify scientific/Office editing behavior.
The caller freezes the source, ROI, score profile and reuse size before either
candidate is compared. Both saved PPTs and render receipts are re-read here.
"""

import math
from pathlib import Path

from PIL import Image
from pptx import Presentation

from .common import digest, identity, read_json, safe_path, write_json
from .render import verify_alternative_render
from .saved_ppt_diff import snapshot
from .score_profiles import evaluate as score_candidates
from .typed_metrics import DEFAULT_PROTOCOL, image_measure


METRICS = {
    "ssim": lambda m: m["ssim"],
    "residual_rgb_255": lambda m: m["regional_residual_rgb_255"]["mean"],
    "color_delta_e": lambda m: m["color_delta_e_2000"]["mean"],
    "internal_edge_f1": lambda m: m["internal_edges"]["f1"],
    "mse_rgb_01": lambda m: m["mse_rgb_01"],
    "texture_histogram_tv": lambda m: m["texture_color_histogram_tv"],
}
ALLOWED_ROLES = {"environment_asset", "material_texture", "decorative_texture"}


def _rendered(pptx, receipt_file):
    receipt_file = Path(receipt_file)
    receipt = read_json(receipt_file)
    if receipt.get("backend") != "LibreOffice":
        # The Office importer has a separate observation route. Until it has
        # a same-protocol pair, this bounded comparator refuses to guess.
        raise ValueError("Raster exception currently requires two verified LibreOffice receipts")
    verify_alternative_render(pptx, receipt_file)
    return receipt, receipt_file.parent / "Final_preview.png"


def _owner_records(pptx, path):
    record = snapshot(pptx)
    names = [name for name in record["objects"] if name == path or name.startswith(path + "/")]
    if path not in record["objects"] or not names:
        raise ValueError("Declared representation owner is absent from saved PPT")
    return record, names


def _same_context(vector, raster, vector_path, raster_path):
    def outside(snapshot_record, owner):
        return {key: value for key, value in snapshot_record["objects"].items()
                if key != owner and not key.startswith(owner + "/")}
    if (vector["slide_count"] != raster["slide_count"]
            or vector["slide_shells"] != raster["slide_shells"]
            or outside(vector, vector_path) != outside(raster, raster_path)):
        raise ValueError("Raster trial changed unrelated saved shapes or slide context")
    vector_theme = {key: value for key, value in vector["shared_parts"].items() if key.startswith("ppt/theme/")}
    raster_theme = {key: value for key, value in raster["shared_parts"].items() if key.startswith("ppt/theme/")}
    if vector_theme != raster_theme:
        raise ValueError("Raster trial changed the shared presentation theme")


def _effective_resolution(pptx, owner_path, source_asset, max_reuse_inches):
    if (not isinstance(max_reuse_inches, list) or len(max_reuse_inches) != 2
            or any(type(x) not in (int, float) or not math.isfinite(x) or x <= 0
                   for x in max_reuse_inches)):
        raise ValueError("Maximum intended reuse size needs two positive inches")
    with Image.open(source_asset) as image:
        native_pixels = image.size
    prs = Presentation(pptx)
    def visit(shapes, prefix):
        for shape in shapes:
            key = prefix + "/" + shape.name
            if key == owner_path:
                return shape
            if shape.shape_type == 6:
                found = visit(shape.shapes, key)
                if found is not None:
                    return found
        return None
    matches = [visit(slide.shapes, "slide-" + str(i)) for i, slide in enumerate(prs.slides, 1)]
    pictures = [shape for shape in matches if shape is not None]
    if len(pictures) != 1 or pictures[0].shape_type != 13:
        raise ValueError("Raster owner must be one separately saved picture")
    if (pictures[0].rotation or any(getattr(pictures[0], field) for field in
                                    ("crop_left", "crop_right", "crop_top", "crop_bottom"))):
        raise ValueError("Cropped or rotated raster needs a separate effective-resolution protocol")
    media_pixels = pictures[0].image.size
    displayed_inches = [pictures[0].width / 914400, pictures[0].height / 914400]
    if any(max_reuse_inches[i] + 1e-8 < displayed_inches[i] for i in (0, 1)):
        raise ValueError("Declared maximum reuse size is smaller than current displayed size")
    information_pixels = [min(a, b) for a, b in zip(native_pixels, media_pixels)]
    effective_ppi = [information_pixels[i] / max_reuse_inches[i] for i in (0, 1)]
    return {"source_information_pixels": list(native_pixels), "saved_media_pixels": list(media_pixels),
            "effective_information_pixels": information_pixels, "displayed_inches": displayed_inches,
            "max_reuse_inches": max_reuse_inches,
            "effective_ppi": effective_ppi}


def compare(contract, base, vector_pptx, vector_receipt, raster_pptx, raster_receipt):
    """Return a bound trial decision; an incomplete gate is never auto-approved."""
    if contract.get('appearance_policy') == 'semantic_hd_replace':
        from .appearance import semantic_raster_comparison
        return semantic_raster_comparison(contract, base, vector_pptx, vector_receipt, raster_pptx, raster_receipt)
    if contract.get("schema_version") != "raster-exception-1":
        raise ValueError("Unknown raster-exception contract")
    if (contract.get("role") not in ALLOWED_ROLES
            or contract.get("scientific_content_excluded") is not True
            or contract.get("semantic_mark") is not False
            or contract.get("native_internal_edit_required") is not False):
        raise ValueError("Scientific, semantic or internally editable nodes cannot use this raster exception")
    if not isinstance(contract.get("plan_hash"), str) or len(contract["plan_hash"]) != 64:
        raise ValueError("Exception must be bound to a frozen drawing plan")
    plan_file = safe_path(base, contract["plan"]["path"])
    if digest(plan_file) != contract["plan"]["sha256"]:
        raise ValueError("Frozen raster-exception subplan changed")
    plan = read_json(plan_file)
    source = safe_path(base, contract["source"]["path"])
    original_asset = safe_path(base, contract["raster_information_source"]["path"])
    profile_file = safe_path(base, contract["profile"]["path"])
    calibration_file = safe_path(base, contract["calibration"]["path"])
    for path, ref in ((source, contract["source"]), (original_asset, contract["raster_information_source"]),
                      (profile_file, contract["profile"]), (calibration_file, contract["calibration"])):
        if digest(path) != ref["sha256"]:
            raise ValueError("Frozen exception input changed")
    profile = read_json(profile_file)
    calibration = read_json(calibration_file)
    if (plan.get("schema_version") != "raster-exception-plan-1"
            or identity(plan) != contract["plan_hash"]
            or plan.get("source_hash") != digest(source)
            or plan.get("scope_id") != contract.get("scope_id")
            or plan.get("roi") != contract.get("roi")
            or plan.get("owner_id") != contract["vector_owner_path"].split("/")[-1]
            or plan.get("native_method") != "adaptive_native_rectangles"
            or plan.get("raster_method") != "separate_picture"
            or not isinstance(plan.get("canvas_background_rgb"), list)
            or len(plan["canvas_background_rgb"]) != 3
            or any(type(value) is not int or not 0 <= value <= 255
                   for value in plan["canvas_background_rgb"])
            or type(plan.get("native_shape_budget")) is not int
            or plan["native_shape_budget"] < 2
            or not isinstance(plan.get("source_observation"), str)
            or not plan["source_observation"].strip()):
        raise ValueError("Raster exception is not bound to an input-specific frozen subplan")
    if (profile.get("source_hash") != digest(source)
            or profile.get("scope_id") != contract.get("scope_id")
            or contract.get("calibration_status") != "CALIBRATED"):
        raise ValueError("A predeclared calibrated profile for this source/scope is required")
    primary = contract.get("primary_metric")
    if (primary not in profile.get("metrics", {}) or primary not in METRICS
            or type(contract.get("minimum_primary_gain")) not in (int, float)
            or not math.isfinite(contract["minimum_primary_gain"])
            or contract["minimum_primary_gain"] <= 0
            or type(contract.get("minimum_native_composite")) not in (int, float)
            or not 0 <= contract["minimum_native_composite"] <= 1):
        raise ValueError("Primary gain and native comparison quality floor must be frozen")
    if (calibration.get("schema_version") != "raster-exception-calibration-1"
            or calibration.get("source_hash") != digest(source)
            or calibration.get("scope_id") != contract["scope_id"]
            or calibration.get("roi") != contract.get("roi")
            or calibration.get("profile_id") != profile["profile_id"]
            or calibration.get("primary_metric") != primary
            or calibration.get("minimum_primary_gain") != contract["minimum_primary_gain"]
            or calibration.get("minimum_native_composite") != contract["minimum_native_composite"]
            or calibration.get("measurement_protocol_hash") != identity(contract.get("measurement_protocol", DEFAULT_PROTOCOL))
            or not isinstance(calibration.get("anchors"), list)
            or len(calibration["anchors"]) < 2):
        raise ValueError("Frozen exception thresholds lack matching pre-candidate calibration anchors")
    vector_pptx, raster_pptx = Path(vector_pptx), Path(raster_pptx)
    expected_background = "".join(f"{value:02X}" for value in plan["canvas_background_rgb"])
    for pptx in (vector_pptx, raster_pptx):
        deck = Presentation(pptx)
        if (len(deck.slides) != 1
                or deck.slides[0].background.fill.type is None
                or str(deck.slides[0].background.fill.fore_color.rgb) != expected_background):
            raise ValueError("Saved slide background differs from frozen input-specific vector plan")
    v_receipt, v_preview = _rendered(vector_pptx, vector_receipt)
    r_receipt, r_preview = _rendered(raster_pptx, raster_receipt)
    comparable = ("backend", "backend_version", "executable", "environment", "mapping")
    if (any(v_receipt.get(key) != r_receipt.get(key) for key in comparable)
            or v_receipt.get("fontconfig") != r_receipt.get("fontconfig")
            or v_receipt["pages"][0]["canvas"] != r_receipt["pages"][0]["canvas"]):
        raise ValueError("Vector and raster renders used different measurement protocols")
    with Image.open(source) as image:
        canvas = image.size
        roi = contract.get("roi")
        if (not isinstance(roi, list) or len(roi) != 4 or any(type(x) is not int for x in roi)):
            raise ValueError("Frozen source ROI must use integer pixels")
        x, y, width, height = roi
        if x < 0 or y < 0 or width < 7 or height < 7 or x + width > canvas[0] or y + height > canvas[1]:
            raise ValueError("Frozen source ROI is empty or outside source")
        reference = image.convert("RGB").crop((x, y, x + width, y + height))
    with Image.open(v_preview) as image:
        if image.size != canvas:
            raise ValueError("Vector preview differs from fixed source canvas")
        vector_image = image.convert("RGB").crop((x, y, x + width, y + height))
    with Image.open(r_preview) as image:
        if image.size != canvas:
            raise ValueError("Raster preview differs from fixed source canvas")
        raster_image = image.convert("RGB").crop((x, y, x + width, y + height))
    vector, vector_names = _owner_records(vector_pptx, contract["vector_owner_path"])
    raster, raster_names = _owner_records(raster_pptx, contract["raster_owner_path"])
    _same_context(vector, raster, contract["vector_owner_path"], contract["raster_owner_path"])
    for protected in plan.get("protected_foreground_paths", []):
        if (protected.startswith(contract["vector_owner_path"] + "/")
                or protected not in vector["objects"] or protected not in raster["objects"]
                or vector["objects"][protected] != raster["objects"][protected]):
            raise ValueError("Protected native foreground is missing, inside raster owner or changed")
    if any(vector["objects"][name]["shape_type"] == "pic" or vector["objects"][name]["media_sha256"]
           for name in vector_names):
        raise ValueError("Native comparator has an embedded picture")
    native_leaves = sum(vector["objects"][name]["shape_type"] != "grpSp" for name in vector_names)
    if (type(contract.get("minimum_native_leaves")) is not int
            or native_leaves < contract["minimum_native_leaves"]
            or native_leaves > plan["native_shape_budget"]
            or contract["minimum_native_leaves"] < 2):
        raise ValueError("Vector comparator lacks the predeclared natural-part detail")
    if raster["objects"][contract["raster_owner_path"]]["shape_type"] != "pic" or len(raster_names) != 1:
        raise ValueError("Raster exception must apply to one separate picture layer")
    resolution = _effective_resolution(raster_pptx, contract["raster_owner_path"],
                                       original_asset, contract.get("max_reuse_inches"))
    minimum_ppi = contract.get("minimum_effective_ppi")
    if type(minimum_ppi) not in (int, float) or not math.isfinite(minimum_ppi) or minimum_ppi <= 0:
        raise ValueError("A positive effective-information PPI floor is required")
    protocol = contract.get("measurement_protocol", DEFAULT_PROTOCOL)
    measured = {"vector": image_measure(reference, vector_image, protocol),
                "raster": image_measure(reference, raster_image, protocol)}
    metric_names = set(profile["metrics"])
    if not metric_names <= set(METRICS):
        raise ValueError("Profile requests a metric not recomputed by this route")
    vectors = {kind: {name: METRICS[name](measurement) for name in metric_names}
               for kind, measurement in measured.items()}
    batch = {"schema_version": "score-batch-1", "profile_hash": identity(profile),
             "source_hash": digest(source), "scope_id": contract["scope_id"],
             "baseline": {"id": "native", "vector": vectors["vector"]},
             "candidates": [{"id": "raster", "vector": vectors["raster"]}]}
    scored = score_candidates(profile, batch)
    row = scored["candidates"][0]
    reasons = []
    if scored["baseline"]["composite"] < contract["minimum_native_composite"]:
        reasons.append("NATIVE_COMPARATOR_BELOW_PREDECLARED_QUALITY_FLOOR")
    if row["decision"] != "PROMOTABLE_SCORE_ONLY":
        reasons.append("RASTER_HAS_NO_VALID_NONCOMPENSATING_SCORE_GAIN")
    elif row["metric_gains"][primary] < contract["minimum_primary_gain"]:
        reasons.append("PRIMARY_VISUAL_GAIN_BELOW_PREDECLARED_DELTA")
    if min(resolution["effective_ppi"]) < minimum_ppi:
        reasons.append("RASTER_EFFECTIVE_INFORMATION_RESOLUTION_TOO_LOW")
    return {"schema_version": "raster-exception-comparison-1",
            "status": "REJECTED_RASTER_EXCEPTION" if reasons else "ELIGIBLE_FOR_SEMANTIC_AND_EDITOR_REVIEW",
            "reasons": reasons, "contract_hash": identity(contract), "plan_hash": contract["plan_hash"],
            "source_hash": digest(source), "profile_hash": identity(profile), "roi": roi,
            "calibration_hash": digest(calibration_file),
            "vector_pptx_hash": digest(vector_pptx), "raster_pptx_hash": digest(raster_pptx),
            "vector_render_receipt_hash": digest(vector_receipt),
            "raster_render_receipt_hash": digest(raster_receipt),
            "native_leaf_count": native_leaves, "resolution": resolution,
            "raw_metrics": vectors, "scoring": scored,
            "scope": "actual same-protocol saved-PPT render comparison; semantic correctness, native editing and full-slide acceptance remain separate"}


def compare_file(contract_file, vector_pptx, vector_receipt, raster_pptx, raster_receipt, output):
    contract_file = Path(contract_file)
    result = compare(read_json(contract_file), contract_file.parent,
                     vector_pptx, vector_receipt, raster_pptx, raster_receipt)
    write_json(output, result)
    return result
