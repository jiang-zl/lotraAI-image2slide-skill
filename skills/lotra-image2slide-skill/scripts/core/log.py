"""One workbook summary from actual receipts. Missing values stay blank."""

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def _details(value, prefix=""):
    if isinstance(value,dict):
        for key,item in value.items():
            yield from _details(item,prefix+str(key)+".")
    elif isinstance(value,list):
        for index,item in enumerate(value):
            yield from _details(item,prefix+str(index)+".")
    else:
        yield prefix.rstrip("."),value


def export_log(scene, qa, output, run_info=None):
    output = Path(output)
    if output.exists():
        raise FileExistsError(output)
    workbook = Workbook()
    objects = workbook.active
    objects.title = "Objects"
    objects.append(["object_id","type","content_revision","structure_revision","style_revision","placement_revision","content_lock","structure_lock","style_lock","placement_lock","status"])
    for o in scene["objects"]:
        objects.append([o["id"],o["semantic_type"],*[o["versions"][f] for f in ("content","structure","style","placement")],
                        *[o["locks"][f] for f in ("content","structure","style","placement")],qa["status"]])
    measurements = workbook.create_sheet("Measurements")
    measurements.append(["object_or_window_id","metric_id","raw_value","normalized_score","threshold","status","evidence_ref","category"])
    for t in qa["windows"]:
        measurements.append([t["id"],"foreground_iou",t["foreground_iou"],None,None,t["status"],"qa.json#/windows/"+t["id"]])
    for c in qa["structure"]["checks"]:
        measurements.append([c["object_id"],c["check_id"],None,None,None,c["status"],c["evidence_ref"]])
    for item in qa.get("local_measurements",[]):
        measurements.append([item["object_or_window_id"],item["metric_id"],item["raw_value"],
                             item.get("normalized_score"),item.get("threshold"),item["status"],item["evidence_ref"],item.get("category")])
    global_ = workbook.create_sheet("Global")
    global_.append(["dimension","weight","score","status","raw_rubric_level_0_to_3","raw_rubric_level_divided_by_3"])
    for k,w in qa["global_weights"].items():
        levels={item["level"] for report in qa.get("visual_reviews",[]) for item in report["observations"]
                if item["target_id"]=="global:"+k and item["criterion"]==k and item["level"] is not None}
        level=next(iter(levels)) if len(levels)==1 else None
        global_.append([k,w,qa["global_dimensions"][k],"UNVERIFIED" if qa["global_dimensions"][k] is None else qa["global_aesthetic"]["status"],
                        level,level/3 if level is not None else None])
    global_.append(["total",None,qa["global_aesthetic"]["score"],qa["global_aesthetic"]["status"]])
    run = workbook.create_sheet("Run")
    run.append(["field","value"])
    run.append(["status",qa["status"]])
    for key,value in (run_info or {}).items():
        run.append([key,value])
    for key,value in qa["binding"].items():
        if isinstance(value,(str,int,float)):
            run.append([key,value])
    usage=qa.get('usage_report',{})
    run.append(["token_usage_status",usage.get('status','PROVIDER_USAGE_UNAVAILABLE')])
    for field in ('input_tokens','cached_input_tokens','output_tokens','reasoning_tokens','total_tokens'):
        run.append([field,usage.get('token_usage',{}).get(field,{}).get('value')])
    if qa.get("visual_reviews"):
        reviews=workbook.create_sheet("Reviews")
        reviews.append(["target_id","criterion","rubric_level_0_to_3","raw_level_divided_by_3",
                        "evaluator","evaluator_version","status","note","evidence_ref","review_sha256"])
        for report in qa["visual_reviews"]:
            for item in report["observations"]:
                reviews.append([item["target_id"],item["criterion"],item["level"],item["normalized_rubric"],
                                report["evaluator"]["id"],report["evaluator"]["version"],item["status"],item["note"],
                                report["artifact"]["path"]+"#/observations",report["artifact"]["sha256"]])
        coverage=qa.get("review_coverage",{})
        run.append(["review_planned_targets",coverage.get("planned_target_count")])
        run.append(["review_covered_targets",coverage.get("covered_target_count")])
        run.append(["review_missing_criteria",len(coverage.get("missing",[]))])
        run.append(["review_conflicts",len(coverage.get("conflicts",[]))])
    if qa.get("native_verification"):
        native=workbook.create_sheet("Native")
        native.append(["object_id","action","check_status","binding_verified","detail_field","detail_value","evidence_ref","report_sha256"])
        receipt=qa["native_verification"]
        for check in receipt["report"].get("checks",[]):
            details=list(_details(check.get("detail",{}))) or [("",None)]
            for key,value in details:
                native.append([check["object_id"],check.get("action"),check["status"],receipt["binding_verified"],
                               key,value,receipt["artifact"]["path"]+"#/checks",receipt["artifact"]["sha256"]])
        for key in ("source_pptx_hash","edited_pptx_hash","operations_hash","office_version","office_build","plan_hash"):
            run.append(["native_"+key,receipt["report"].get(key)])
        run.append(["native_binding_verified",receipt["binding_verified"]])
    if qa.get("stage_history"):
        stages=workbook.create_sheet("Stages")
        stages.append(["kind","stage_id","name","status","started_at","ended_at","elapsed_seconds","cpu_seconds",
                       "candidate","scene_revision","evidence_ref","artifact_sha256"])
        for history in qa["stage_history"]:
            record=history["record"]
            stages.append([history["kind"],record.get("id",record.get("task_id")),record.get("name"),record.get("status"),
                           record.get("started_at"),record.get("ended_at"),record.get("elapsed_seconds"),record.get("cpu_seconds"),
                           record.get("candidate"),record.get("scene_revision"),history["evidence_ref"],history["artifact"]["sha256"]])
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF",bold=True)
            cell.fill = PatternFill("solid",fgColor="17324D")
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type="s"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(70,max(16,max(len(str(c.value or "")) for c in column)+2))
    output.parent.mkdir(parents=True,exist_ok=True)
    workbook.save(output)
