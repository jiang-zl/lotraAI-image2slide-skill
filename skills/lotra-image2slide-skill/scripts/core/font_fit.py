"""Native font candidate gallery and fixed isolated-glyph comparison.

This is a screening experiment, not a replacement for full-slide rendering.
No font files are copied into the Skill or delivery.
"""
import copy
from pathlib import Path
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from .common import digest,identity,read_json,write_json,safe_path,raster_rgb
from .build import text_body
from .typed_metrics import image_measure,DEFAULT_PROTOCOL


def prepare(scene,base,object_id,roi,variants,output,quality_unit=None):
    output=Path(output)
    if output.exists():raise FileExistsError(output)
    obj=next(o for o in scene['objects'] if o['id']==object_id)
    if obj['render_type']!='text' or any('ast' in r for p in obj['content'].get('paragraphs',[]) for r in p):raise ValueError('Font gallery currently isolates plain native text; inline math stays on the full native route')
    if not isinstance(variants,list) or any(not isinstance(v,dict) for v in variants):
        raise ValueError('Font gallery variants must be typed style objects')
    styles=[copy.deepcopy(obj['style'])]+[dict(obj['style'],**v) for v in variants]
    for style in styles:
        if style.get('valign','top') not in ('top','middle','bottom'):
            raise ValueError('Font gallery style.valign must be top, middle or bottom')
    source=safe_path(base,scene['source']['path'])
    if digest(source)!=scene['source']['sha256']:raise ValueError('Source changed')
    image=raster_rgb(Image.open(source),scene['canvas'].get('background_rgb',[255,255,255]));x,y,w,h=roi
    sample=image.crop((x,y,x+w,y+h));data=np.asarray(sample);bg=np.median(data.reshape(-1,3),axis=0).astype(int).tolist()
    quality=None
    if quality_unit is not None:
        from .source_quality import _mask
        from .source_ocr import source_literal
        if (quality_unit['profile']!='text' or quality_unit['objects']!=[object_id]
                or list(roi)!=quality_unit['roi'] or source_literal(quality_unit) is None):
            raise ValueError('Native gallery needs the exact confirmed original word unit')
        quality={'unit':quality_unit['id'],'unit_hash':identity(quality_unit),
            'protocol':copy.deepcopy(quality_unit['protocol']),
            'metrics':copy.deepcopy(quality_unit['metrics']),
            'literal':source_literal(quality_unit),'masks':{}}
        if quality['protocol'].get('background_rgb'):
            bg=list(quality['protocol']['background_rgb'])
        # This gallery has only this one actual native word on each page.
        # Source masks still keep the exact original acceptance window.
        quality['protocol']['candidate_native_text_domain']=False
        for key in ('reference_mask','evaluation_mask'):
            mask=_mask(base,quality_unit.get(key))
            if mask is not None:
                if mask.shape!=(round(h),round(w)):raise ValueError('Use the original word-mask frame')
                quality['masks'][key]=mask
    pad=40;cw=max(100,w+pad*2);ch=max(100,h+pad*2)
    reference=Image.new('RGB',(cw,ch),tuple(bg));reference.paste(sample,(pad,pad))
    output.mkdir(parents=True);reference.save(output/'reference.png')
    if quality is not None:
        for key,mask in list(quality['masks'].items()):
            target=output/(key+'.png');Image.fromarray(mask.astype('uint8')*255).save(target)
            quality['masks'][key]={'path':target.name,'sha256':digest(target)}
    scale=scene['canvas']['target_width_inches']*914400/scene['canvas']['width']
    prs=Presentation();prs.slide_width=round(cw*scale);prs.slide_height=round(ch*scale)
    unique=[]
    for s in styles:
        if s not in unique:unique.append(s)
    for i,style in enumerate(unique):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(*bg)
        bx,by,bw,bh=obj['bbox'];dx,dy=style.get('bbox_offset_source_px',[0,0]);bx+=dx;by+=dy
        if abs(dx)>4 or abs(dy)>4:raise ValueError('Native glyph fitting is bounded by its located source evidence')
        shape=slide.shapes.add_textbox(round((bx-x+pad)*scale),round((by-y+pad)*scale),round(bw*scale),round(bh*scale));shape.name=object_id
        text_body(shape,obj['content'],style)
    prs.save(output/'Font_candidates.pptx')
    record={'version':'native-font-screen-1','source_hash':digest(source),'source_roi':roi,'object_id':object_id,
            'scene_hash':identity(scene),'original_style':obj['style'],'candidate_styles':unique,
            'reference_hash':digest(output/'reference.png'),'pptx_hash':digest(output/'Font_candidates.pptx'),
            'canvas':[cw,ch],'background_rgb':bg,'padding':pad,
            'scope':'Fixed cropped source glyphs on a native solid-background gallery; full context must be rendered after selection',
            'quality_unit':quality}
    write_json(output/'gallery.json',record);return record


def rank(directory,render_directory,output):
    directory,render_directory=map(Path,(directory,render_directory));plan=read_json(directory/'gallery.json');rr=read_json(render_directory/'render.json')
    if plan['pptx_hash']!=digest(directory/'Font_candidates.pptx') or plan['reference_hash']!=digest(directory/'reference.png') or rr['pptx_hash']!=plan['pptx_hash']:raise ValueError('Stale font gallery evidence')
    reference=Image.open(directory/'reference.png');results=[]
    quality=plan.get('quality_unit')
    protocol=(quality['protocol'] if quality else
              dict(DEFAULT_PROTOCOL,background_rgb=plan['background_rgb'],foreground_distance_rgb=.45))
    masks={}
    if quality:
        from .source_quality import _mask
        masks={key:_mask(directory,value) for key,value in quality['masks'].items()}
    saved=Presentation(directory/'Font_candidates.pptx') if quality else None
    a=np.asarray(reference.convert('RGB')).astype(float);a_mask=np.mean(a,axis=2)<128
    # Resolve fonts from saved rendered PDF: a requested font name is not proof
    # that the renderer used it (notably Aptos in bundled LibreOffice).
    import re
    import pymupdf
    pdf_record=rr.get('pdf') or {}
    pdf=Path(rr.get('pdf_path') or pdf_record.get('path') or 'Final.pdf')
    if not pdf.is_absolute():pdf=render_directory/pdf
    recorded_pdf_hash=rr.get('pdf_hash') or pdf_record.get('sha256')
    actual_pdf_hash=digest(pdf) if pdf.exists() else None
    if actual_pdf_hash and recorded_pdf_hash and actual_pdf_hash!=recorded_pdf_hash:
        raise ValueError('Font PDF differs from its saved render receipt')
    pdf_bound=bool(actual_pdf_hash and recorded_pdf_hash)
    pdf_evidence={'path':str(pdf),'sha256':actual_pdf_hash,
        'status':('BOUND_SAVED_PDF' if pdf_bound else 'UNBOUND_PDF' if actual_pdf_hash
                  else pdf_record.get('status') or 'MISSING_PDF'),
        'reason':pdf_record.get('reason'),
        'limits':'A missing or failed PDF export is not evidence of font substitution or rasterized glyphs.'}
    doc=pymupdf.open(pdf) if pdf_bound else None
    def family(s):
        return re.sub(r'[^a-z0-9]','',re.sub(r'(?i)(bold|italic|regular|semibold|light|medium)','',s))
    for i,style in enumerate(plan['candidate_styles']):
        path=render_directory/('Final_preview.png' if i==0 else 'Page_'+str(i+1)+'.png')
        if not path.exists():raise ValueError('Unrendered font candidate')
        image=Image.open(path).convert('RGB')
        # One fixed canvas mapping, not candidate registration. Higher render
        # resolution is normalized to the gallery's original source canvas.
        if abs(image.width/image.height-reference.width/reference.height)>1/reference.height:
            raise ValueError('Font render is outside its declared canvas')
        if image.size!=reference.size:image=image.resize(reference.size,Image.Resampling.LANCZOS)
        b=np.asarray(image).astype(float);b_mask=np.mean(b,axis=2)<128
        # Padding is only a native-host drawing viewport. It cannot dilute
        # glyph weight/quality or select a style that fails in the real ROI.
        pad=plan['padding'];w,h=plan['source_roi'][2:]
        window=(pad,pad,pad+round(w),pad+round(h))
        measure=image_measure(reference.crop(window),image.crop(window),protocol,
                              masks.get('reference_mask'),masks.get('evaluation_mask'))
        constraints=[]
        if quality:
            import unicodedata
            from .typed_metrics import scalar_vector
            from .source_quality import assess
            owned=[shape for shape in saved.slides[i].shapes if shape.name==plan['object_id']]
            actual=owned[0].text if len(owned)==1 and owned[0].has_text_frame else None
            norm=lambda value:''.join(unicodedata.normalize('NFKC',value).split())
            values=scalar_vector(measure)
            values.update(foreground_f1=measure['foreground_match']['f1'],
                foreground_mass_ratio=measure['foreground_mass_ratio'],
                foreground_area_ratio=measure['foreground_area_ratio'],
                color_histogram_tv=measure['texture_color_histogram_tv'],
                native_text_match=float(norm(actual)==norm(quality['literal'])) if actual is not None else None)
            constraints=assess(values,quality['metrics'])
        xor=int((a_mask[pad:pad+round(h),pad:pad+round(w)]^b_mask[pad:pad+round(h),pad:pad+round(w)]).sum())
        page_available=doc is not None and i<len(doc)
        fonts=sorted({s['font'] for b in doc[i].get_text('dict')['blocks'] if 'lines' in b for line in b['lines'] for s in line['spans']}) if page_available else []
        exact=bool(fonts) and all(family(f)==family(style['font']) for f in fonts)
        results.append({'index':i,'style':style,'glyph_pixel_disagreement':xor,'measurements':measure,'preview_hash':digest(path),'status':'UNVERIFIED',
                        'rendered_fonts':fonts,'font_match':'MATCH' if exact else 'SUBSTITUTED_OR_UNVERIFIED',
                        'font_evidence_status':('REQUESTED_FAMILY_MATCH' if exact else 'SAVED_PDF_OTHER_FAMILY' if fonts
                            else 'SAVED_PDF_NO_TEXT_SPANS' if page_available else 'PDF_PAGE_MISSING' if doc is not None
                            else pdf_evidence['status']),
                        'source_constraints':constraints})
    if doc is not None:doc.close()
    # PDF ToUnicode/family aliases can be broken while the actual Office glyphs
    # are correct. They stay provenance diagnostics, never an appearance oracle.
    eligible=[r for r in results if r.get('measurements')]
    def key(r):
        if quality:
            deficits=[c['deficit'] if c['deficit'] is not None else 1. for c in r['source_constraints']]
            return (max(deficits,default=1.),sum(deficits),r['glyph_pixel_disagreement'],r['index'])
        m=r['measurements'];mass=m['foreground_mass_ratio'];presence=m['foreground_match']['f1']
        deficits=[max(0,.85-mass)/.85 if mass is not None else 1,
                  max(0,mass-1.2)/1.2 if mass is not None else 1,
                  max(0,.95-presence)/.95]
        return (max(deficits),sum(deficits),r['glyph_pixel_disagreement'],r['index'])
    selected=min(eligible or results,key=key)
    result={'plan_hash':identity(plan),'source_hash':plan['source_hash'],'object_id':plan['object_id'],'candidates':results,
            'selected_index':selected['index'],'selected_style':selected['style'],'status':'SCREENED_REQUIRES_FULL_CONTEXT',
            'objective':'noncompensating glyph presence/weight screening, then fixed pixel disagreement; full-context source targets still apply',
            'quality_unit':{'id':quality['unit'],'hash':quality['unit_hash']} if quality else None,
            'measurement_basis':'fixed_original_word_targets' if quality else 'legacy_screen_proxy',
            'selected_family_unchanged':selected['style']['font']==plan['original_style']['font'],
            'font_selection_verified':False,'font_warning':'PDF family labels do not certify actual glyph typefaces. Ranking screens actual glyphs; full-context saved-native evaluation is mandatory.',
            'pdf_evidence':pdf_evidence,
            'renderer':{k:rr.get(k) for k in ('backend','backend_version','office_build')},'render_receipt_hash':digest(render_directory/'render.json')}
    write_json(output,result);return result
