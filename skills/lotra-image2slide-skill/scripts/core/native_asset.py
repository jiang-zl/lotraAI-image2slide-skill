"""Reuse the actual saved native artwork, including user edits, as a symbol.

The recipe is optional. A snapshot is data; no code, image substitution or
stale reverse-conversion is involved. Ancestor coordinate frames are retained.
"""
import copy
from pathlib import Path
from pptx import Presentation
from .common import digest, safe_path
from .mathml import NS
from .saved_components import inventory, _tag, _id


def load(base, reference,with_source=False):
    path=safe_path(base,reference['path'])
    if digest(path)!=reference['sha256']:raise ValueError('Native asset snapshot changed')
    deck=Presentation(path);number=reference.get('slide',1)
    slide=deck.slides[number-1]
    rows=inventory(path,number)['objects']
    found=[r for r in rows if r['name']==reference['root_name']]
    if len(found)!=1 or found[0]['kind'] not in ('grpSp','pic','sp'):
        raise ValueError('A reusable component needs one actual named artwork root')
    row=found[0];element=next(e for e in slide.shapes._spTree.iter()
                            if _tag(e) in ('grpSp','pic','sp') and _id(e)==row['shape_id'])
    clone=copy.deepcopy(element)
    # Preserve scaled/rotated ancestors while pruning all unrelated siblings.
    ancestor=element.getparent()
    while _tag(ancestor)=='grpSp':
        wrapper=copy.deepcopy(ancestor)
        for sub in list(wrapper):
            if _tag(sub) in ('sp','grpSp','pic','cxnSp','graphicFrame'):wrapper.remove(sub)
        wrapper.append(clone);clone=wrapper;ancestor=ancestor.getparent()
    forbidden=('{'+NS['p']+'}graphicFrame',
               '{'+NS['a']+'}hlinkClick','{'+NS['a']+'}hlinkHover')
    if next(clone.iter(*forbidden),None) is not None:
        raise ValueError('Saved artwork cannot hide active or external content')
    for element in clone.iter():
        for key,value in element.attrib.items():
            if key.startswith('{'+NS['r']+'}'):
                from pptx.opc.constants import RELATIONSHIP_TYPE as RT
                rel=slide.part.rels[value]
                from .vector_asset import MASTER_REL,MASTER_NS,master_bytes
                if rel.is_external or rel.reltype not in (RT.IMAGE,MASTER_REL):
                    raise ValueError('Saved artwork depends on unsupported external content')
                if rel.reltype==MASTER_REL:
                    from hashlib import sha256
                    part=rel.target_part
                    if element.tag!='{'+MASTER_NS+'}vectorMaster' or part.content_type not in ('application/xml','application/gzip'):
                        raise ValueError('Saved vector master must remain explicitly bound pure source SVG')
                    master_bytes(part.blob,element.get('compression'),element.get('sha256'))
    ids={_id(e) for e in clone.iter(*('{'+NS['p']+'}'+tag for tag in ('sp','grpSp','cxnSp','pic')))}
    for connection in clone.iter('{'+NS['a']+'}stCxn','{'+NS['a']+'}endCxn'):
        if int(connection.get('id')) not in ids:
            raise ValueError('Capture external connections separately before native symbol reuse')
    return (clone,row['bbox_emu'],slide.part) if with_source else (clone,row['bbox_emu'])


def instantiate(group,base,reference,bbox_emu,prefix,allocate):
    clone,source_box,source=load(base,reference,with_source=True)
    from .resource_transfer import transfer
    transfer(clone,source,group.part)
    mapping={}
    for element in clone.iter(*('{'+NS['p']+'}'+tag for tag in ('sp','grpSp','cxnSp','pic'))):
        props=element.find('.//p:cNvPr',NS);old=props.get('id');new=allocate()
        mapping[old]=str(new);props.set('id',str(new))
        props.set('name',prefix+'__saved_'+old)
    for connection in clone.iter('{'+NS['a']+'}stCxn','{'+NS['a']+'}endCxn'):
        connection.set('id',mapping[connection.get('id')])
    group._element.append(clone)
    xf=group._element.find('p:grpSpPr/a:xfrm',NS)
    for tag,values in [('off',bbox_emu[:2]),('ext',bbox_emu[2:]),
                       ('chOff',source_box[:2]),('chExt',source_box[2:])]:
        item=xf.find('a:'+tag,NS);keys=('x','y') if tag in ('off','chOff') else ('cx','cy')
        for key,value in zip(keys,values):item.set(key,str(round(value)))
    return len(mapping)
