"""Source-relative spacing of observed ink masks, independent of shape frames.

Mask provenance still needs semantic review. This checker is useful only when
the two objects' masks correspond to actual source/rendered visible ink.
"""

import math
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation
from scipy.ndimage import distance_transform_edt

from .common import digest, read_json, safe_path, write_json
from .render import verify_alternative_render
from .saved_ppt_diff import snapshot


def _bound_file(base, reference):
    if not isinstance(reference, dict) or set(reference) != {"path", "sha256"}:
        raise ValueError("Spacing input needs a frozen path and SHA-256")
    path = safe_path(base, reference["path"])
    if digest(path) != reference["sha256"]:
        raise ValueError("Frozen spacing input changed: " + reference["path"])
    return path


def _mask(base, reference, size):
    path = _bound_file(base, reference)
    with Image.open(path) as image:
        if image.mode != "L" or image.size != size:
            raise ValueError("Ink mask must be grayscale and match the fixed canvas")
        pixels = np.asarray(image)
    if not np.isin(pixels, [0, 255]).all() or not np.any(pixels):
        raise ValueError("Ink mask must be nonempty and binary")
    return pixels == 255


def _distance(first, second):
    overlap = int(np.count_nonzero(first & second))
    if overlap:
        return {"gap_px": 0.0, "overlap_pixels": overlap}
    centers = float(distance_transform_edt(~first)[second].min())
    return {"gap_px": max(0.0, centers - 1.0), "overlap_pixels": 0}


def _saved_bounds_within(pptx, saved, paths, roi, size, projected=None):
    if (not isinstance(paths, list) or not paths or
            any(not isinstance(path, str) or path not in saved["objects"] for path in paths)
            or not isinstance(roi, list) or len(roi) != 4 or
            any(type(value) is not int for value in roi)):
        raise ValueError("Candidate ink owner needs saved object paths and a fixed extraction ROI")
    x0, y0, x1, y1 = roi
    if not (0 <= x0 < x1 <= size[0] and 0 <= y0 < y1 <= size[1]):
        raise ValueError("Candidate ink extraction ROI is outside the preview")
    deck = Presentation(pptx)
    for path in paths:
        item = saved["objects"][path]
        box = (projected or {}).get(path)
        if box is None:
            return False  # Never treat child-local coordinates as slide bounds.
        left, top = box[0] / deck.slide_width * size[0], box[1] / deck.slide_height * size[1]
        right = (box[0] + box[2]) / deck.slide_width * size[0]
        bottom = (box[1] + box[3]) / deck.slide_height * size[1]
        if left < x0 - 1 or top < y0 - 1 or right > x1 + 1 or bottom > y1 + 1:
            return False
    return True


def check(contract, base):
    if contract.get("schema_version") != "ink-spacing-1":
        raise ValueError("Unknown ink spacing contract")
    source = _bound_file(base, contract["source_preview"])
    candidate = _bound_file(base, contract["candidate_preview"])
    with Image.open(source) as image:
        size = image.size
    with Image.open(candidate) as image:
        if image.size != size:
            raise ValueError("Source and candidate ink canvases differ")
    render_status = "MASK_ONLY_DIAGNOSTIC"
    saved_pptx, saved_objects = None, None
    if contract.get("candidate_pptx") or contract.get("candidate_render_receipt"):
        pptx = _bound_file(base, contract["candidate_pptx"])
        receipt = _bound_file(base, contract["candidate_render_receipt"])
        if candidate != receipt.parent / "Final_preview.png":
            raise ValueError("Candidate ink preview is not the saved PPT render output")
        host = read_json(receipt).get("backend")
        if host == "PowerPoint":
            from .host_render import verify_render
            verify_render(pptx, candidate, receipt)
            render_status = "VERIFIED_POWERPOINT_RENDER"
        else:
            verify_alternative_render(pptx, receipt)
            render_status = "VERIFIED_ALTERNATIVE_RENDER"
        saved_pptx, saved_objects = pptx, snapshot(pptx)
        # Compose real native group rotations/flips/scales once. Paths keep
        # duplicate names in different parents distinct.
        from .saved_components import inventory
        deck=Presentation(pptx);projected={}
        for number in range(1,len(deck.slides)+1):
            records=inventory(pptx,number,presentation=deck)["objects"]
            by_id={row["shape_id"]:row for row in records}
            def native_path(row):
                parent=row["parent_id"]
                return (native_path(by_id[parent]) if parent is not None else "slide-"+str(number))+"/"+row["name"]
            projected.update({native_path(row):row["bbox_emu"] for row in records})
    pairs = contract.get("pairs")
    if (not isinstance(pairs, list) or not pairs or
            any(not isinstance(p, dict) or not isinstance(p.get("id"), str) or not p["id"] for p in pairs)
            or len({p["id"] for p in pairs}) != len(pairs)):
        raise ValueError("Distinct source-relative spacing pairs are required")
    results = []
    all_bound_to_saved_owners = True
    for pair in pairs:
        if (pair.get("intent") not in ("separate", "observed_overlap")
                or not isinstance(pair.get("mask_provenance"), str)
                or not pair["mask_provenance"].strip()
                or any(type(pair.get(key)) not in (int, float) or not math.isfinite(pair[key]) or pair[key] < 0
                       for key in ("max_crowding_px", "max_expansion_px"))
                or type(pair.get("max_overlap_increase_pixels")) is not int
                or pair["max_overlap_increase_pixels"] < 0):
            raise ValueError("Pair needs source-observed intent, mask provenance and frozen tolerances")
        masks = pair.get("masks", {})
        if set(masks) != {"source_a", "source_b", "candidate_a", "candidate_b"}:
            raise ValueError("Both entities need source and candidate ink masks")
        source_a, source_b, candidate_a, candidate_b = (
            _mask(base, masks[key], size) for key in
            ("source_a", "source_b", "candidate_a", "candidate_b"))
        old = _distance(source_a, source_b)
        new = _distance(candidate_a, candidate_b)
        failures = []
        owner_binding = pair.get("candidate_object_paths")
        roi_binding = pair.get("candidate_mask_roi_xyxy")
        if owner_binding is None and roi_binding is None:
            all_bound_to_saved_owners = False
        elif saved_pptx is None or not isinstance(owner_binding, dict) or not isinstance(roi_binding, dict) or set(owner_binding) != {"a", "b"} or set(roi_binding) != {"a", "b"}:
            raise ValueError("Saved-object ink bounds need both owner paths and fixed extraction ROIs")
        else:
            for entity in ("a", "b"):
                if not _saved_bounds_within(saved_pptx, saved_objects,
                                            owner_binding[entity], roi_binding[entity], size, projected):
                    failures.append("CANDIDATE_ENTITY_ESCAPES_FROZEN_INK_ROI:" + entity)
        if pair["intent"] == "separate":
            if old["overlap_pixels"]:
                raise ValueError("Declared separate source masks already overlap")
            if new["overlap_pixels"]:
                failures.append("NEW_INK_COLLISION")
            elif new["gap_px"] + pair["max_crowding_px"] < old["gap_px"] - 1e-8:
                failures.append("NEW_CROWDING_RELATIVE_TO_SOURCE")
            elif new["gap_px"] > old["gap_px"] + pair["max_expansion_px"] + 1e-8:
                failures.append("SPACING_EXPANDED_RELATIVE_TO_SOURCE")
        else:
            if old["overlap_pixels"] == 0:
                raise ValueError("Observed-overlap control has no source ink overlap")
            if new["overlap_pixels"] > old["overlap_pixels"] + pair["max_overlap_increase_pixels"]:
                failures.append("INTENTIONAL_OVERLAP_EXCESSIVELY_EXPANDED")
        results.append({"id": pair["id"], "intent": pair["intent"],
                        "source": old, "candidate": new,
                        "status": "FAIL" if failures else "PASS_PAIR_ONLY",
                        "failures": failures,
                        "mask_provenance": pair["mask_provenance"]})
    failed = any(item["status"] == "FAIL" for item in results)
    return {"schema_version": "ink-spacing-report-1",
            "status": ("FAIL" if failed else "PASS_POWERPOINT_RENDER_SPACING_ONLY"
                       if render_status == "VERIFIED_POWERPOINT_RENDER" and all_bound_to_saved_owners
                       else "PASS_ALTERNATIVE_RENDER_SPACING_ONLY"
                       if render_status == "VERIFIED_ALTERNATIVE_RENDER" and all_bound_to_saved_owners
                       else "MASK_DIAGNOSTIC_ONLY"),
            "render_status": render_status, "source_preview_hash": digest(source),
            "candidate_preview_hash": digest(candidate), "canvas": list(size),
            "pairs": results,
            "scope": "actual per-object ink-mask gaps; semantic mask quality, full layout and target-editor behavior separate"}


def check_file(contract_file, output):
    contract_file = Path(contract_file)
    result = check(read_json(contract_file), contract_file.parent)
    write_json(output, result)
    return result
