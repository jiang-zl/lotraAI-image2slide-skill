"""Bounded, source-derived variants for illustrative photo stacks.

The host must identify the donor's category from the source. This operation
changes appearance within that category; it does not prove semantic realism or
recover the hidden original. Scientific/data imagery is deliberately excluded.
"""

from itertools import combinations
from pathlib import Path
import io
import hashlib
import math
import numpy as np
from PIL import Image
from scipy.ndimage import map_coordinates
from scipy.ndimage import binary_dilation
from skimage.metrics import structural_similarity
from pptx import Presentation
from .common import digest, identity, read_json, safe_path, write_json
from .saved_components import inventory


def _image(path):
    with Image.open(path) as image:
        return np.asarray(image.convert('RGB'), dtype=np.uint8)


def _similarity(a, b):
    if a.shape != b.shape:
        a = np.asarray(Image.fromarray(a).resize((64, 64), Image.Resampling.BICUBIC))
        b = np.asarray(Image.fromarray(b).resize((64, 64), Image.Resampling.BICUBIC))
    return float(structural_similarity(a, b, data_range=255, channel_axis=2))


def _lineage(layer, base, asset_hash):
    """Resolve a verifiable donor chain; an image hash alone is not an identity."""
    root = asset_hash
    method = 'unverified_direct_asset'
    completion_ref = layer.get('completion_receipt')
    if completion_ref:
        receipt_path = safe_path(base, completion_ref['path'])
        if digest(receipt_path) != completion_ref['sha256']:
            raise ValueError('Completion receipt changed')
        receipt = read_json(receipt_path)
        if receipt.get('output_hash') != asset_hash or not receipt.get('source_hash'):
            raise ValueError('Completion receipt is not bound to the layer asset')
        root = receipt['source_hash']
        method = receipt.get('method', 'unknown_completion')
    variant_ref = layer.get('variant_receipt')
    if variant_ref:
        receipt_path = safe_path(base, variant_ref['path'])
        if digest(receipt_path) != variant_ref['sha256']:
            raise ValueError('Variant receipt changed')
        receipt = read_json(receipt_path)
        expected = read_json(safe_path(base, completion_ref['path'])).get('candidate_hash') if completion_ref else asset_hash
        if receipt.get('output_hash') != expected or not receipt.get('source_hash'):
            raise ValueError('Variant receipt is not bound to the candidate image')
        if completion_ref and receipt['source_hash'] != root:
            raise ValueError('Variant donor and completed source disagree')
        root = receipt['source_hash']
        method = receipt.get('method', 'unknown_variant')
    return {'root_donor_sha256': root, 'method': method,
            'receipt_bound': bool(completion_ref or variant_ref)}


def diversity(layers, base, *, near_threshold=0.9, repetition_policy=None,
              policy_evidence_refs=None, semantic_review=None):
    """Measure appearance and provenance without mistaking affine change for a new instance.

    `distinct_samples` rejects exact copies, candidate-observed same identities
    and shared bounded-transform donors. A front card does not establish the
    hidden rear cards' source categories: `category` may describe a broad stack
    such as training samples, while per-card source and candidate observations
    remain separate. `observed_repeat` permits repetition only with located
    evidence; neither policy produces semantic PASS itself.
    """
    if not isinstance(near_threshold, (int, float)) or not 0.5 < near_threshold < 1:
        raise ValueError('Near-duplicate SSIM threshold must be between 0.5 and 1')
    if len(layers) < 2:
        raise ValueError('Diversity needs at least two layers')
    if repetition_policy not in (None, 'distinct_samples', 'observed_repeat'):
        raise ValueError('Unknown photo-stack repetition policy')
    if repetition_policy and (not isinstance(policy_evidence_refs, list) or not policy_evidence_refs
                              or any(not isinstance(ref, str) or not ref for ref in policy_evidence_refs)):
        raise ValueError('Repetition policy needs located source evidence')
    arrays = []
    for layer in layers:
        if layer.get('kind') != 'illustrative_photo' or not layer.get('category'):
            raise ValueError('Each layer needs an illustrative stack category')
        source_status = layer.get('source_subject_status')
        if source_status is not None:
            if source_status not in ('observed', 'inferred', 'unresolved'):
                raise ValueError('Unknown source subject observation status')
            if (not isinstance(layer.get('source_evidence_refs'), list)
                    or not layer['source_evidence_refs']
                    or any(not isinstance(ref, str) or not ref for ref in layer['source_evidence_refs'])
                    or not layer.get('candidate_subject_category')
                    or not isinstance(layer.get('candidate_visual_evidence_refs'), list)
                    or not layer['candidate_visual_evidence_refs']
                    or any(not isinstance(ref, str) or not ref
                           for ref in layer['candidate_visual_evidence_refs'])):
                raise ValueError('Per-card source and candidate observations need located evidence')
            if (source_status == 'observed') != bool(layer.get('source_subject_category')):
                raise ValueError('Only a source-observed card may assert its source subject category')
        path = safe_path(base, layer['asset']['path'])
        if digest(path) != layer['asset']['sha256']:
            raise ValueError('Layer asset changed')
        arrays.append((layer, path, _image(path), _lineage(layer, base, layer['asset']['sha256'])))
    if len({row[0]['category'] for row in arrays}) != 1:
        raise ValueError('Do not mix semantic categories in a diversity group')
    pairs = []
    for left, right in combinations(range(len(arrays)), 2):
        a, b = arrays[left], arrays[right]
        aa, bb = a[2], b[2]
        if aa.shape != bb.shape:
            aa = np.asarray(Image.fromarray(aa).resize((64, 64), Image.Resampling.BICUBIC))
            bb = np.asarray(Image.fromarray(bb).resize((64, 64), Image.Resampling.BICUBIC))
        left_identity, right_identity = a[0].get('semantic_instance_id'), b[0].get('semantic_instance_id')
        pairs.append({'left': left, 'right': right,
                      'exact_duplicate': digest(a[1]) == digest(b[1]) or bool(np.array_equal(a[2], b[2])),
                      'ssim': _similarity(a[2], b[2]),
                      'mean_rgb_difference_255': float(np.abs(aa.astype(float) - bb.astype(float)).mean()),
                      'shared_donor_lineage': a[3]['root_donor_sha256'] == b[3]['root_donor_sha256'],
                      'same_declared_semantic_instance': bool(left_identity and right_identity and left_identity == right_identity),
                      'bounded_same_donor_transform':
                          a[3]['root_donor_sha256'] == b[3]['root_donor_sha256'] and
                          'bounded_source_affine_and_tone' in (a[3]['method'], b[3]['method'])})
    findings = []
    policy_status = 'NOT_DECLARED'
    if repetition_policy == 'distinct_samples':
        if any(pair['exact_duplicate'] for pair in pairs):
            findings.append('EXACT_DUPLICATE_IN_DISTINCT_SAMPLES')
        if any(pair['same_declared_semantic_instance'] for pair in pairs):
            findings.append('SAME_DECLARED_INSTANCE_IN_DISTINCT_SAMPLES')
        if any(pair['bounded_same_donor_transform'] for pair in pairs):
            findings.append('BOUNDED_TRANSFORM_IS_NOT_NEW_INSTANCE')
        if any(not layer[0].get('semantic_instance_id') for layer in arrays):
            findings.append('INSTANCE_IDENTITY_NOT_OBSERVED')
        if any(pair['shared_donor_lineage'] and not pair['bounded_same_donor_transform'] for pair in pairs):
            findings.append('SHARED_DONOR_REQUIRES_SEMANTIC_REVIEW')
        policy_status = 'FAIL' if any(item in findings for item in (
            'EXACT_DUPLICATE_IN_DISTINCT_SAMPLES',
            'SAME_DECLARED_INSTANCE_IN_DISTINCT_SAMPLES',
            'BOUNDED_TRANSFORM_IS_NOT_NEW_INSTANCE')) else 'REQUIRES_SEMANTIC_REVIEW'
    elif repetition_policy == 'observed_repeat':
        policy_status = 'DECLARED_REPEAT_REQUIRES_SEMANTIC_REVIEW'
    semantic_status = 'UNVERIFIED'
    review_hash = None
    if semantic_review is not None:
        if repetition_policy is None:
            raise ValueError('Semantic diversity review requires a source-grounded repetition policy')
        review = read_json(semantic_review)
        review_hash = digest(semantic_review)
        expected_policy = identity({'repetition_policy': repetition_policy,
                                    'policy_evidence_refs': policy_evidence_refs or []})
        if (review.get('schema_version') != 'illustrative-diversity-review-1'
                or review.get('asset_hashes') != [entry[0]['asset']['sha256'] for entry in arrays]
                or review.get('policy_hash') != expected_policy
                or review.get('reviewer', {}).get('kind') not in ('human', 'host_model')):
            raise ValueError('Semantic diversity review is not bound to these separated assets and source policy')
        cards = review.get('cards', [])
        reviewed_pairs = review.get('pairs', [])
        if (len(cards) != len(arrays) or len(reviewed_pairs) != len(pairs)
                or [card.get('index') for card in cards] != list(range(len(arrays)))
                or [(pair.get('left'), pair.get('right')) for pair in reviewed_pairs] !=
                    [(pair['left'], pair['right']) for pair in pairs]):
            raise ValueError('Every separated card and pair needs one located semantic review')
        statuses = []
        for index, card in enumerate(cards):
            if (card.get('candidate_subject_category') != arrays[index][0].get('candidate_subject_category')
                    or card.get('whole_subject_continuity') not in ('PASS', 'FAIL', 'UNVERIFIED')
                    or card.get('single_subject') not in ('PASS', 'FAIL', 'UNVERIFIED')
                    or not card.get('evidence_refs')):
                raise ValueError('Card review lacks whole-subject observation or located evidence')
            statuses.extend((card['whole_subject_continuity'], card['single_subject']))
        for pair in reviewed_pairs:
            decision = pair.get('status')
            expected = 'PASS_DISTINCT' if repetition_policy == 'distinct_samples' else 'PASS_OBSERVED_REPEAT'
            if decision not in (expected, 'FAIL', 'UNVERIFIED') or not pair.get('evidence_refs'):
                raise ValueError('Pair review lacks a policy-specific identity decision and evidence')
            statuses.append('PASS' if decision == expected else decision)
        if any(item in findings for item in ('EXACT_DUPLICATE_IN_DISTINCT_SAMPLES',
                                             'SAME_DECLARED_INSTANCE_IN_DISTINCT_SAMPLES',
                                             'BOUNDED_TRANSFORM_IS_NOT_NEW_INSTANCE')):
            semantic_status = 'FAIL'
        elif 'FAIL' in statuses:
            semantic_status = 'FAIL'
        elif 'UNVERIFIED' in statuses:
            semantic_status = 'UNVERIFIED'
        else:
            semantic_status = 'PASS'
        policy_status = ('PASS_SEMANTIC_INSTANCE_POLICY' if semantic_status == 'PASS'
                         else semantic_status)
    return {'schema_version': 'illustrative-diversity-1', 'category': arrays[0][0]['category'],
            'layer_count': len(arrays), 'exact_duplicate_pairs': sum(p['exact_duplicate'] for p in pairs),
            'near_duplicate_ssim_threshold': float(near_threshold),
            'near_duplicate_pairs': sum(p['ssim'] >= near_threshold for p in pairs),
            'pairs': pairs, 'provenance': [entry[3] for entry in arrays],
            'per_card_observations': [{key: entry[0].get(key) for key in
                                       ('source_subject_status', 'source_subject_category',
                                        'source_evidence_refs', 'candidate_subject_category',
                                        'candidate_visual_evidence_refs', 'semantic_instance_id')}
                                      for entry in arrays],
            'repetition_policy': repetition_policy, 'policy_evidence_refs': policy_evidence_refs or [],
            'policy_status': policy_status, 'policy_findings': findings,
            'semantic_status': semantic_status, 'semantic_review_hash': review_hash,
            'scope': 'Appearance and bound lineage are diagnostics; whole-object host semantic review and PPT context remain required'}


def generate_variant(request, base, output):
    """Make one deterministic, modest affine/color variation of an observed donor.

    The caller supplies independently located same-category evidence. A variant
    must subsequently pass visible-pixel preservation, composition and semantic
    review; this function cannot promote a slide candidate.
    """
    if request.get('schema_version') != 'illustrative-variant-1':
        raise ValueError('Unknown illustrative variant request')
    if request.get('kind') != 'illustrative_photo' or request.get('category') != request.get('donor_category'):
        raise ValueError('Only a class-matched illustrative donor can be varied')
    if not request.get('category') or not request.get('evidence_refs'):
        raise ValueError('Variant needs located category observations')
    seed = request.get('seed')
    if type(seed) is not int or not 0 <= seed < 2**32:
        raise ValueError('Variant needs a bounded deterministic integer seed')
    donor = safe_path(base, request['donor']['path'])
    if digest(donor) != request['donor']['sha256']:
        raise ValueError('Donor changed')
    source = _image(donor)
    height, width = source.shape[:2]
    if min(height, width) < 24 or max(height, width) > 2048:
        raise ValueError('Unsupported illustrative thumbnail dimensions')
    if request.get('target_size') != [width, height]:
        raise ValueError('Variant cannot change the frozen thumbnail canvas')
    rng = np.random.default_rng(seed)
    angle = float(rng.uniform(-6, 6)) * math.pi / 180
    scale = float(rng.uniform(0.89, 0.95))
    x_shift = float(rng.uniform(-0.06, 0.06) * width)
    y_shift = float(rng.uniform(-0.06, 0.06) * height)
    yy, xx = np.mgrid[:height, :width]
    x = xx - (width - 1) / 2 - x_shift
    y = yy - (height - 1) / 2 - y_shift
    sx = (math.cos(angle) * x + math.sin(angle) * y) * scale + (width - 1) / 2
    sy = (-math.sin(angle) * x + math.cos(angle) * y) * scale + (height - 1) / 2
    # Reflect the known image border rather than inventing unobserved content.
    transformed = np.stack([map_coordinates(source[:, :, channel], [sy, sx], order=1, mode='reflect')
                            for channel in range(3)], axis=-1)
    gain = float(rng.uniform(0.965, 1.035))
    bias = float(rng.uniform(-4, 4))
    result = np.clip(transformed.astype(float) * gain + bias, 0, 255).astype(np.uint8)
    similarity = _similarity(source, result)
    exact_fraction = float(np.mean(np.all(source == result, axis=2)))
    if similarity < 0.35 or similarity > 0.98 or exact_fraction > 0.5:
        raise ValueError('Variant is too distorted or too close to an exact copy; choose another observed donor/seed')
    output = Path(output)
    if output.exists():
        raise FileExistsError(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    Image.fromarray(result).save(output)
    receipt = {'schema_version': 'illustrative-variant-result-1', 'category': request['category'],
               'source_hash': digest(donor), 'output_hash': digest(output), 'seed': seed,
               'method': 'bounded_source_affine_and_tone', 'source_ssim': similarity,
               'exact_pixel_fraction': exact_fraction, 'angle_degrees': angle * 180 / math.pi,
               'scale': scale, 'evidence_refs': request['evidence_refs'],
               'semantic_status': 'UNVERIFIED', 'status': 'GENERATED_REQUIRES_COMPOSITION_AND_SEMANTIC_REVIEW'}
    write_json(output.with_suffix('.variant.json'), receipt)
    return receipt


def plan_independent_layers(request, base):
    """Choose distinct full-subject donor routes before any hidden-pixel fill.

    This plan never invents an unobserved source category or calls an image
    model. The host generates or selects full candidates, then the completion
    and separated-card review gates decide whether they are usable.
    """
    from .completion import validate as validate_completion
    if (request.get('schema_version') != 'illustrative-stack-plan-1'
            or request.get('repetition_policy') not in ('distinct_samples', 'observed_repeat')
            or not request.get('policy_evidence_refs')
            or not isinstance(request.get('cards'), list) or not request['cards']):
        raise ValueError('Photo-stack plan needs cards and a source-grounded repetition policy')
    donors = []
    for donor in request.get('donors', []):
        ref = donor.get('asset')
        path = safe_path(base, ref['path']) if isinstance(ref, dict) and ref.get('path') else None
        review = donor.get('whole_subject_review', {})
        if (path is None or digest(path) != ref.get('sha256')
                or not donor.get('category') or not donor.get('semantic_instance_id')
                or review.get('status') != 'PASS' or not review.get('evidence_refs')):
            raise ValueError('A donor must be a hash-bound independently reviewed whole subject')
        with Image.open(path) as image:
            size = image.size
        donors.append((donor, size))
    used_hashes, used_instances, layers = set(), set(), []
    for card in request['cards']:
        contract_ref = card.get('completion_contract')
        if not isinstance(contract_ref, dict) or not contract_ref.get('path'):
            raise ValueError('Each card needs its frozen completion contract')
        contract_file = safe_path(base, contract_ref['path'])
        if digest(contract_file) != contract_ref.get('sha256'):
            raise ValueError('Completion contract changed')
        contract = read_json(contract_file)
        if contract.get('kind') != 'illustrative_photo':
            raise ValueError('Photo-stack completion cannot synthesize data or native icons')
        source, _, _, _ = validate_completion(contract, contract_file.parent)
        with Image.open(source) as image:
            canvas = image.size
        source_status = card.get('source_subject_status')
        source_category = card.get('source_subject_category')
        candidate_category = card.get('candidate_subject_category')
        if (source_status not in ('observed', 'inferred', 'unresolved')
                or (source_status == 'observed') != bool(source_category)
                or not card.get('source_evidence_refs')
                or (source_category and candidate_category != source_category)):
            raise ValueError('Do not project a hidden candidate category backward onto the source')
        if not candidate_category:
            layers.append({'layer_id': card['layer_id'],
                           'route': 'NEEDS_HOST_CATEGORY_DECISION',
                           'source_subject_status': source_status,
                           'reason': 'The source does not identify this hidden card; choose and record a plausible candidate category first'})
            continue
        available = [(donor, size) for donor, size in donors
                     if donor['category'] == candidate_category and size == canvas
                     and (request['repetition_policy'] != 'distinct_samples'
                          or donor['asset']['sha256'] not in used_hashes
                          and donor['semantic_instance_id'] not in used_instances)]
        if available:
            donor, _ = available[0]
            route = 'INDEPENDENT_FULL_SUBJECT_DONOR'
            used_hashes.add(donor['asset']['sha256'])
            used_instances.add(donor['semantic_instance_id'])
            donor_ref = donor['asset']
            instance_id = donor['semantic_instance_id']
        else:
            route = 'HOST_GENERATE_ONE_FULL_SUBJECT'
            donor_ref = None
            instance_id = None
        layers.append({'layer_id': card['layer_id'], 'route': route,
                       'source_subject_status': source_status,
                       'source_evidence_refs': card['source_evidence_refs'],
                       'candidate_subject_category': candidate_category,
                       'canvas': list(canvas), 'completion_contract': contract_ref,
                       'donor': donor_ref, 'semantic_instance_id': instance_id,
                       'host_brief': ('Create or select one complete ' + candidate_category +
                                      ' on this fixed canvas. Continue the visible subject as one coherent object; '
                                      'do not add another subject, an occluder, a collage, or scientific data. '
                                      'Match the observed viewpoint and style. Preserve observed pixels only via '
                                      'complete-asset, then inspect the isolated full card and the recomposed slide.')})
    return {'schema_version': 'illustrative-stack-execution-plan-1',
            'request_hash': identity(request), 'repetition_policy': request['repetition_policy'],
            'policy_evidence_refs': request['policy_evidence_refs'], 'layers': layers,
            'status': 'HOST_CANDIDATES_AND_SEMANTIC_REVIEW_REQUIRED',
            'scope': 'source-bound route selection, not hidden truth or a generated image'}


def freeze_stack_visible_mask(pptx, target_id, occluder_ids, output, *, margin=6):
    """Freeze a conservative source-pixel mask for opaque axis-aligned PPT photos.

    Every listed occluder must be a later sibling. Omitted later artwork only
    overprotects the donor; it never licenses extra changes. Margin accounts for
    drawing/picture resampling before the candidate is generated.
    """
    if type(margin) is not int or not 0 <= margin <= 16:
        raise ValueError('Visible-mask safety margin must be 0..16 source pixels')
    if not occluder_ids or target_id in occluder_ids or len(occluder_ids) != len(set(occluder_ids)):
        raise ValueError('Require distinct later picture occluders')
    deck = Presentation(pptx)
    if len(deck.slides) != 1:
        raise ValueError('Stack mask currently requires one slide')

    def rotated_or_flipped(shape):
        transforms = shape._element.xpath('./p:spPr/a:xfrm | ./p:grpSpPr/a:xfrm')
        return any(int(x.get('rot', '0')) % 21600000 != 0
                   or x.get('flipH') in ('1', 'true') or x.get('flipV') in ('1', 'true')
                   for x in transforms)

    def locate(shapes, transformed=False):
        ids = [shape.shape_id for shape in shapes]
        if target_id in ids:
            if transformed:
                raise ValueError('Transformed ancestor needs explicit pixel visibility analysis')
            index = ids.index(target_id)
            if any(oid not in ids[index + 1:] for oid in occluder_ids):
                raise ValueError('Occluders must be later siblings in z order')
            selected = [shapes[index]] + [shapes[ids.index(oid)] for oid in occluder_ids]
            return selected
        for shape in shapes:
            if shape.shape_type == 6:
                found = locate(shape.shapes, transformed or rotated_or_flipped(shape))
                if found:
                    return found
        return None

    shapes = locate(deck.slides[0].shapes)
    if shapes is None:
        raise ValueError('Target picture or later sibling occluder is absent')
    for shape in shapes:
        if shape.shape_type != 13 or rotated_or_flipped(shape):
            raise ValueError('Stack mask only supports unrotated, unflipped pictures')
        if any(abs(getattr(shape, 'crop_' + side)) > 1e-9 for side in ('left', 'right', 'top', 'bottom')):
            raise ValueError('Stack picture has a crop; use an explicit host-observed mask')
        with Image.open(io.BytesIO(shape.image.blob)) as image:
            if np.any(np.asarray(image.convert('RGBA').getchannel('A')) != 255):
                raise ValueError('Transparent stack picture needs explicit visibility analysis')
    rows = {row['shape_id']: row for row in inventory(pptx)['objects']}
    x, y, width, height = rows[target_id]['bbox_emu']
    if min(width, height) <= 0:
        raise ValueError('Degenerate picture bbox')
    with Image.open(io.BytesIO(shapes[0].image.blob)) as image:
        pixels_w, pixels_h = image.size
    yy, xx = np.mgrid[:pixels_h, :pixels_w]
    px = x + (xx + 0.5) * width / pixels_w
    py = y + (yy + 0.5) * height / pixels_h
    visible = np.ones((pixels_h, pixels_w), dtype=bool)
    for shape_id in occluder_ids:
        ox, oy, ow, oh = rows[shape_id]['bbox_emu']
        visible &= ~((px >= ox) & (px < ox + ow) & (py >= oy) & (py < oy + oh))
    if margin:
        visible = binary_dilation(visible, iterations=margin)
    output = Path(output)
    if output.exists():
        raise FileExistsError(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    Image.fromarray(visible.astype(np.uint8) * 255).save(output)
    receipt = {'schema_version': 'photo-stack-visible-mask-1', 'pptx_hash': digest(pptx),
               'target_id': target_id, 'occluder_ids': list(occluder_ids), 'margin_source_pixels': margin,
               'target_media_hash': hashlib.sha256(shapes[0].image.blob).hexdigest(),
               'geometry_emu': {str(sid): rows[sid]['bbox_emu'] for sid in [target_id] + list(occluder_ids)},
               'mask_hash': digest(output), 'visible_pixels': int(visible.sum()),
               'status': 'FROZEN_BEFORE_CANDIDATE',
               'scope': 'Opaque unrotated sibling pictures; alternate renderer check still required'}
    write_json(output.with_suffix('.freeze.json'), receipt)
    return receipt
