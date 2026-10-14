"""Bound, read-only crop plans and evidence-bearing human/host review imports.

``prepare_review`` compares the frozen reference with a saved-file render. Its
optional actual_objects argument is a parsed-PPT receipt, never a second Scene:
{pptx_hash, parser_version, objects: {id: {bbox, text, ast, evidence_ref, ...}}}.
Bboxes and evidence roi fields are canvas [x, y, width, height]; crop rectangles
are half-open [left, top, right, bottom]. Only one whole-canvas resize is allowed.

``import_review`` consumes {binding, manifest_hash, evaluator: {id, version},
observations: [{target_id, criterion, level, note, evidence_refs, defects}]}.
Each evidence ref is {path, sha256, roi}; roi must equal that target's core or
context. Defects are {id, severity: critical|major|minor, note}. Levels use the
fixed 0--3 rubric below. Null observations remain unverified. Imports describe
review coverage and raw rubric values; they cannot accept a reconstruction.
"""

import math
import time
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw, ImageFont

from .common import ROOT, digest, identity, read_json, safe_path, write_json, raster_rgb
from .grids import coverage, windows
from .metrics import balance, cer, foreground, iou, layout, render_canvas_mapping
from .mathml import canonical_ast

CHECKER_VERSION = "bound-review-1.0.0"
GLOBAL_WEIGHTS = {"layout": .25, "visual_balance": .15, "alignment_spacing": .20,
                  "layer_occlusion": .10, "color_style": .10, "readability": .20}
RUBRIC = {"0": "Clear severe error", "1": "Visible repairable defect",
          "2": "Readable with localized shortcomings", "3": "No locatable defect found"}


def prepare_loop_views(workdir, scene_file, pptx, preview, scopes, issues, candidate_id, previous=()):
    """Prepare bounded host observation batches, including every embedded bitmap.

    Unlike an object inventory alone, a coarse source scan can reveal an owner
    missing from the Scene. Pixel differences are hints, never defect verdicts.
    """
    import io
    from pptx import Presentation
    workdir, scene_file = Path(workdir), Path(scene_file)
    scene = read_json(scene_file)
    source_path = safe_path(scene_file.parent, scene['source']['path'])
    if digest(source_path) != scene['source']['sha256']:
        raise ValueError('Observation source changed')
    with Image.open(source_path) as img:
        source = raster_rgb(img, scene['canvas'].get('background_rgb', [255, 255, 255]))
    with Image.open(preview) as img:
        candidate = img.convert('RGB').resize(source.size, Image.Resampling.LANCZOS)
    width, height = source.size
    folder = workdir / 'visual-views' / candidate_id
    folder.mkdir(parents=True, exist_ok=False)
    views = []
    old = {v['id']: v for v in previous}
    def emit(key, kind, roi, result=None, owner=None):
        rect = _rect(roi, width, height)
        if rect is None:
            return
        images = {'source': source.crop(rect), 'candidate': result if result is not None else candidate.crop(rect)}
        refs = {}
        for role, image in images.items():
            path = folder / (identity(key)[:16] + '-' + role + '.png')
            image.save(path)
            refs[role] = {'path': path.relative_to(workdir).as_posix(), 'sha256': digest(path)}
        # File names include the candidate id, so compare content identities.
        signature = identity({'images': {k: v['sha256'] for k, v in refs.items()},
                              'kind': kind, 'owner': owner, 'roi': roi})
        record = {'id': key, 'kind': kind, 'roi': roi, 'owner': owner, **refs,
                  'signature': signature, 'reuse': False,
                  'prompt': 'Compare source meaning and actual appearance; locate omissions, extras, spacing, boundary, clarity, texture and reuse defects. Unknown defect kinds are allowed.'}
        before = old.get(key)
        if (kind != 'whole' and before and before.get('signature') == signature
                and before.get('review', {}).get('status') == 'satisfactory'):
            record.update(reuse=True, review=before['review'])
        if images['source'].size == images['candidate'].size:
            record['difference_hint'] = _pixels(images['source'], images['candidate'])
        if kind == 'bitmap_full':
            record['prompt'] = ('Source shows visible placement; candidate is the complete saved picture. '
                                'Inspect semantic match, whole-subject continuity, hidden completion, natural sharpness, '
                                'mosaic/blur, boundaries and independent reuse. Reconstructed details may differ from a blurry source.')
        views.append(record)
    emit('whole', 'whole', [0, 0, width, height])
    for row in range(3):
        for col in range(3):
            l, t = width*col//3, height*row//3
            r, b = width*(col+1)//3, height*(row+1)//3
            emit(f'source-scan:{row}:{col}', 'source_scan', [l, t, r-l, b-t])
    for owner, scope in scopes.items():
        parent = scope.get('parent')
        if scope.get('semantic') and (not parent or not scopes.get(parent, {}).get('parent')):
            emit('module:' + owner, 'module', scope['roi'], owner=owner)
    for issue in issues:
        if issue.get('status') == 'open':
            emit('issue:' + issue['id'], 'issue', issue['roi'], owner=issue.get('owner'))
    objects = {o['id']: o for o in scene['objects']}
    prs = Presentation(pptx)
    def visit(shapes, prefix, slide_width, slide_height):
        for shape in shapes:
            key = prefix + '/' + shape.name
            if shape.shape_type == 6:
                visit(shape.shapes, key, slide_width, slide_height)
            elif shape.shape_type == 13:
                obj = objects.get(shape.name)
                roi = obj['bbox'] if obj else [float(shape.left)*width/slide_width, float(shape.top)*height/slide_height,
                                               float(shape.width)*width/slide_width, float(shape.height)*height/slide_height]
                with Image.open(io.BytesIO(shape.image.blob)) as image:
                    asset = raster_rgb(image)
                emit('bitmap:' + key, 'bitmap_full', roi, asset, shape.name)
                emit('bitmap-context:' + key, 'bitmap_context', roi, owner=shape.name)
    for index, slide in enumerate(prs.slides, 1):
        visit(slide.shapes, 'slide-' + str(index), prs.slide_width, prs.slide_height)
    return views


def implementation_hash():
    files = [p for folder in ("scripts", "config", "schemas")
             for p in (ROOT / folder).rglob("*") if p.is_file() and p.suffix in (".py", ".ps1", ".json")]
    return identity({p.relative_to(ROOT).as_posix(): digest(p) for p in sorted(files)})


def current_binding(scene, base, pptx, preview, render_receipt, *, implementation_version):
    """Recompute all bound identities; callers must use this before an import."""
    if not isinstance(implementation_version, str) or not implementation_version.strip():
        raise ValueError("Implementation version required")
    for key in ("source", "evidence", "profile"):
        if digest(safe_path(base, scene[key]["path"])) != scene[key]["sha256"]:
            raise ValueError("Frozen " + key + " hash changed")
    receipt = read_json(render_receipt)
    if receipt.get("pptx_hash") != digest(pptx) or receipt.get("preview_hash") != digest(preview):
        raise ValueError("Render receipt does not bind saved PPT and preview")
    return {"source_hash": scene["source"]["sha256"], "evidence_hash": scene["evidence"]["sha256"],
            "scene_hash": identity(scene), "scene_revision": scene["revision"],
            "profile_hash": scene["profile"]["sha256"], "pptx_hash": digest(pptx),
            "preview_hash": digest(preview), "render_receipt_hash": digest(render_receipt),
            "implementation_version": implementation_version, "implementation_hash": implementation_hash(),
            "checker_version": CHECKER_VERSION,
            "component_versions": {o["id"]: dict(o["versions"]) for o in scene["objects"]},
            "dependency_hash": identity({"relations": scene["relations"], "prototypes": scene["prototypes"],
                                          "parents": {o["id"]: o.get("parent") for o in scene["objects"]}}),
            "render_environment": {k: receipt.get(k) for k in ("backend", "backend_version", "environment")}}


def _box(value):
    if not isinstance(value, (list, tuple)) or len(value) != 4 or any(
            isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) for v in value):
        raise ValueError("Expected finite xywh bbox")
    if value[2] <= 0 or value[3] <= 0:
        raise ValueError("Expected positive bbox extents")
    return list(value)


def _rect(box, width, height, padding=0):
    x, y, w, h = _box(box)
    rect = [max(0, math.floor(x-padding)), max(0, math.floor(y-padding)),
            min(width, math.ceil(x+w+padding)), min(height, math.ceil(y+h+padding))]
    if rect[0] >= rect[2] or rect[1] >= rect[3]:
        return None
    return rect


def _pixels(reference, candidate, threshold=245):
    a, b = np.asarray(reference.convert("RGB"), dtype=np.int16), np.asarray(candidate.convert("RGB"), dtype=np.int16)
    am, bm = foreground(reference, threshold), foreground(candidate, threshold)
    difference = np.abs(a-b)
    union = am | bm
    return {"foreground_iou": iou(am, bm), "mae_rgb_255": float(difference.mean()),
            "foreground_mae_rgb_255": float(difference[union].mean()) if union.any() else None,
            "reference_foreground_pixels": int(am.sum()), "candidate_foreground_pixels": int(bm.sum()),
            "added_foreground_pixels": int((bm & ~am).sum()), "missing_foreground_pixels": int((am & ~bm).sum()),
            "reference_blank": not bool(am.any()), "unexpected_foreground_on_blank": bool(bm.any()) if not am.any() else None,
            "status": "UNVERIFIED"}


def object_measurements(reference, actual, width, height):
    """Compare observed source evidence with parser output, preserving unknowns."""
    observation = reference.get("observation", reference.get("content", {}))
    observed = reference.get("evidence_status") == "observed"
    ref_box = reference.get("roi", reference.get("bbox"))
    result = {"actual_present": actual is not None, "position": None, "text_cer": None,
              "formula_ast_exact": None, "endpoint_error_canvas": None, "missing": [],
              "status": "UNVERIFIED", "evidence_status": reference.get("evidence_status")}
    if not isinstance(observation, dict):
        # Prose is source evidence for a host review, not a structured text,
        # formula or endpoint annotation. Never parse words such as "text"
        # or "ast" in that prose as measurement fields.
        observation = {}
        result["missing"].append("structured_source_observation")
    if not observed:
        result["missing"].append("observed_reference_evidence")
        return result
    if actual is None:
        result["missing"].append("saved_ppt_object")
        return result
    if not actual.get("evidence_ref"):
        raise ValueError("Saved PPT measurement lacks object evidence_ref")
    if ref_box is not None and actual.get("bbox") is not None:
        r, c = _box(ref_box), _box(actual["bbox"])
        dx, dy = c[0]+c[2]/2-r[0]-r[2]/2, c[1]+c[3]/2-r[1]-r[3]/2
        result["position"] = {"center_dx_px": dx, "center_dy_px": dy,
                              "center_error_canvas": math.hypot(dx/width, dy/height),
                              "width_error_px": c[2]-r[2], "height_error_px": c[3]-r[3],
                              "reference_bbox": r, "actual_bbox": c}
    else:
        result["missing"].append("reference_or_actual_bbox")
    if "text" in observation:
        value = actual.get("text")
        if value is not None and not isinstance(value, str):
            raise ValueError("Parsed text must be a string")
        result.update({"reference_text": observation["text"], "actual_text": value,
                       "text_cer": cer(observation["text"], value),
                       "text_exact": value == observation["text"] if value is not None else None,
                       "text_normalization": "none; case, whitespace and scientific signs preserved"})
        if value is None:
            result["missing"].append("saved_ppt_text")
        if "baseline" in observation and actual.get("baseline") is not None:
            result["baseline_error_px"] = actual["baseline"]-observation["baseline"]
        else:
            result["missing"].append("reference_or_actual_text_baseline")
    if "ast" in observation:
        result["formula_ast_exact"] = canonical_ast(observation["ast"]) == canonical_ast(actual["ast"]) if actual.get("ast") is not None else None
        if actual.get("ast") is None:
            result["missing"].append("saved_ppt_math_ast")
    if "endpoints" in observation and actual.get("endpoints") is not None:
        a, b = np.asarray(observation["endpoints"], dtype=float), np.asarray(actual["endpoints"], dtype=float)
        if a.shape != (2, 2) or b.shape != (2, 2) or not np.isfinite(a).all() or not np.isfinite(b).all():
            raise ValueError("Directed endpoint measurement needs two finite canvas points")
        result["endpoint_error_canvas"] = float(np.linalg.norm(a-b, axis=1).mean()/math.hypot(width, height))
    return result


def connection_corridors(scene):
    """Return full connector/control-hull crops, independent of grid boundaries."""
    width, height = scene["canvas"]["width"], scene["canvas"]["height"]
    result = []
    for obj in scene["objects"]:
        content = obj["content"]
        if obj["render_type"] not in ("connector", "bezier") or not content.get("points"):
            continue
        points = np.asarray(content["points"], dtype=float)
        if points.ndim != 2 or points.shape[1] != 2 or len(points) < 2 or not np.isfinite(points).all():
            raise ValueError("Connection points require finite pairs")
        space = content.get("coordinate_space", "canvas" if obj["render_type"] == "connector" else "normalized")
        if space == "normalized":
            x, y, w, h = _box(obj["bbox"])
            points = points*np.array([w, h])+np.array([x, y])
        elif space != "canvas":
            raise ValueError("Unsupported connection coordinate_space")
        lo, hi = points.min(axis=0), points.max(axis=0)
        # A cubic lies inside its control hull, so this includes every segment.
        pad = max(6, 3*obj.get("style", {}).get("line_width", 1))
        rect = _rect([lo[0], lo[1], max(1, hi[0]-lo[0]), max(1, hi[1]-lo[1])], width, height, pad)
        if rect is not None:
            result.append({"id": "corridor:"+obj["id"], "kind": "corridor", "object_ids": [obj["id"]],
                           "core": rect, "context": rect, "points_canvas": points.tolist(),
                           "required_criteria": ["direction_topology", "occlusion_continuity"]})
    return result


def _contact_sheets(targets, output):
    pages = []
    font = ImageFont.load_default(size=14)
    for offset in range(0, len(targets), 6):
        batch = targets[offset:offset+6]
        sheet = Image.new("RGB", (1536, 420*math.ceil(len(batch)/2)), "#e5e7eb")
        draw = ImageDraw.Draw(sheet)
        for idx, target in enumerate(batch):
            ox, oy = (idx % 2)*768, (idx//2)*420
            draw.text((ox+8, oy+5), target["id"], fill="black", font=font)
            for row, field in enumerate(("core", "context")):
                for col, role in enumerate(("reference", "candidate")):
                    x, y = ox+8+col*380, oy+28+row*194
                    draw.text((x, y), role+" "+field, fill="black", font=font)
                    with Image.open(output/target["crops"][field][role]["path"]) as crop:
                        zoom = min(4, 368/crop.width, 164/crop.height)
                        size = (max(1, round(crop.width*zoom)), max(1, round(crop.height*zoom)))
                        display = crop.resize(size, Image.Resampling.NEAREST if zoom > 1 else Image.Resampling.LANCZOS)
                        sheet.paste(display, (x, y+22))
        path = f"contact-sheets/page-{offset//6+1:03d}.png"
        (output/path).parent.mkdir(exist_ok=True)
        sheet.save(output/path)
        pages.append({"path": path, "sha256": digest(output/path), "target_ids": [t["id"] for t in batch]})
    return pages


def prepare_review(scene, base, pptx, preview, render_receipt, output_dir, *, implementation_version,
                   actual_objects=None, max_crop_pixels=200_000_000):
    """Write a new, portable comparison directory; refuse overwrites and drift."""
    started, timestamp = time.monotonic(), datetime.now(timezone.utc).isoformat()
    binding = current_binding(scene, base, pptx, preview, render_receipt,
                              implementation_version=implementation_version)
    profile = read_json(safe_path(base, scene["profile"]["path"]))
    if profile["global_aesthetic"]["weights"] != GLOBAL_WEIGHTS:
        raise ValueError("User's six global weights changed")
    evidence = read_json(safe_path(base, scene["evidence"]["path"]))
    if evidence.get("source_hash") != binding["source_hash"]:
        raise ValueError("Evidence does not bind the frozen source")
    known = {o["id"]: o for o in evidence.get("objects", [])}
    if not known or len(known) != len(evidence["objects"]):
        raise ValueError("Source inventory must be nonempty with unique IDs")
    objects = {o["id"]: o for o in scene["objects"]}
    if not objects or len(objects) != len(scene["objects"]):
        raise ValueError("Scene inventory must be nonempty with unique IDs")
    actual = {}
    if actual_objects is not None:
        if actual_objects.get("pptx_hash") != binding["pptx_hash"] or not actual_objects.get("parser_version"):
            raise ValueError("Actual-object measurements need saved PPT hash and parser version")
        actual = actual_objects["objects"]
    width, height = scene["canvas"]["width"], scene["canvas"]["height"]
    if width*height > 40_000_000 or max_crop_pixels <= 0:
        raise TimeoutError("Review exceeds explicit pixel budget")
    with Image.open(safe_path(base, scene["source"]["path"])) as img:
        reference = raster_rgb(img,scene["canvas"].get("background_rgb",[255,255,255]))
    with Image.open(preview) as img:
        candidate = img.convert("RGB")
    if reference.size != (width, height):
        raise ValueError("Source canvas mismatch")
    canvas_mapping = render_canvas_mapping((width,height),candidate.size)
    render_size = list(candidate.size)
    candidate = candidate.resize((width, height), Image.Resampling.LANCZOS)
    tiles = windows(width, height, profile["grid"]["halo"])
    targets = [t | {"kind": "grid", "required_criteria": ["visual_fidelity", "omission_extras"]} for t in tiles]
    missing = []
    for obj_id in sorted(set(known) | set(objects)):
        ref = known.get(obj_id)
        if ref is None:
            missing.append({"object_id": obj_id, "item": "frozen_source_object_evidence"})
            continue
        box = ref.get("roi", ref.get("bbox"))
        rect = _rect(box, width, height) if box is not None else None
        if rect is None:
            missing.append({"object_id": obj_id, "item": "nonempty_reference_roi"})
            continue
        criteria = ["semantic_content", "visual_fidelity"]
        kind = objects.get(obj_id, {}).get("render_type")
        if kind == "group":
            criteria.append("natural_parts")
        if kind in ("text", "equation"):
            criteria.append("readability")
        measurement = object_measurements(ref, actual.get(obj_id), width, height)
        targets.append({"id": "object:"+obj_id, "kind": "object", "object_ids": [obj_id],
                        "core": rect, "context": _rect(box, width, height, max(6, .12*max(box[2:]))),
                        "required_criteria": criteria, "object_measurements": measurement,
                        "source_evidence_status": ref.get("evidence_status")})
    targets.extend(connection_corridors(scene))
    full = [0, 0, width, height]
    targets.extend({"id": "global:"+key, "kind": "global", "core": full, "context": full,
                    "required_criteria": [key]} for key in GLOBAL_WEIGHTS)
    targets.append({"id": "source-only", "kind": "source_inventory", "core": full, "context": full,
                    "required_criteria": ["source_omission_review"]})
    crop_pixels = sum(2*(t[field][2]-t[field][0])*(t[field][3]-t[field][1])
                      for t in targets for field in ("core", "context"))
    if len(targets) > 10_000 or crop_pixels > max_crop_pixels:
        raise TimeoutError("Review crops exceed explicit resource budget")
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=False)
    (output/"crops").mkdir()
    crop_cache = {}
    for index, target in enumerate(targets):
        target["object_ids"] = target.get("object_ids", [k for k, obj in known.items()
            if (box := obj.get("roi", obj.get("bbox"))) is not None
            and box[0] < target["core"][2] and box[0]+box[2] > target["core"][0]
            and box[1] < target["core"][3] and box[1]+box[3] > target["core"][1]])
        target["crops"] = {}
        for field in ("core", "context"):
            pair = {}
            for role, image in (("reference", reference), ("candidate", candidate)):
                cache_key = (tuple(target[field]), role)
                if cache_key not in crop_cache:
                    relative = f"crops/{index:04d}-{field}-{role}.png"
                    image.crop(target[field]).save(output/relative)
                    crop_cache[cache_key] = {"path": relative, "sha256": digest(output/relative), "roi": target[field]}
                pair[role] = crop_cache[cache_key]
            target["crops"][field] = pair
        target["measurements"] = _pixels(reference.crop(target["core"]), candidate.crop(target["core"]))
    ref_boxes = {k: r.get("roi", r.get("bbox")) for k, r in known.items()
                 if r.get("evidence_status") == "observed" and r.get("roi", r.get("bbox")) is not None}
    actual_boxes = {k: a["bbox"] for k, a in actual.items() if k in ref_boxes and a.get("bbox") is not None}
    relations = [(r["source"], r["target"]) for r in evidence.get("relations", [])
                 if r.get("evidence_status") == "observed" and r.get("source") in ref_boxes and r.get("target") in ref_boxes]
    manifest = {"checker_version": CHECKER_VERSION, "binding": binding,
                "started_at": timestamp, "profile_status": profile["status"], "rubric": RUBRIC,
                "global_weights": GLOBAL_WEIGHTS, "targets": targets, "grid_coverage": coverage(tiles, width, height),
                "planned_grid_count": len(tiles), "planned_object_ids": sorted(set(known) | set(objects)),
                "reference_only_object_ids": sorted(set(known)-set(objects)),
                "candidate_only_object_ids": sorted(set(objects)-set(known)),
                "actual_only_object_ids": sorted(set(actual)-set(known)),
                "source_missing_items": missing, "contact_sheets": _contact_sheets(targets, output),
                "measurement_protocol": {"version": CHECKER_VERSION, "canvas": [width, height],
                    "input_render_size": render_size, "canvas_mapping":canvas_mapping,"canvas_transform": "whole_canvas_lanczos_resize",
                    "local_registration": "none", "foreground": "any RGB channel < 245; white-background raw diagnostic",
                    "pixel_error": "mean absolute RGB error; range 0..255; foreground union separately",
                    "empty_foreground": "IoU null when both empty; additions reported separately",
                    "masks": "full fixed source ROI including actual occlusion; no candidate-dependent mask",
                    "raw_metrics_only": True, "thresholds": None, "grid_halo": profile["grid"]["halo"]},
                "raw_global": {"reference_balance_white_background": balance(reference),
                    "candidate_balance_white_background": balance(candidate),
                    "layout": layout(ref_boxes, actual_boxes, relations, width, height)},
                "actual_objects_receipt_hash": identity(actual_objects) if actual_objects is not None else None,
                "status": "UNVERIFIED", "reason": "Measurements and crop coverage do not establish calibrated quality or native editing",
                "elapsed_seconds": time.monotonic()-started, "crop_pixels_upper_bound": crop_pixels,
                "unique_crop_count": len(crop_cache)}
    if current_binding(scene, base, pptx, preview, render_receipt,
                       implementation_version=implementation_version) != binding:
        raise ValueError("Review inputs or implementation changed during computation")
    manifest["manifest_hash"] = identity(manifest)
    write_json(output/"review-manifest.json", manifest)
    return manifest


def import_review(manifest, review, *, current_binding, review_directory=None):
    """Validate observations without supplying absent observations or accepting."""
    if manifest.get("manifest_hash") != identity({k: v for k, v in manifest.items() if k != "manifest_hash"}):
        raise ValueError("Review manifest integrity failure")
    if review.get("binding") != manifest["binding"] or current_binding != manifest["binding"]:
        raise ValueError("Stale review binding")
    if review.get("manifest_hash") != manifest["manifest_hash"]:
        raise ValueError("Review is not bound to this crop plan")
    if manifest.get("global_weights") != GLOBAL_WEIGHTS or manifest.get("rubric") != RUBRIC:
        raise ValueError("Fixed weights or rubric changed")
    width, height = manifest["measurement_protocol"]["canvas"]
    expected_grid = windows(width, height, manifest["measurement_protocol"]["grid_halo"])
    actual_grid = [{k: t[k] for k in ("id", "scheme", "core", "context")}
                   for t in manifest["targets"] if t["kind"] == "grid"]
    if actual_grid != expected_grid or manifest["planned_grid_count"] != len(expected_grid):
        raise ValueError("Review plan must preserve all 94 fixed grid windows")
    evaluator = review.get("evaluator", {})
    if not evaluator.get("id") or not evaluator.get("version"):
        raise ValueError("Evaluator identity and version required")
    targets = {t["id"]: t for t in manifest["targets"]}
    if len(targets) != len(manifest["targets"]):
        raise ValueError("Duplicate planned target")
    if any("global:"+k not in targets for k in GLOBAL_WEIGHTS) or "source-only" not in targets:
        raise ValueError("Required global or source-omission review target missing")
    for target in targets.values():
        fixed = {"grid": ["visual_fidelity", "omission_extras"],
                 "corridor": ["direction_topology", "occlusion_continuity"],
                 "source_inventory": ["source_omission_review"]}
        required = target.get("required_criteria", [])
        if not required or len(required) != len(set(required)):
            raise ValueError("Empty or duplicate review checklist")
        if target["kind"] in fixed and required != fixed[target["kind"]]:
            raise ValueError("Fixed review checklist changed")
        if target["kind"] == "global" and (not target["id"].startswith("global:")
                or target["id"][7:] not in GLOBAL_WEIGHTS or required != [target["id"][7:]]):
            raise ValueError("Global dimension checklist changed")
        if target["kind"] == "object" and not {"semantic_content", "visual_fidelity"}.issubset(required):
            raise ValueError("Object semantic and visual review cannot be omitted")
    observations, seen, defects = [], set(), {}
    for item in review.get("observations", []):
        target = targets.get(item.get("target_id"))
        if target is None or item.get("criterion") not in target["required_criteria"]:
            raise ValueError("Observation references an unplanned target or criterion")
        key = (item["target_id"], item["criterion"])
        if key in seen:
            raise ValueError("Duplicate target criterion; reviews cannot be averaged")
        seen.add(key)
        level = item.get("level")
        if level is not None and (type(level) is not int or level not in (0, 1, 2, 3)):
            raise ValueError("Rubric level must be null or an integer from 0 to 3")
        if not isinstance(item.get("note"), str) or not item["note"].strip():
            raise ValueError("Located observation or missing-measurement reason required")
        refs = item.get("evidence_refs", [])
        if level is not None and not refs:
            raise ValueError("Scored rubric requires located crop evidence")
        if level is not None and review_directory is None:
            raise ValueError("Scored rubric requires a review directory to verify actual crop bytes")
        allowed = [target["crops"][field][role] for field in ("core", "context")
                   for role in ("reference", "candidate")]
        for ref in refs:
            if ref not in allowed:
                raise ValueError("Evidence must bind an exact target crop hash and canvas ROI")
            if review_directory is not None and digest(safe_path(review_directory, ref["path"])) != ref["sha256"]:
                raise ValueError("Review crop artifact changed")
        if level is not None:
            # A source-only inventory must still see the actual saved candidate;
            # looking only at the source cannot establish that nothing vanished.
            required_roles = ("reference", "candidate")
            if any(not any(ref == target["crops"][field][role] for field in ("core", "context") for ref in refs)
                   for role in required_roles):
                raise ValueError("Scored comparison needs both source and candidate evidence")
        for defect in item.get("defects", []):
            if not defect.get("id") or defect.get("severity") not in ("critical", "major", "minor") or not defect.get("note"):
                raise ValueError("Defect requires stable ID, severity and located description")
            if level == 3:
                raise ValueError("No-defect rubric conflicts with a located defect")
            prior = defects.get(defect["id"])
            if prior and (prior["severity"] != defect["severity"] or prior["note"] != defect["note"]):
                raise ValueError("Repeated defect identity has inconsistent meaning")
            defects.setdefault(defect["id"], dict(defect, target_ids=[]))["target_ids"].append(item["target_id"])
        observations.append(dict(item, level=level, normalized_rubric=level/3 if level is not None else None,
                                 status="OBSERVED" if level is not None else "UNVERIFIED"))
    completed = {(o["target_id"], o["criterion"]) for o in observations if o["level"] is not None}
    missing = [{"target_id": t["id"], "criterion": c} for t in targets.values()
               for c in t["required_criteria"] if (t["id"], c) not in completed]
    reviewed_targets = [t["id"] for t in targets.values() if all((t["id"], c) in completed for c in t["required_criteria"])]
    dimensions = {key: next((o["normalized_rubric"] for o in observations if o["target_id"] == "global:"+key), None)
                  for key in GLOBAL_WEIGHTS}
    return {"status": "UNVERIFIED", "scope": "located_fixed_rubric_review_only", "binding": current_binding,
            "manifest_hash": manifest["manifest_hash"], "review_hash": identity(review), "evaluator": evaluator,
            "observations": observations, "missing_observations": missing,
            "source_missing_items": manifest["source_missing_items"], "defects": list(defects.values()),
            "critical_defect_count": sum(d["severity"] == "critical" for d in defects.values()),
            "planned_target_count": len(targets), "reviewed_target_count": len(reviewed_targets),
            "reviewed_target_ids": reviewed_targets, "planned_object_ids": manifest["planned_object_ids"],
            "unreviewed_object_ids": [k for k in manifest["planned_object_ids"] if "object:"+k not in reviewed_targets],
            "global_weights": GLOBAL_WEIGHTS, "global_dimensions_rubric_only": dimensions,
            "global_aesthetic_score": None, "local_quality_scores": None,
            "reason": "Raw rubric mapping is not a calibrated composite; native, semantic and quantitative gates remain separate",
            "reviewed_at": datetime.now(timezone.utc).isoformat()}
