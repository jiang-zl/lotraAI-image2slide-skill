"""Make an existing flat native owner reusable without rasterizing it.

Grouping is only a structural candidate until the saved presentation has
passed full-page render parity and declared editor operations.
"""

import re
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation

from .common import digest, read_json, write_json
from .render import verify_alternative_render
from .saved_ppt_diff import snapshot


def _matches(name, selector):
    if selector.endswith("*"):
        return name.startswith(selector[:-1])
    return name == selector


def group_native_owner(contract, baseline_pptx, candidate_pptx):
    """Move selected top-level native shapes into one semantic PPT group."""
    if contract.get("schema_version") != "native-owner-grouping-1":
        raise ValueError("Unknown native owner grouping contract")
    baseline_pptx, candidate_pptx = Path(baseline_pptx), Path(candidate_pptx)
    if candidate_pptx.exists() or candidate_pptx.resolve() == baseline_pptx.resolve():
        raise FileExistsError("Native grouping needs a new task-owned output path")
    if digest(baseline_pptx) != contract.get("baseline_pptx_sha256"):
        raise ValueError("Historical PPT changed after grouping contract was frozen")
    group_name = contract.get("group_name")
    selectors = contract.get("top_level_selectors")
    expected = contract.get("expected_native_shapes")
    if (not isinstance(group_name, str) or not re.fullmatch(r"[A-Za-z][A-Za-z0-9_-]{0,79}", group_name)
            or not isinstance(selectors, list) or not selectors
            or any(not isinstance(item, str) or not item.strip() for item in selectors)
            or type(expected) is not int or expected < 2):
        raise ValueError("Grouping needs a frozen native owner name, selectors and shape count")
    deck = Presentation(baseline_pptx)
    if len(deck.slides) != 1:
        raise ValueError("Bounded grouping currently supports one-slide PPTs")
    slide = deck.slides[0]
    if any(shape.name == group_name for shape in slide.shapes):
        raise ValueError("New semantic owner name already exists")
    selected = [shape for shape in slide.shapes if any(_matches(shape.name, rule) for rule in selectors)]
    if len(selected) != expected or any(not any(_matches(shape.name, rule) for shape in selected)
                                       for rule in selectors):
        raise ValueError("Saved native shape selection differs from frozen grouping contract")
    if any(shape.shape_type in (6, 13) or shape._element.xpath(".//a:blip")
           for shape in selected):
        raise ValueError("Native owner grouping cannot hide pictures or nested groups")
    names = [shape.name for shape in selected]
    if len(set(names)) != len(names):
        raise ValueError("Selected native shapes have duplicate names")
    tree = slide.shapes._spTree
    first_index = min(list(tree).index(shape._element) for shape in selected)
    group = slide.shapes.add_group_shape()
    group.name = group_name
    tree.remove(group._element)
    for shape in selected:
        tree.remove(shape._element)
        group.shapes._spTree.append(shape._element)
    tree.insert(first_index, group._element)
    group.shapes._recalculate_extents()
    candidate_pptx.parent.mkdir(parents=True, exist_ok=True)
    deck.save(candidate_pptx)
    before, after = snapshot(baseline_pptx), snapshot(candidate_pptx)
    moved = {"slide-1/" + name: "slide-1/" + group_name + "/" + name for name in names}
    if (set(after["objects"]) != (set(before["objects"]) - set(moved)) |
            set(moved.values()) | {"slide-1/" + group_name}):
        raise ValueError("Grouping changed the saved shape inventory beyond the semantic owner")
    for old_path, old_record in before["objects"].items():
        new_path = moved.get(old_path, old_path)
        if after["objects"][new_path] != old_record:
            raise ValueError("Grouping changed a native part or unrelated saved object: " + old_path)
    return {"schema_version": "native-owner-grouping-result-1",
            "status": "SAVED_NATIVE_GROUP_NEEDS_RENDER_AND_EDITOR_REVIEW",
            "baseline_pptx_hash": before["pptx_hash"],
            "candidate_pptx_hash": after["pptx_hash"],
            "group_path": "slide-1/" + group_name,
            "native_child_count": len(selected), "native_child_names": names,
            "group_bbox_emu": [int(group.left), int(group.top), int(group.width), int(group.height)],
            "scope": "saved-PPT native grouping and object preservation only"}


def group_file(contract_file, baseline_pptx, candidate_pptx, output):
    result = group_native_owner(read_json(contract_file), baseline_pptx, candidate_pptx)
    write_json(output, result)
    return result


def check_render_parity(grouping_report, baseline_pptx, baseline_receipt,
                        candidate_pptx, candidate_receipt):
    if grouping_report.get("schema_version") != "native-owner-grouping-result-1":
        raise ValueError("Unknown native grouping report")
    if (digest(baseline_pptx) != grouping_report.get("baseline_pptx_hash")
            or digest(candidate_pptx) != grouping_report.get("candidate_pptx_hash")):
        raise ValueError("Native grouping render check has stale PPT identities")
    receipts = []
    previews = []
    for pptx, receipt_file in ((baseline_pptx, baseline_receipt), (candidate_pptx, candidate_receipt)):
        receipt = read_json(receipt_file)
        if receipt.get("backend") != "LibreOffice":
            raise ValueError("Native grouping parity needs same-protocol verified LibreOffice renders")
        verify_alternative_render(pptx, receipt_file)
        receipts.append(receipt)
        with Image.open(Path(receipt_file).parent / "Final_preview.png") as image:
            previews.append(np.asarray(image.convert("RGB")))
    keys = ("backend", "backend_version", "executable", "environment", "mapping", "fontconfig")
    if any(receipts[0].get(key) != receipts[1].get(key) for key in keys):
        raise ValueError("Grouping preview protocol changed")
    if previews[0].shape != previews[1].shape:
        raise ValueError("Grouping preview canvas changed")
    different = int(np.count_nonzero(np.any(previews[0] != previews[1], axis=2)))
    return {"schema_version": "native-owner-grouping-parity-1",
            "status": "PASS_ALTERNATIVE_RENDER_IDENTICAL" if different == 0 else "FAIL_VISUAL_GROUPING_CHANGE",
            "changed_pixels": different, "canvas": list(previews[0].shape[:2][::-1]),
            "baseline_render_receipt_hash": digest(baseline_receipt),
            "candidate_render_receipt_hash": digest(candidate_receipt),
            "grouping_candidate_hash": grouping_report["candidate_pptx_hash"],
            "scope": "same-protocol LibreOffice full-page parity; target-editor operation remains separate"}


def check_render_parity_file(grouping_report_file, baseline_pptx, baseline_receipt,
                             candidate_pptx, candidate_receipt, output):
    result = check_render_parity(read_json(grouping_report_file), baseline_pptx,
                                 baseline_receipt, candidate_pptx, candidate_receipt)
    write_json(output, result)
    return result
