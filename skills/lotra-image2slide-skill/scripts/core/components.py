"""Extract native semantic groups from the saved deck; never replace them with exports."""

import copy
import io
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from .common import digest, identity, write_json, writer
from .inspect import check
from .mathml import NS


def export_components(scene, base, pptx, output):
    if check(scene,base,pptx)["status"] != "PASS":
        raise ValueError("Cannot extract components from structurally invalid deck")
    output = Path(output)
    if output.exists():
        raise FileExistsError(output)
    source = Presentation(pptx)
    components = [o for o in scene["objects"] if o["render_type"] == "group" and o.get('component',not o.get('parent'))]
    if not components:
        raise ValueError("No applicable semantic groups; record explicit N/A externally")
    deck = Presentation()
    deck.slide_width,deck.slide_height = source.slide_width,source.slide_height
    entries = []
    with writer(output):
        for i,o in enumerate(components,1):
            def find(collection,name):
                for candidate in collection:
                    if candidate.name==name:return candidate
                    if hasattr(candidate,'shapes'):
                        found=find(candidate.shapes,name)
                        if found is not None:return found
                return None
            shape=find(source.slides[0].shapes,o['id'])
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            element = copy.deepcopy(shape._element)
            rasters=element.xpath('.//a:blip')
            for blip in rasters:
                if blip.get('{'+NS['r']+'}link'):raise ValueError('External linked media cannot be exported')
                rid=blip.get('{'+NS['r']+'}embed')
                part=source.slides[0].part.related_part(rid)
                if not part.content_type.startswith('image/'):raise ValueError('Unsupported component relationship')
                _,new_rid=slide.part.get_or_add_image_part(io.BytesIO(part.blob))
                blip.set('{'+NS['r']+'}embed',new_rid)
            slide.shapes._spTree.insert_element_before(element,"p:extLst")
            saved = slide.shapes[0]
            saved.left,saved.top = Inches(.3),Inches(.3)
            entries.append({"semantic_id":o["id"],"slide":i,"versions":o["versions"],"representation":"mixed" if rasters else "vector",
                            "source_native_id":shape.shape_id,"source_structure_hash":identity(shape._element.xml),
                            "png_status":"UNVERIFIED","svg_status":"UNVERIFIED","native_behavior":"UNVERIFIED"})
        out = output / "Components_editable.pptx"
        deck.save(out)
        index = {"scene_revision":scene["revision"],"scene_hash":identity(scene),"source_pptx_hash":digest(pptx),
                 "native_pptx_hash":digest(out),"components":entries,"status":"EXTRACTED_NOT_ACCEPTED"}
        write_json(output / "index.json",index)
    return index
