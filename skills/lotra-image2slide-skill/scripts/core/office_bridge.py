"""Use installed Windows Office from Windows or WSL, independently of GUI input."""
import os
import ntpath
import platform
import subprocess
import json
import shutil
from pathlib import Path
from .common import digest, read_json, write_json


def discover(configured=None):
    configured = configured or {}
    system = platform.system()
    if configured.get('disabled'):
        return {'available': False, 'reason': 'explicitly_disabled'}
    if system == 'Windows':
        roots = [Path(os.environ.get(k, 'C:/Program Files')) for k in ('ProgramFiles', 'ProgramFiles(x86)')]
        powershell = Path(os.environ.get('WINDIR', 'C:/Windows'))/'System32/WindowsPowerShell/v1.0/powershell.exe'
    elif system == 'Linux' and (os.environ.get('WSL_DISTRO_NAME') or Path('/proc/sys/fs/binfmt_misc/WSLInterop').exists()):
        mount = Path(configured.get('windows_mount', '/mnt/c'))
        roots = [mount/'Program Files', mount/'Program Files (x86)']
        powershell = mount/'Windows/System32/WindowsPowerShell/v1.0/powershell.exe'
    else:
        return {'available': False, 'reason': 'not_a_local_windows_or_wsl_host'}
    if configured.get('powershell'):
        powershell = Path(configured['powershell'])
    office = Path(configured['powerpoint']) if configured.get('powerpoint') else next(
        (root/relative for root in roots for relative in ('Microsoft Office/root/Office16/POWERPNT.EXE',
           'Microsoft Office/Office16/POWERPNT.EXE') if (root/relative).is_file()), None)
    return {'available': bool(powershell.is_file() and office and office.is_file()),
            'powershell': str(powershell), 'powerpoint': str(office) if office else None,
            'scope': 'installed_program_discovery_actual_export_required',
            'host': 'Windows', 'caller_platform': system}


def _host_path(path):
    path = Path(path).resolve()
    if platform.system() == 'Windows':
        return str(path)
    suffix = []
    while not path.exists():
        suffix.insert(0, path.name)
        path = path.parent
    anchor = subprocess.run(['wslpath', '-w', str(path)], check=True,
                            capture_output=True, text=True, stdin=subprocess.DEVNULL).stdout.strip()
    return ntpath.join(anchor, *suffix) if suffix else anchor


def _local_path(path):
    """Resolve a recorded Windows job path through the same WSL bridge."""
    spelling = str(path)
    if platform.system() != 'Windows' and (ntpath.splitdrive(spelling)[0] or spelling.startswith('\\\\')):
        spelling = subprocess.run(['wslpath', '-u', spelling], check=True,
            capture_output=True, text=True, stdin=subprocess.DEVNULL).stdout.strip()
    return Path(spelling).resolve()


def render(pptx, output, capability, width=2000, timeout=300,save_native_copy=False):
    pptx, output = Path(pptx).resolve(), Path(output).resolve()
    if output.exists():
        raise FileExistsError(output)
    if not capability.get('available'):
        raise RuntimeError('Windows Office bridge is not discovered')
    # A migrated task may carry a WSL spelling into a Windows process. Discover
    # the actual caller's installed bridge before launching; no executable
    # alias or old-host availability flag is evidence of a current export.
    if capability.get('caller_platform')!=platform.system():
        actual=discover()
        if not actual.get('available'):raise RuntimeError('Task moved to another caller; discover its actual installed Office bridge')
        capability=actual
    before = digest(pptx)
    wrapper = Path(__file__).resolve().parent.parent/'run_office_job.ps1'
    arguments = [capability['powershell'], '-NoProfile', '-ExecutionPolicy', 'Bypass',
                 '-File', _host_path(wrapper), '-Kind', 'render', '-InputPptx', _host_path(pptx),
                 '-OutDir', _host_path(output), '-PreviewWidth', str(width),
                 '-TimeoutSeconds', str(timeout)]
    if save_native_copy:arguments.append('-SaveNativeCopy')
    if capability.get('svg_converter'):
        arguments.append('-ExportSvg')
    # The wrapper stages a byte-identical local copy; only its own document is
    # closed. It never quits/kills shared Office or changes machine policies.
    process = subprocess.run(arguments, capture_output=True, timeout=timeout+30,
                             stdin=subprocess.DEVNULL, encoding='utf-8', errors='replace')
    if process.returncode:
        failure_mark='LOTRAAI_OFFICE_FAILURE:'
        lines=[line[len(failure_mark):] for line in process.stdout.splitlines() if line.startswith(failure_mark)]
        if len(lines)==1:
            job=json.loads(lines[0])
            if job.get('input_pptx_hash')==before:
                try:
                    return recover_pdf_preview(pptx,output,job['staging_directory'],width=width)
                except (ValueError,OSError):
                    pass
        output.parent.mkdir(parents=True, exist_ok=True)
        write_json(output.parent/(output.name+'-office-failure.json'),
                   {'stage': 'actual_office_export', 'pptx_hash': before,
                    'exit_code': process.returncode, 'stdout': process.stdout[-2500:],
                    'stderr': process.stderr[-2500:], 'retry': 'inspect preserved staging before retry'})
        raise RuntimeError('Actual Windows Office export failed: '+process.stderr[-600:])
    receipt = read_json(output/'render.json')
    if receipt.get('pptx_hash') != before or digest(pptx) != before:
        raise RuntimeError('Office export input binding changed')
    receipt['mapping'] = 'fixed_canvas_no_content_registration'
    receipt['route'] = 'Windows_Office_program_interface_from_'+platform.system()
    receipt['native_editing_claim'] = 'not_inferred_from_export'
    from .transient_export import clean_office_stage,register
    receipt['private_staging_cleanup']=clean_office_stage(output,pptx)
    write_json(output/'render.json', receipt)
    register(output,pptx,receipt)
    return receipt


def recover_pdf_preview(pptx, output, staging, width=2000):
    """Resume the bound actual Office PDF after its PNG export failed.

    No Office reopen, PPT rebuild, alternate PPT layout, or content alignment.
    The original failed job and direct-PNG failure remain explicit evidence.
    """
    from pptx import Presentation
    from .render import rasterize_pdf
    pptx,output=Path(pptx).resolve(),Path(output).resolve()
    if output.exists():raise FileExistsError(output)
    staging=_local_path(staging);before=digest(pptx)
    job=read_json(staging/'job-receipt.json')
    execution=read_json(staging/'execute-status.json')
    failure=read_json(staging/'result/failure.json')
    if (job.get('kind')!='render' or job.get('status')!='FAILED'
            or job.get('input_pptx_hash')!=before or digest(staging/'Input.pptx')!=before
            or execution.get('status')!='FAILED' or failure.get('stage')!='png'
            or (staging/'result/pdf-failure.json').exists()):
        raise ValueError('Resume only a hash-bound successful Office PDF with failed direct PNG')
    pdf=staging/'result/Final.pdf'
    if not pdf.is_file() or not pdf.stat().st_size:raise ValueError('No actual saved Office PDF')
    import pymupdf
    prs=Presentation(pptx)
    with pymupdf.open(pdf) as document:
        if len(document)!=len(prs.slides):raise ValueError('Actual PDF page count mismatch')
        for page in document:
            if abs(page.rect.width/page.rect.height-prs.slide_width/prs.slide_height)>1e-4:
                raise ValueError('Actual PDF canvas differs from saved native slide')
        pdf_metadata=document.metadata
    output.mkdir(parents=True)
    shutil.copyfile(pdf,output/'Final.pdf')
    pages=rasterize_pdf(output/'Final.pdf',output,width,len(prs.slides),
        round(width*prs.slide_height/prs.slide_width))
    session_file=staging/'result/office-session.json'
    session=read_json(session_file) if session_file.is_file() else {}
    receipt={'status':'PASS','scope':'actual_office_pdf_layout_only','backend':'PowerPoint',
        'backend_version':session.get('backend_version','NOT_RECORDED'),
        'office_build':session.get('office_build','NOT_RECORDED'),'platform':'Windows',
        'pptx_hash':before,'pdf':{'status':'EXPORTED','path':'Final.pdf'},
        'pdf_hash':digest(output/'Final.pdf'),'preview_hash':pages[0]['sha256'],'pages':pages,
        'preview_export':{'backend':'PyMuPDF','version':pymupdf.VersionBind,
            'source':'actual hash-bound PowerPoint PDF; direct Office PNG failed'},
        'mapping':'fixed_canvas_no_content_registration','pdf_metadata':pdf_metadata,
        'route':'Windows_Office_PDF_with_bound_raster_preview_from_'+platform.system(),
        'direct_png_failure':failure,'original_job':job,
        'renderer_script_hash':digest(staging/'render_office.ps1'),
        'text_metrics_status':'NOT_RECORDED_AFTER_PNG_FAILURE',
        'native_editing_claim':'not_inferred_from_export'}
    if digest(pptx)!=before:raise ValueError('Original PPT changed while resuming actual PDF')
    write_json(output/'render.json',receipt)
    write_json(output/'resumed-actual-office-pdf.json',{
        'pptx_hash':before,'pdf_hash':receipt['pdf_hash'],'staging':str(staging),
        'scope':'Saved actual drawing continuation; no rebuild or Office retry',
        'failure':failure,'job':job})
    return receipt
