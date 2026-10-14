"""Bind an actual host-UI Office PDF export to a saved candidate.

This ingests a host observation, not a claim that Python controlled Office.
An exported PDF proves rendering only; edit/save/reopen needs separate evidence.
"""
from pathlib import Path
import tempfile
from pptx import Presentation
from .common import digest, write_json, environment
from .render import rasterize_pdf

RECEIPT_VERSION = 'host-render-0.2.1'
POWERPOINT_APPLICATIONS = {'PowerPoint', 'Microsoft PowerPoint', 'Microsoft PowerPoint macOS'}


def _validate_observation(pptx,pdf,observation):
    required=('method','application','document_path','pptx_hash','pdf_hash','observed_actions','export_mode')
    if not isinstance(observation,dict) or any(not observation.get(k) for k in required):
        raise ValueError('Incomplete actual host render observation')
    if observation['method']!='host_ui_observed':raise ValueError('Not an actual host-UI observation')
    if observation['application'] not in POWERPOINT_APPLICATIONS:
        raise ValueError('Host target render requires an observed PowerPoint application')
    actions=observation['observed_actions']
    if not isinstance(actions,list) or any(not isinstance(a,str) or not a.strip() for a in actions):
        raise ValueError('Observed actions must be a nonempty list of actual UI observations')
    if Path(observation['document_path']).resolve()!=pptx.resolve():raise ValueError('Observed a different document')
    if observation['pptx_hash']!=digest(pptx) or observation['pdf_hash']!=digest(pdf):raise ValueError('Changed host render artifacts')
    if observation['export_mode']!='local_print_pdf':raise ValueError('Expected local PDF export')


def import_render(pptx,pdf,observation,output,width):
    pptx,pdf,output=map(Path,(pptx,pdf,output))
    if output.exists():raise FileExistsError(output)
    _validate_observation(pptx,pdf,observation)
    prs=Presentation(pptx)
    pages=rasterize_pdf(pdf,output,width,len(prs.slides),round(width*prs.slide_height/prs.slide_width))
    _validate_observation(pptx,pdf,observation)
    receipt={'status':'PASS','scope':'host_observed_target_render_only',
             'receipt_version':RECEIPT_VERSION,'backend':'PowerPoint',
             'application':observation['application'],'office_build':observation.get('office_build'),
             'pptx_hash':digest(pptx),'pdf_hash':digest(pdf),'pdf_path':str(pdf.resolve()),
             'preview_hash':pages[0]['sha256'],'pages':pages,'environment':environment(),
             'host_observation':observation,'native_edit_operations':'UNVERIFIED',
             'mapping':'fixed_canvas_no_content_registration'}
    write_json(output/'render.json',receipt)
    return receipt


def verify_render(pptx,preview,receipt):
    """Recheck persisted render evidence, never infer UI actions from PDF metadata.

    The host observation remains an explicit trust boundary. macOS print PDFs
    can have an empty creator and a generic Quartz producer. Their metadata
    alone neither proves nor disproves that PowerPoint exported them.
    """
    pptx,preview=Path(pptx),Path(preview)
    if receipt.get('receipt_version')!=RECEIPT_VERSION or receipt.get('scope')!='host_observed_target_render_only':
        raise ValueError('Unsupported host target render receipt')
    if receipt.get('status')!='PASS' or receipt.get('backend')!='PowerPoint':
        raise ValueError('Unsuccessful host target render receipt')
    pdf=Path(receipt['pdf_path']);observation=receipt.get('host_observation')
    _validate_observation(pptx,pdf,observation)
    if (receipt.get('pptx_hash')!=digest(pptx) or receipt.get('pdf_hash')!=digest(pdf)
            or receipt.get('preview_hash')!=digest(preview)
            or receipt.get('application')!=observation['application']
            or receipt.get('office_build')!=observation.get('office_build')):
        raise ValueError('Host render receipt has inconsistent artifact or application binding')
    pages=receipt.get('pages')
    if not isinstance(pages,list) or not pages or pages[0].get('sha256')!=digest(preview):
        raise ValueError('Host render receipt has no matching first-page preview')
    width=pages[0]['canvas'][0];prs=Presentation(pptx)
    with tempfile.TemporaryDirectory(prefix='lotraai-host-render-check-') as directory:
        fresh=rasterize_pdf(pdf,directory,width,len(prs.slides),round(width*prs.slide_height/prs.slide_width))
    if fresh!=pages:
        raise ValueError('Host preview does not match fresh fixed-canvas PDF rasterization')
    _validate_observation(pptx,pdf,observation)
    return {'status':'PASS','scope':'target_render_only','receipt_version':RECEIPT_VERSION,
            'host_observation_required':True,'native_edit_operations':'UNVERIFIED'}
