"""Capability-qualified local presentation applications, with explicit receipts.

Discovery is not readiness. Only a real PDF and an editable save/reopen of an
owned fixture qualify an adapter. Each delivery still gets its own roundtrip.
"""
import json
import os
import platform
from pathlib import Path
import shutil
import subprocess
import uuid
import zipfile
from .common import digest,identity,read_json,write_json

POLICY='presentation-backend-1.36.0'
NATIVE=('PowerPoint','WPS')
NAMES={'auto':'auto','powerpoint':'PowerPoint','wps':'WPS','libreoffice':'LibreOffice'}


def preferences(project,**updates):
    path=Path(project).resolve()/'.lotraAI-build/Presentation-preferences.json'
    value=read_json(path) if path.is_file() else {'backend':'auto','output_root':None}
    if updates:
        if updates.get('backend') not in (None,*NAMES):raise ValueError('Unknown presentation backend preference')
        for key,item in updates.items():
            if key not in ('backend','output_root'):raise ValueError('Unknown presentation preference')
            if item is not None:value[key]=str(Path(item).expanduser().resolve()) if key=='output_root' else item
        write_json(path,value)
    return value


def discover():
    result=[]
    if platform.system()=='Windows':
        import winreg
        shell=Path(os.environ.get('WINDIR','C:/Windows'))/'System32/WindowsPowerShell/v1.0/powershell.exe'
        for name,prog in (('PowerPoint','PowerPoint.Application'),('WPS','KWPP.Application')):
            try:
                with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT,prog+'\\CLSID') as key:clsid=winreg.QueryValue(key,None)
                command=None
                for view in (winreg.KEY_WOW64_64KEY,winreg.KEY_WOW64_32KEY):
                    try:
                        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT,'CLSID\\'+clsid+'\\LocalServer32',0,winreg.KEY_READ|view) as key:command=winreg.QueryValue(key,None)
                        break
                    except OSError:pass
                if command is None:continue
                # Registry values contain a quoted executable plus optional switches.
                exe=command.split('"')[1] if command.startswith('"') else command.split(' /')[0].strip()
                if not Path(exe).is_file() or not shell.is_file():continue
                stat=Path(exe).stat()
                # Security labels can change without modifying program bytes.
                acl=subprocess.run(['icacls',exe],capture_output=True,timeout=15).stdout.hex()
                token=subprocess.run(['whoami','/groups','/fo','csv','/nh'],capture_output=True,timeout=15).stdout.hex()
                result.append({'backend':name,'executable':exe,'powershell':str(shell),'prog_id':prog,
                    'platform':'Windows','stamp':[stat.st_size,stat.st_mtime_ns,stat.st_ctime_ns],
                    'security_context':identity([os.environ.get('USERNAME'),acl,token]),
                    'scope':'DISCOVERED_NOT_QUALIFIED'})
            except (OSError,subprocess.SubprocessError):continue
    elif platform.system()=='Darwin':
        import plistlib
        for app in (Path('/Applications/Microsoft PowerPoint.app'),Path.home()/'Applications/Microsoft PowerPoint.app'):
            meta=app/'Contents/Info.plist'
            if meta.is_file():
                value=plistlib.loads(meta.read_bytes())
                result.append({'backend':'PowerPoint','application':str(app),'platform':'Darwin',
                    'executable':str(app/'Contents/MacOS/Microsoft PowerPoint'),
                    'stamp':[digest(meta),value.get('CFBundleVersion')],
                    'scope':'EXPERIMENTAL_ADAPTER_REQUIRES_THIS_MAC_PROBE'})
                break
    return result


def _code():
    root=Path(__file__).parent
    return {p.name:digest(p) for p in (Path(__file__),root.parent/'native_presentation.ps1',
            root.parent/'native_powerpoint.applescript',root/'office_compatibility.py',root/'render.py',root/'geometry.py',root/'mathml.py')}


def fixture(path):
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
    from .geometry import set_native_route
    from .mathml import insert_math
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    box=slide.shapes.add_textbox(Inches(.4),Inches(.3),Inches(8),Inches(.6))
    box.name='lotra-native-fixture';box.text='lotraAI editable native presentation 123'
    box.text_frame.paragraphs[0].font.size=Pt(24)
    shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.5),Inches(1.3),Inches(2),Inches(.7));shape.text='Native shape'
    conn=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(.5),Inches(2.4),Inches(4),Inches(3))
    set_native_route(conn,[[.5,2.4],[1.5,2.5],[2,2.2],[4,3]],914400)
    math=slide.shapes.add_textbox(Inches(.5),Inches(3.3),Inches(4),Inches(.8))
    insert_math(math,{'op':'fraction','numerator':{'op':'token','text':'x'},'denominator':{'op':'token','text':'2'}},size=28)
    table=slide.shapes.add_table(2,2,Inches(5),Inches(3.3),Inches(3),Inches(1)).table
    for row in range(2):
        for column in range(2):table.cell(row,column).text='Cell '+str(row+1)+','+str(column+1)
    prs.save(path)


def _math_count(path):
    from lxml import etree as E
    from .mathml import NS
    with zipfile.ZipFile(path) as z:
        return sum(len(E.fromstring(z.read(n)).xpath('//m:oMath',namespaces=NS))
            for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml'))


def _object_names(path):
    from collections import Counter
    from lxml import etree as E
    from .mathml import NS
    result={}
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not name.startswith('ppt/slides/slide') or not name.endswith('.xml'):continue
            tree=E.fromstring(z.read(name))
            for alternate in tree.xpath('//mc:AlternateContent',namespaces=NS):
                branches=list(alternate)
                chosen=next((n for n in branches if E.QName(n).localname=='Choice'),branches[0] if branches else None)
                for branch in branches:
                    if branch is not chosen:alternate.remove(branch)
            tables=[[[''.join(cell.xpath('.//a:t/text()',namespaces=NS)) for cell in row.findall('a:tc',NS)]
                for row in table.findall('a:tr',NS)] for table in tree.findall('.//a:tbl',NS)]
            result[name]={'names':Counter(tree.xpath('//p:cNvPr/@name',namespaces=NS)),'table_cells':tables}
    return result


def _native(pptx,output,capability,width=1440,timeout=120,roundtrip=False,require_pdf=True):
    from pptx import Presentation
    from .render import rasterize_pdf
    from .office_compatibility import validate
    pptx=Path(pptx).resolve(strict=True);output=Path(output).resolve()
    if output.exists():raise FileExistsError(output)
    validate(pptx);before=digest(pptx)
    output.mkdir(parents=True);job_id=uuid.uuid4().hex
    stage=output/'native-job';stage.mkdir()
    staged=stage/('i-'+job_id[:12]+'.pptx');shutil.copyfile(pptx,staged)
    if digest(staged)!=before:raise ValueError('Input changed while staging native validation')
    roundtrip_file=stage/('r-'+job_id[:12]+'.pptx')
    request={'job_id':job_id,'backend':capability['backend'],'input':str(staged),
        'pdf':str(output/'Final.pdf') if require_pdf else '', 'roundtrip':str(roundtrip_file) if roundtrip else '',
        'result':str(stage/'Result.json')}
    write_json(stage/'Request.json',request)
    try:
        if os.name=='nt':
            import ctypes
            def office_path(path):
                path=Path(path)
                if len(str(path))<=240:return str(path)
                # Preserve .pptx: a DOS filename such as INPUT~1.PPT can select
                # an unintended converter in Office even though its bytes are OOXML.
                existing=path.parent
                buffer=ctypes.create_unicode_buffer(32768)
                count=ctypes.windll.kernel32.GetShortPathNameW(str(existing),buffer,len(buffer))
                short=Path(buffer.value) if count and count<len(buffer) else existing
                answer=str(short/path.name)
                if len(answer)>240:raise ValueError('Office path exceeds its safe native limit; choose a shorter project/output directory')
                return answer
            for name in ('input','pdf','roundtrip'):
                if request[name]:request[name]=office_path(request[name])
            write_json(stage/'Request.json',request)
        if capability['platform']=='Windows':
            command=[capability['powershell'],'-NoProfile','-ExecutionPolicy','Bypass','-File',
                str(Path(__file__).parent.parent/'native_presentation.ps1'),'-Request',str(stage/'Request.json')]
        else:
            # Arguments are data, never interpolated into AppleScript source.
            command=['/usr/bin/osascript',str(Path(__file__).parent.parent/'native_powerpoint.applescript'),
                str(staged),request['pdf'],request['roundtrip'],job_id]
        # File-backed logs keep the deadline finite if a shared GUI application
        # inherits helper handles. Kill only the owned helper, never Office.
        with (stage/'Helper.log').open('wb') as log:
            proc=subprocess.Popen(command,stdout=log,stderr=log,stdin=subprocess.DEVNULL,
                creationflags=subprocess.CREATE_NO_WINDOW if os.name=='nt' else 0)
            try:proc.wait(timeout=timeout)
            except subprocess.TimeoutExpired:
                proc.kill();proc.wait(timeout=5);raise
        helper_log=(stage/'Helper.log').read_text(encoding='utf-8',errors='replace')[-3000:]
        if capability['platform']=='Darwin' and proc.returncode==0:
            write_json(stage/'Result.json',{'status':'PASS','backend':'PowerPoint','stage':'complete',
                'job_id':job_id,'native_save':'PASS' if roundtrip else 'NOT_RUN',
                'edit_save_reopen':'PASS' if roundtrip else 'NOT_RUN','backend_version':capability['stamp'][1]})
        result=read_json(stage/'Result.json') if (stage/'Result.json').is_file() else {}
        if digest(staged)!=before:raise ValueError('Native application changed its immutable staged input')
        if proc.returncode or result.get('status')!='PASS' or result.get('job_id')!=job_id:
            raise RuntimeError(str(result or {'stage':'helper','exit_code':proc.returncode,
                'helper_log':helper_log}))
        prs=Presentation(pptx)
        if roundtrip:
            saved=roundtrip_file;validate(saved)
            if (len(Presentation(saved).slides)!=len(prs.slides) or _math_count(saved)<_math_count(pptx)
                    or _object_names(saved)!=_object_names(pptx)):
                raise ValueError('Native roundtrip lost slides, named native descendants or genuine Office Math')
            result.update(roundtrip_sha256=digest(saved),native_math_preserved=True,native_descendants_preserved=True,native_table_cells_preserved=True)
        if digest(pptx)!=before:raise ValueError('Native backend changed immutable input')
        if not require_pdf:
            receipt={'schema':POLICY,'status':'PASS','scope':'actual_native_save_reopen',
                'backend':capability['backend'],'pptx_hash':before,'native_job':result,
                'native_editing_claim':'PASS_OWNED_COPY_EDIT_SAVE_REOPEN','pdf_export':'NOT_RUN',
                'renderer_script_hash':identity(_code())}
            write_json(output/'Native-save.json',receipt);return receipt
        pdf=output/'Final.pdf'
        if not pdf.is_file() or not pdf.stat().st_size:raise ValueError('Native application produced no PDF')
        pages=rasterize_pdf(pdf,output,width,len(prs.slides),round(width*prs.slide_height/prs.slide_width))
        receipt={'schema':POLICY,'status':'PASS','scope':'actual_native_backend_render','backend':capability['backend'],
            'backend_version':result.get('backend_version'),'office_build':result.get('office_build'),
            'platform':capability['platform'],'backend_identity':identity(capability),'pptx_hash':before,
            'pdf':{'status':'EXPORTED','path':'Final.pdf'},'pdf_hash':digest(pdf),
            'preview_hash':pages[0]['sha256'],'pages':pages,'mapping':'fixed_canvas_no_content_registration',
            'native_editing_claim':'PASS_OWNED_COPY_EDIT_SAVE_REOPEN' if roundtrip else 'NOT_RUN',
            'native_job':result,'renderer_script_hash':identity(_code())}
        from .render import font_environment
        receipt.update(font_environment(output))
        metrics=output/'text-bounds.json'
        if metrics.is_file():
            value=read_json(metrics);value['pptx_hash']=before;write_json(metrics,value)
        write_json(output/'render.json',receipt)
        from .transient_export import register
        register(output,pptx,receipt)
        return receipt
    except Exception as error:
        write_json(output/'Backend-failure.json',{'status':'FAILED','backend':capability['backend'],
            'input_sha256':before,'stage':request,'reason':str(error)[-4000:],
            'progress':read_json(stage/'Progress.json') if (stage/'Progress.json').is_file() else None,
            'timeout':isinstance(error,subprocess.TimeoutExpired),'user_input_unchanged':digest(pptx)==before})
        raise


def probe(project,capability,refresh=False):
    project=Path(project).resolve()
    key=identity({'application':capability,'code':_code(),'policy':POLICY})
    root=project/'.lotraAI-build/backend-probes'/key[:12]
    proof=root/'Qualification.json'
    if proof.is_file() and not refresh:
        old=read_json(proof)
        if old.get('status')=='READY' and old.get('probe_sha256')==digest(root/'render/render.json'):
            return old
    if root.exists():root=root.with_name(key[:12]+'-'+uuid.uuid4().hex[:8]);proof=root/'Qualification.json'
    root.mkdir(parents=True);fixture(root/'Fixture.pptx')
    try:
        receipt=_native(root/'Fixture.pptx',root/'render',capability,roundtrip=True)
        import pymupdf
        with pymupdf.open(root/'render/Final.pdf') as pdf:
            text=''.join(p.get_text() for p in pdf)
        import unicodedata
        text=unicodedata.normalize('NFKC',text)
        if 'lotraAI' not in text or 'x' not in text or '2' not in text:
            raise ValueError('Fixture text/native Math missing from actual PDF')
        value={**capability,'status':'READY','policy':POLICY,'native_math':True,
            'backend_version':receipt.get('backend_version'),'office_build':receipt.get('office_build'),
            'pdf_export':True,'native_save':True,'edit_save_reopen':True,'generator_profile':'office-safe-native-1.36.0',
            'qualification':str(proof),'probe_sha256':digest(root/'render/render.json'),
            'scope':'this host/application/fixture; each final file needs independent checks'}
    except Exception as error:
        value={**capability,'status':'CAPABILITY_BLOCKED','policy':POLICY,'reason':str(error)[-1800:]}
    write_json(proof,value)
    return value


def select(project,requested=None,refresh=False,libreoffice=None):
    project=Path(project).resolve();prefs=preferences(project)
    choice=NAMES.get(requested or prefs.get('backend','auto'))
    if choice is None:raise ValueError('Unknown requested presentation backend')
    candidates=discover();failures=[]
    if candidates:
        from .render import font_environment
        fonts=font_environment(project/'.lotraAI-build')
        candidates=[{**c,**fonts} for c in candidates]
    # An explicit choice is binding. Auto prefers the OS .pptx association among native adapters.
    if choice=='auto' and platform.system()=='Windows':
        try:
            import winreg
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER,r'Software\Microsoft\Windows\CurrentVersion\Explorer\FileExts\.pptx\UserChoice') as key:
                associated=winreg.QueryValueEx(key,'ProgId')[0].lower()
            preferred='WPS' if 'wpp' in associated or 'kwpp' in associated else 'PowerPoint'
            candidates.sort(key=lambda c:c['backend']!=preferred)
        except OSError:pass
    for cap in candidates:
        if choice not in ('auto',cap['backend']):continue
        value=probe(project,cap,refresh)
        if value['status']=='READY':
            value={**value,'selection':{'requested':choice,'failed_candidates':failures}}
            write_json(project/'.lotraAI-build/Presentation-backend.json',value);return value
        failures.append(value)
    if choice in ('auto','LibreOffice') and libreoffice:
        # LibreOffice's own bounded Math/export checks remain in its adapter.
        return {'backend':'LibreOffice','status':'READY','policy':POLICY,'selection':{'requested':choice,
                'failed_candidates':failures},'scope':'declared fallback; actual PDF/native save checks still required',**libreoffice}
    return {'status':'CAPABILITY_BLOCKED','policy':POLICY,'backend':choice,'failed_candidates':failures,
        'reason':'Requested local application is absent or its real PDF/edit/save/reopen probe failed'}


def render(pptx,output,capability,width=1440,timeout=120):
    return _native(pptx,output,capability,width,timeout)


def verify_render(pptx,receipt_file):
    import tempfile
    from pptx import Presentation
    from .render import rasterize_pdf
    receipt_file=Path(receipt_file);r=read_json(receipt_file);root=receipt_file.parent
    if (r.get('schema')!=POLICY or r.get('scope')!='actual_native_backend_render' or r.get('status')!='PASS'
            or r.get('backend') not in NATIVE or r.get('pptx_hash')!=digest(pptx)
            or r.get('renderer_script_hash')!=identity(_code())):raise ValueError('Unsupported/stale native backend proof')
    if digest(root/'Final.pdf')!=r['pdf_hash'] or digest(root/'Final_preview.png')!=r['preview_hash']:
        raise ValueError('Native render artifacts changed')
    prs=Presentation(pptx);width=r['pages'][0]['canvas'][0]
    with tempfile.TemporaryDirectory(prefix='lotra-native-check-') as folder:
        pages=rasterize_pdf(root/'Final.pdf',folder,width,len(prs.slides),round(width*prs.slide_height/prs.slide_width))
    if pages!=r['pages']:raise ValueError('Native preview differs from actual PDF')
    return {'status':'PASS','scope':'actual_native_backend_render','backend':r['backend'],
        'native_edit_verified':r.get('native_editing_claim')=='PASS_OWNED_COPY_EDIT_SAVE_REOPEN'}


def validate_delivery(pptx,audits,caps,*,destination=None):
    """Fresh input-specific qualification, separate from the environment fixture."""
    cap=caps.get('presentation_backend',caps)
    audits=Path(audits).resolve();audits.mkdir(parents=True,exist_ok=True)
    if destination is not None:
        chosen=Path(destination).resolve(strict=True)
        work=chosen/('.lotra-native-validation-'+uuid.uuid4().hex[:12])
        if Path(pptx).resolve().parent!=chosen or work.exists():raise ValueError('Unbound final-file validation destination')
    else:chosen=None;work=None
    def retain(actual,final):
        if actual==final:return
        actual,final=Path(actual),Path(final)
        if actual.is_symlink() or actual.resolve().parent!=chosen or final.exists():raise ValueError('Unexpected native evidence transfer boundary')
        from .transient_export import release_task_pdfs,VERSION as PDF_POLICY
        # These validation PDFs have no remaining consumers. Release them while
        # their original ownership paths and saved inputs are still available.
        release_task_pdfs(actual)
        files={p.relative_to(actual).as_posix():digest(p) for p in actual.rglob('*') if p.is_file()}
        if any(p.is_symlink() for p in actual.rglob('*')):raise ValueError('Linked native evidence must be retained for inspection')
        shutil.copytree(actual,final)
        if files!={p.relative_to(final).as_posix():digest(p) for p in final.rglob('*') if p.is_file()}:
            raise ValueError('Cross-volume native evidence transfer changed bytes')
        rebound=[]
        for marker in final.rglob('Pdf-lifecycle.json'):
            owner=read_json(marker);old_root=actual/marker.parent.relative_to(final)
            if owner.get('schema')!=PDF_POLICY:continue
            if owner.get('status')!='DISPOSED' or owner.get('root')!=str(old_root):
                raise ValueError('Transferred PDF ownership is not the exact disposed validation lease')
            original=Path(owner['pptx'])
            if original.is_relative_to(actual):owner['pptx']=str(final/original.relative_to(actual))
            owner['root']=str(marker.parent)
            owner['evidence_transfer']={'from':str(old_root),'scope':'owned validation evidence; immutable hashes preserved'}
            write_json(marker,owner);rebound.append(marker.relative_to(final).as_posix())
        write_json(final/'Evidence-transfer.json',{'status':'HASH_VERIFIED_OWNED_EVIDENCE_TRANSFER',
            'original_files':files,'disposed_pdf_roots_rebound':rebound})
        # Only this newly owned child on the explicitly selected destination is
        # released, after the application exited and every copied hash matched.
        if actual.resolve().parent!=chosen or not actual.name.startswith('.lotra-native-validation-'):
            raise ValueError('Native evidence cleanup escaped selected destination')
        shutil.rmtree(actual)
    if cap.get('backend') not in NATIVE:
        from .office_compatibility import validate
        validate(pptx)
        executable=cap.get('executable') or caps.get('libreoffice',{}).get('executable')
        if not executable:raise ValueError('Final file has no usable save/reopen backend')
        from .libreoffice_backend import save_reopen
        final=audits/'libreoffice-delivery';actual=work or final
        proof=save_reopen(pptx,actual,executable)
        retain(actual,final)
        return {'status':'ALTERNATIVE_SAVE_REOPEN_PASS','backend':'LibreOffice','input_sha256':digest(pptx),
            'target_powerpoint':'UNVERIFIED','proof':proof,'tested_destination_root':str(chosen.parent) if chosen else str(audits),
            'scope':'LibreOffice save/reopen; target native application unavailable; no Office editability qualification'}
    final=audits/'native-delivery';actual=work or final
    result=_native(pptx,actual,cap,roundtrip=True,require_pdf=False)
    retain(actual,final)
    return {'status':'PASS','backend':cap['backend'],'input_sha256':digest(pptx),
        'receipt':str(final/'Native-save.json'),
        'receipt_sha256':digest(final/'Native-save.json'),'tested_destination_root':str(chosen.parent) if chosen else str(audits),
        'scope':'owned copy native save, edit/save/reopen; canonical bytes unchanged; PDF has a separate actual-render gate',
        'native_job':result['native_job']}
