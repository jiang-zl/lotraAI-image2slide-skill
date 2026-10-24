"""Read-only saved-package checker. Structure PASS never implies Office editing."""

from pathlib import Path
from zipfile import ZipFile
import hashlib
import posixpath
from lxml import etree as E
from .common import digest, environment, identity, read_json, safe_path
from .mathml import NS,canonical_ast
from .scene import validate

CHECKER_VERSION = "native-structure-0.2.3"


def xml(path):
    from .resource_facts import xml_part
    return xml_part(path, 'ppt/slides/slide1.xml')


def active_children(tree):
    from .native_markup import active_children as logical_children
    yield from logical_children(tree)


def inventory(root):
    result, duplicate = {}, []

    def walk(tree, parent=None):
        for n in active_children(tree):
            props = n.find(".//p:cNvPr", NS)
            if props is None:
                continue
            name = props.get("name")
            if name in result:
                duplicate.append(name)
            result[name] = {"node": n, "parent": parent, "id": props.get("id")}
            if E.QName(n).localname == "grpSp":
                walk(n, name)
    walk(root.find("p:cSld/p:spTree", NS))
    return result, duplicate


def parse_math_container(container):
    def seq(nodes):
        items = [parse(n) for n in nodes if E.QName(n).namespace == NS["m"] and not E.QName(n).localname.endswith("Pr")]
        if not items:
            raise ValueError("Empty math expression")
        return items[0] if len(items) == 1 else {"op": "seq", "items": items}

    def arg(n, tag):
        c = n.find("m:" + tag, NS)
        if c is None:
            raise ValueError("Missing math argument " + tag)
        return seq(c)

    def parse(n):
        tag = E.QName(n).localname
        if tag == "r":
            t = n.find("m:t", NS)
            if t is None or not t.text:
                raise ValueError("Empty math token")
            value={"op": "token", "text": t.text}
            normal=n.find('m:rPr/m:nor',NS)
            if normal is not None and normal.get('{'+NS['m']+'}val','1') in ('1','true','on'):
                # A source-typeface variable remains explicitly italic on a
                # normal-font OMML run. Upright normal tokens retain normal.
                pr=n.find('a:rPr',NS)
                if pr is None or pr.get('i','0') not in ('1','true','on'):value['normal']=True
            elif all(c in '0123456789., ' for c in t.text):
                # Upright mathematical digits/punctuation use m:sty, not the
                # normal-text mode that bypasses native subscript scaling.
                value['normal']=True
            return value
        specs = {"sSubSup": ("scripts", [("base", "e"), ("sub", "sub"), ("sup", "sup")]),
                 "sSub": ("sub",[("base","e"),("sub","sub")]),
                 "sSup": ("sup",[("base","e"),("sup","sup")]),
                 "f": ("fraction", [("numerator", "num"), ("denominator", "den")]),
                 "rad": ("sqrt", [("body", "e")]),
                 "nary": ("sum", [("lower", "sub"), ("upper", "sup"), ("body", "e")]),
                 "acc": ("accent", [("body", "e")])}
        if tag in specs:
            op, args = specs[tag]
            value = {"op": op, **{key: arg(n, element) for key, element in args}}
            if tag == "rad" and (n.find("m:radPr/m:degHide", NS).get(f"{{{NS['m']}}}val") not in ("1","true","on") or len(n.find("m:deg",NS))):
                raise ValueError("Unexpected radical degree")
            if tag == "nary":
                operator=n.find("m:naryPr/m:chr", NS)
                # MS-OE376 2.1.1811 documents U+222B when nary chr is omitted;
                # current PowerPoint also emits this form for actual integrals.
                symbol=operator.get(f"{{{NS['m']}}}val") if operator is not None else '∫'
                if symbol=='∏':value['op']='product'
                elif symbol!='∑':
                    # Office builds a literal integral/product with attached
                    # bounds into nary. Preserve its exact symbol and bounds.
                    return {'op':'seq','items':[{'op':'scripts','base':{'op':'token','text':symbol},
                        'sub':value['lower'],'sup':value['upper']},value['body']]}
            if tag == "acc":
                value["mark"] = n.find("m:accPr/m:chr", NS).get(f"{{{NS['m']}}}val")
            return value
        if tag == "m":
            return {"op": "matrix", "rows": [[seq(c) for c in r.findall("m:e", NS)] for r in n.findall("m:mr", NS)]}
        if tag == "d":
            left=n.find('m:dPr/m:begChr',NS);right=n.find('m:dPr/m:endChr',NS)
            left=left.get('{'+NS['m']+'}val') if left is not None else '('
            right=right.get('{'+NS['m']+'}val') if right is not None else ')'
            matrix = arg(n, "e")
            if left=='{' and right=='' and matrix['op']=='matrix':return {"op": "cases", "rows": matrix["rows"]}
            return {'op':'delimited','left':left,'right':right,'body':matrix}
        if tag=='func':return {'op':'seq','items':[arg(n,'fName'),arg(n,'e')]}
        raise ValueError("Unsupported/invalid math structure " + tag)
    return seq(container)


def check(scene, base, pptx, owners=None, phase=None, validated=False):
    if not validated:validate(scene, base)
    root = xml(pptx)
    nodes, duplicates = inventory(root)
    checks = []
    with ZipFile(pptx) as package:
        rel_root=E.fromstring(package.read("ppt/slides/_rels/slide1.xml.rels"),E.XMLParser(resolve_entities=False,no_network=True))
        media_relationships={r.get("Id"):r for r in rel_root}

    def add(id, name, passed, detail):
        if owners is not None and id not in owners:return
        placement_checks={'placement','group_bounds','trace_group_bounds','connector_geometry','connector_flip',
                          'routed_canvas_bounds','routed_transform','stCxn','endCxn'}
        if phase=='body' and name in placement_checks:return
        if phase=='placement' and name not in placement_checks|{'presence','ownership','group_transform'}:return
        checks.append({"object_id": id, "check_id": name, "status": "PASS" if passed else "FAIL", "detail": detail,
                       "evidence_ref": "ppt/slides/slide1.xml#" + id})

    add("scene", "unique_names", not duplicates, duplicates)
    expected = {o["id"]: o for o in scene["objects"]}
    backend_children=set()
    for obj in scene['objects']:
        if obj['render_type']=='trace':
            data=read_json(safe_path(base,obj['content']['geometry_ref']['path']))
            backend_children.update(obj['id']+'.tone-'+str(i) for i in range(len(data['layers'])))
    unexpected=set(nodes)-set(expected)-backend_children
    add("scene", "unexpected_visible_nodes", not unexpected, sorted(unexpected))
    native_ids=[v['id'] for v in nodes.values()]
    add('scene','unique_native_ids',len(native_ids)==len(set(native_ids)),len(native_ids))
    scale = scene["canvas"]["target_width_inches"] * 914400 / scene["canvas"]["width"]
    for id, o in expected.items():
        if owners is not None and id not in owners:continue
        entry = nodes.get(id)
        add(id, "presence", entry is not None, "Required object in active content branch")
        if entry is None:
            continue
        n, kind = entry["node"], o["render_type"]
        tag = E.QName(n).localname
        add(id, "ownership", entry["parent"] == o.get("parent"), {"actual": entry["parent"], "expected": o.get("parent")})
        hidden = n.find(".//p:cNvPr", NS).get("hidden") in ("1", "true")
        invisible = n.xpath(".//a:alpha[@val='0']", namespaces=NS)
        add(id, "visible_native", not hidden and not invisible, "No hidden/zero-opacity native placeholder")
        media = n.xpath(".//a:blip", namespaces=NS)
        if kind not in ("photo", "group"):
            add(id, "no_raster_substitute", not media and tag != "pic", "Native schematic content required")
        geom = n.find("p:spPr/a:xfrm", NS)
        imported=o.get('content',{}).get('native_snapshot',{})
        verified_import=imported.get('native_xml_sha256')==hashlib.sha256(E.tostring(n,method='c14n',exclusive=True)).hexdigest()
        if kind in ("group","trace"):
            geom = n.find("p:grpSpPr/a:xfrm", NS)
            if geom is not None:
                off, ext = geom.find("a:off", NS), geom.find("a:ext", NS)
                coff, cext = geom.find("a:chOff", NS), geom.find("a:chExt", NS)
                same_origin = off is not None and coff is not None and dict(off.attrib) == dict(coff.attrib)
                same_scale = ext is not None and cext is not None and dict(ext.attrib) == dict(cext.attrib)
                if verified_import:
                    from .boundary_geometry import _transform
                    try:
                        import numpy as np
                        matrix=_transform(geom,group=True)
                        valid_transform=np.isfinite(matrix).all() and abs(np.linalg.det(matrix[:2,:2]))>1e-15
                    except (ValueError,TypeError):valid_transform=False
                    add(id,'group_transform',bool(valid_transform),'Verified native snapshot has a finite nonsingular affine group transform; source placement/anatomy require separate checks')
                else:
                    add(id, "group_transform", same_origin and same_scale, "Generated absolute-coordinate groups must preserve child origin and scale")
                if kind=='trace' and off is not None and ext is not None:
                    actual=[int(off.get('x')),int(off.get('y')),int(ext.get('cx')),int(ext.get('cy'))]
                    add(id,'trace_group_bounds',all(abs(a-b)<=2 for a,b in zip(actual,[round(v*scale) for v in o['bbox']])),'Trace display layers share the declared coordinate frame')
                def belongs(obj):
                    p=obj.get("parent")
                    while p:
                        if p==id:return True
                        p=expected[p].get("parent")
                    return False
                leaves=[]
                for obj in expected.values():
                    if obj['render_type']=='group' or not belongs(obj):continue
                    if obj['render_type']=='connector' and obj['content'].get('points'):
                        from .geometry import native_route_spec,connector_points
                        spec=native_route_spec(connector_points(expected,obj),scale)
                        leaves.append([v/scale for v in spec['off']+spec['ext']])
                    else:leaves.append(obj['bbox'])
                if leaves and off is not None and ext is not None:
                    x0,y0=min(b[0] for b in leaves),min(b[1] for b in leaves)
                    x1,y1=max(b[0]+b[2] for b in leaves),max(b[1]+b[3] for b in leaves)
                    actual=[int(off.get("x")),int(off.get("y")),int(ext.get("cx")),int(ext.get("cy"))]
                    wanted=[round(v*scale) for v in (x0,y0,x1-x0,y1-y0)]
                    add(id,"group_bounds",all(abs(a-b)<=2 for a,b in zip(actual,wanted)),{"actual":actual,"expected":wanted})
        if geom is not None:
            rotation=round(o['style'].get('rotation',0)*60000)%21600000
            if kind=='connector' and o['content'].get('points'):
                from .geometry import connector_points,native_route_spec
                rotation=native_route_spec(connector_points(expected,o),round(scene['canvas']['target_width_inches']*914400)/scene['canvas']['width'])['rot']
            add(id, "transform_rotation", verified_import or int(geom.get("rot", "0"))%21600000 == rotation and (kind=="connector" or (geom.get("flipH", "0") in ("0", "false") and geom.get("flipV", "0") in ("0", "false"))), "Declared rotation or exact native snapshot; source instance affine placement is independently checked")
        if geom is not None and kind not in ("connector", "group"):
            off, ext = geom.find("a:off", NS), geom.find("a:ext", NS)
            actual = [int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))]
            wanted = [round(v * scale) for v in o["bbox"]]
            add(id, "placement", all(abs(a-b) <= 2 for a, b in zip(actual, wanted)), {"actual_emu": actual, "expected_emu": wanted})
        if kind == "shape":
            g = n.find("p:spPr/a:prstGeom", NS)
            if o['content'].get('connection_sites'):
                sites=n.findall('p:spPr/a:custGeom/a:cxnLst/a:cxn/a:pos',NS)
                actual=[[p.get('x'),p.get('y')] for p in sites]
                wanted=[['portX'+str(i),'portY'+str(i)] for i in range(4+len(o['content']['connection_sites']))]
                wanted_guides=[]
                for i,(u,v) in enumerate([[.5,0],[0,.5],[.5,1],[1,.5]]+o['content']['connection_sites']):
                    wanted_guides.extend([('portX'+str(i),'*/ w '+str(round(u*100000))+' 100000'),('portY'+str(i),'*/ h '+str(round(v*100000))+' 100000')])
                actual_guides=[(g.get('name'),g.get('fmla')) for g in n.findall('p:spPr/a:custGeom/a:gdLst/a:gd',NS)]
                add(id,'native_custom_ports',tag=='sp' and actual==wanted and actual_guides==wanted_guides,{'actual':actual,'expected':wanted,'units':'shape-width/height guides, not path-local literals'})
                # Compare saved contour against the declared rounded/elliptic primitive.
                from pptx import Presentation
                from pptx.enum.shapes import MSO_SHAPE
                from .geometry import ported_shape
                temp=Presentation();sh=temp.slides.add_slide(temp.slide_layouts[6]).shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,100,100)
                expected_geom=ported_shape(sh,o)
                def path_signature(node):
                    return [(E.QName(x).localname,tuple(sorted(x.attrib.items()))) for x in node.iter()]
                actual_path=n.find('p:spPr/a:custGeom/a:pathLst',NS)
                add(id,'ported_primitive_outline',actual_path is not None and path_signature(actual_path)==path_signature(expected_geom.find('a:pathLst',NS)),o['content']['preset'])
            else:
                add(id, "native_preset", tag == "sp" and g is not None and g.get("prst") == o["content"]["preset"], o["content"]["preset"])
                from .presets import guides
                actual_guides=[] if g is None else [(a.get('name'),a.get('fmla')) for a in g.findall('a:avLst/a:gd',NS)]
                add(id,'preset_adjustments',actual_guides==guides(o['content']['preset'],o['content'].get('adjustments',[])),{'actual':actual_guides})
        elif kind == "group":
            children = [name for name, info in nodes.items() if info["parent"] == id]
            add(id, "natural_parts", tag == "grpSp" and bool(children), {"children": children, "semantic_scope": "declared probe parts; independent visual review still required"})
            if o['content'].get('shared_trunk'):
                from .measure import measure
                measured=measure(scene,pptx)['objects'].get(id,{})
                contract=o['content']['shared_trunk'];actual=measured.get('route_points');wanted=contract['reference_points']
                error=max((abs(a-b) for p,q in zip(actual or [],wanted) for a,b in zip(p,q)),default=float('inf'))
                good=(measured.get('geometry_contract',{}).get('members')==contract['members'] and not measured.get('geometry_contract_error')
                      and actual is not None and len(actual)==len(wanted) and error<.1)
                add(id,'shared_trunk_native_routes',good,{'actual_prefix':actual,'expected_prefix':wanted,'max_error_px':error if error!=float('inf') else None})
        elif kind == "text":
            from .text_layout import saved_text,expected_text
            try:actual = saved_text(n,o['content'])
            except ValueError as error:
                add(id,'text_layout_mapping',False,{'reason':str(error)})
                actual=None
            c=o['content'];paragraphs=c.get('paragraphs',[c.get('runs',[])])
            wanted=expected_text(c)
            add(id, "text_exact", actual == wanted, {"actual": actual, "expected": wanted})
            wanted_math=[r['ast'] for p in paragraphs for r in p if 'ast' in r]
            zones=n.xpath('p:txBody/a:p/a14:m/m:oMath | p:txBody/a:p/a14:m/m:oMathPara/m:oMath',namespaces=NS)
            try:parsed_math=[parse_math_container(z) for z in zones]
            except (ValueError,AttributeError,TypeError):parsed_math=None
            add(id,'inline_math_exact',canonical_ast(parsed_math)==canonical_ast(wanted_math),{'expected':wanted_math,'actual':parsed_math})
        elif kind=='path':
            c=o['content'];contours=c.get('contours')
            if contours is None:contours=read_json(safe_path(base,c['geometry_ref']['path']))['contours']
            expected_ops=[]
            for contour in contours:
                from .geometry import path_operations
                expected_ops.extend((kind,tuple(round(v*100000) for p in pts for v in p) if pts else None) for kind,pts in path_operations(contour))
            paths=n.findall('p:spPr/a:custGeom/a:pathLst/a:path',NS);got=[]
            for path in paths:
                for command in path:
                    pts=command.findall('a:pt',NS)
                    got.append((E.QName(command).localname,tuple(int(v) for p in pts for v in (p.get('x'),p.get('y'))) if pts else None))
            dimensions=all(p.get('w')=='100000' and p.get('h')=='100000' for p in paths)
            add(id,'native_path_exact',tag=='sp' and dimensions and got==expected_ops,{'expected_operations':len(expected_ops),'actual_operations':len(got)})
            if c.get('connection_sites'):
                sites=n.findall('p:spPr/a:custGeom/a:cxnLst/a:cxn/a:pos',NS)
                expected_sites=[];expected_guides=[]
                for i,(u,v) in enumerate([[.5,0],[0,.5],[.5,1],[1,.5]]+c['connection_sites']):
                    expected_sites.append(['portX'+str(i),'portY'+str(i)])
                    expected_guides.extend([('portX'+str(i),'*/ w '+str(round(u*100000))+' 100000'),('portY'+str(i),'*/ h '+str(round(v*100000))+' 100000')])
                actual_guides=[(g.get('name'),g.get('fmla')) for g in n.findall('p:spPr/a:custGeom/a:gdLst/a:gd',NS)]
                add(id,'native_visible_path_ports',[[p.get('x'),p.get('y')] for p in sites]==expected_sites and actual_guides==expected_guides,'Real open-path ports in shape w/h coordinates')
            contract=c.get('attachment_contract')
            if contract:
                owner=contract['owner_id'];good=tag=='sp' and entry['parent']==owner and nodes.get(owner,{}).get('node') is not None
                add(id,'rigid_parent_attachment',good,{'owner':owner,'scope':'Native group ownership; ConnectorFormat and group translation are separately verified'})
        elif kind=='trace':
            data=read_json(safe_path(base,o['content']['geometry_ref']['path']))
            add(id,'native_trace_group',tag=='grpSp' and not media and len(list(active_children(n)))==len(data['layers']),'Measured color layers are native paths within one observed trace')
            for index,layer in enumerate(data['layers']):
                key=id+'.tone-'+str(index);item=nodes.get(key);good=item is not None and item['parent']==id
                if good:
                    element=item['node'];fill=element.find('p:spPr/a:solidFill/a:srgbClr',NS)
                    good=fill is not None and fill.get('val')==layer['fill'] and not element.xpath('.//a:blip',namespaces=NS)
                    good=good and element.find('p:nvSpPr/p:cNvPr',NS).get('hidden','0') not in ('1','true')
                    got=[]
                    for path in element.findall('p:spPr/a:custGeom/a:pathLst/a:path',NS):
                        if path.get('w')!='100000' or path.get('h')!='100000':good=False
                        for command in path:
                            pt=command.find('a:pt',NS)
                            got.append((command.tag.rsplit('}',1)[-1],None if pt is None else (int(pt.get('x')),int(pt.get('y')))))
                    wanted=[]
                    for contour in layer['contours']:
                        wanted.append(('moveTo',tuple(round(v*100000) for v in contour['points'][0])))
                        wanted.extend(('lnTo',tuple(round(v*100000) for v in p)) for p in contour['points'][1:])
                        if contour.get('closed'):wanted.append(('close',None))
                    good=good and got==wanted
                    xf=element.find('p:spPr/a:xfrm',NS)
                    actual_box=[int(xf.find('a:off',NS).get(k)) for k in ('x','y')]+[int(xf.find('a:ext',NS).get(k)) for k in ('cx','cy')]
                    good=good and all(abs(a-b)<=2 for a,b in zip(actual_box,[round(v*scale) for v in o['bbox']]))
                add(key,'trace_layer_exact',good,{'source_pixels':layer['pixel_count'],'fill':layer['fill']})
        elif kind=='connector' and o['content'].get('points'):
            from .geometry import connector_points,bounds,native_route_spec
            c=o['content'];points=connector_points(expected,o);bx,by,bw,bh=bounds(points)
            spec=native_route_spec(points,round(scene['canvas']['target_width_inches']*914400)/scene['canvas']['width'])
            add(id,'native_connector',tag==('sp' if spec['preset']=='nativeFreeform' else 'cxnSp'),'Native preset connector or unfilled native source freeform with builtin arrow ends')
            for which,end,default in [('source','stCxn',3),('target','endCxn',1)]:
                conn=n.find('p:nvCxnSpPr/p:cNvCxnSpPr/a:'+end,NS)
                if c.get(which) and c.get(which+'_anchor',True):
                    target=nodes.get(c.get(which+'_anchor_object',c[which]));site=c.get(which+'_site',default)
                    valid=conn is not None and target is not None and conn.get('id')==target['id'] and conn.get('idx')==str(site)
                else:valid=conn is None
                add(id,end,valid,{'semantic_node':c.get(which),'native_anchor':c.get(which+'_anchor',bool(c.get(which)))})
            preset=n.find('p:spPr/a:prstGeom',NS)
            guides=[] if preset is None else [(g.get('name'),g.get('fmla')) for g in preset.findall('a:avLst/a:gd',NS)]
            wanted=[('adj'+str(i),'val '+str(v)) for i,v in enumerate(spec['guides'],1)]
            from .geometry import freeform_route_exact
            valid=freeform_route_exact(n,spec) if spec['preset']=='nativeFreeform' else preset is not None and preset.get('prst')==spec['preset'] and guides==wanted and n.find('p:spPr/a:custGeom',NS) is None
            add(id,'routed_path_exact',valid,{'preset':spec['preset'],'actual_guides':guides,'expected_guides':wanted,'endpoint_quantization_px':spec['endpoint_quantization_px']})
            xf=n.find('p:spPr/a:xfrm',NS)
            actual=[] if xf is None else [int(xf.find('a:off',NS).get(k)) for k in ('x','y')]+[int(xf.find('a:ext',NS).get(k)) for k in ('cx','cy')]
            wanted_box=spec['off']+spec['ext']
            add(id,'routed_canvas_bounds',actual==wanted_box,{'actual':actual,'expected':wanted_box})
            add(id,'routed_transform',xf is not None and int(xf.get('rot','0'))==spec['rot'] and (xf.get('flipH','0') in ('1','true'))==spec['flipH'] and (xf.get('flipV','0') in ('1','true'))==spec['flipV'],'Native preset orientation')
            for end,expected_type in [('tailEnd',c.get('arrow','triangle')),('headEnd',c.get('start_arrow','none'))]:
                head=n.find('p:spPr/a:ln/a:'+end,NS)
                add(id,end,(head.get('type') if head is not None else 'none')==expected_type,expected_type)
        elif kind == "connector":
            add(id, "native_connector", tag == "cxnSp", "p:cxnSp required")
            for end, target_key, site_key, default in [("stCxn", "source", "source_site", 3), ("endCxn", "target", "target_site", 1)]:
                conn = n.find("p:nvCxnSpPr/p:cNvCxnSpPr/a:" + end, NS)
                target = nodes.get(o["content"][target_key])
                valid = conn is not None and target is not None and conn.get("id") == target["id"] and conn.get("idx") == str(o["content"].get(site_key, default))
                add(id, end, valid, {"target": o["content"][target_key]})
            tail = n.find("p:spPr/a:ln/a:tailEnd", NS)
            head = n.find("p:spPr/a:ln/a:headEnd", NS)
            add(id, "arrow_direction", tail is not None and tail.get("type") == "triangle" and (head is None or head.get("type") == "none"), "Built-in arrowhead at target only")
            def anchor(object_id, site):
                x,y,w,h = expected[object_id]["bbox"]
                sites=[(.5,0),(0,.5),(.5,1),(1,.5)]+expected[object_id]['content'].get('connection_sites',[])
                u,v=sites[site];return x+u*w,y+v*h
            c = o["content"]
            start = anchor(c["source"], c.get("source_site",3))
            end = anchor(c["target"], c.get("target_site",1))
            if geom is not None:
                off, ext = geom.find("a:off", NS), geom.find("a:ext", NS)
                actual = [int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))]
                wanted = [round(min(start[0],end[0])*scale), round(min(start[1],end[1])*scale),
                          abs(round(end[0]*scale)-round(start[0]*scale)), abs(round(end[1]*scale)-round(start[1]*scale))]
                add(id,"connector_geometry",all(abs(a-b)<=2 for a,b in zip(actual,wanted)),{"actual":actual,"expected":wanted})
                horizontal=geom.get("flipH","0") in ("1","true")
                vertical=geom.get("flipV","0") in ("1","true")
                add(id,"connector_flip",(wanted[2]==0 or horizontal==(start[0]>end[0])) and (wanted[3]==0 or vertical==(start[1]>end[1])),"Directed endpoint sign; zero extent flip has no geometric effect")
        elif kind == "bezier":
            path = n.find("p:spPr/a:custGeom/a:pathLst/a:path", NS)
            pts = [] if path is None else path.xpath("a:moveTo/a:pt | a:cubicBezTo/a:pt", namespaces=NS)
            got = [[int(p.get("x")), int(p.get("y"))] for p in pts]
            wanted = [[round(v*100000) for v in p] for p in o["content"]["points"]]
            add(id, "cubic_geometry", got == wanted and path is not None and path.find("a:cubicBezTo", NS) is not None, {"actual": got, "expected": wanted})
            paths = n.findall("p:spPr/a:custGeom/a:pathLst/a:path", NS)
            valid_path = len(paths) == 1 and path.get("w") == "100000" and path.get("h") == "100000" and path.get("stroke", "1") in ("1","true") and path.get("fill") == "none"
            valid_path = valid_path and [E.QName(e).localname for e in path] == ["moveTo","cubicBezTo"]
            add(id,"cubic_path_contract",valid_path,"One visible cubic path with fixed coordinate scale and no added segments")
        elif kind == "equation":
            math = n.findall("p:txBody/a:p/a14:m/m:oMathPara/m:oMath", NS)
            parsed = None
            error = None
            try:
                if len(math) != 1:
                    raise ValueError("One active native math container required")
                parsed = parse_math_container(math[0])
            except (ValueError, AttributeError, TypeError) as exc:
                error = str(exc)
            add(id, "math_ast_exact", canonical_ast(parsed) == canonical_ast(o["content"]["ast"]), {"actual_ast": parsed, "error": error})
            extra = n.xpath("p:txBody/a:p/a:r | p:txBody/a:p/a:fld | p:txBody/a:p/a:br", namespaces=NS)
            add(id,"math_exclusive_content",not extra and len(n.findall("p:txBody/a:p/a14:m",NS)) == 1,"Formula shape cannot contain extra ordinary text or math zones")
        elif kind == "photo":
            add(id, "permitted_photo", tag == "pic" and bool(media), "Independent permitted raster object")
            valid=False
            if len(media)==1:
                relationship=media_relationships.get(media[0].get("{"+NS["r"]+"}embed"))
                if relationship is not None and relationship.get("TargetMode")!="External":
                    resource=posixpath.normpath(posixpath.join("ppt/slides",relationship.get("Target","")))
                    if resource.startswith("ppt/media/"):
                        with ZipFile(pptx) as package:
                            valid=hashlib.sha256(package.read(resource)).hexdigest()==o["content"]["sha256"]
            add(id,"photo_payload",valid,"Saved media bytes match authorized photo")
        elif kind == "table":
            table=n.find("a:graphic/a:graphicData/a:tbl",NS)
            rows=[] if table is None else [["\n".join(''.join(p.xpath('a:r/a:t/text()',namespaces=NS)) for p in cell.findall('a:txBody/a:p',NS)) for cell in row.findall("a:tc",NS)] for row in table.findall("a:tr",NS)]
            add(id,"native_table_cells",tag=="graphicFrame" and table is not None and rows==o["content"]["rows"],{"actual":rows,"expected":o["content"]["rows"]})
            from .tables import saved_merges
            add(id,'native_table_merges',saved_merges(n,NS)==o['content'].get('merges',[]),{'actual':saved_merges(n,NS),'expected':o['content'].get('merges',[])})
        elif kind == "unresolved":
            add(id,"unresolved_object",False,"Unresolved semantics cannot pass reconstruction structure")
    fail = any(c["status"] == "FAIL" for c in checks)
    return {"checker_version": CHECKER_VERSION, "status": "FAIL" if fail else "PASS", "scope": "package_structure_only",
            "native_edit_roundtrip": "UNVERIFIED", "visual_attribution": "UNVERIFIED", "checks": checks,
            "binding": {"source_hash": scene["source"]["sha256"], "scene_hash": identity(scene), "scene_revision": scene["revision"],
                        "evidence_hash": scene["evidence"]["sha256"], "profile_hash": scene["profile"]["sha256"],
                        "component_versions": {o["id"]: o["versions"] for o in scene["objects"]},
                        "pptx_hash": digest(pptx), "environment": environment()}}
