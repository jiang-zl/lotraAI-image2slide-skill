"""Owner-based component extraction from actual saved PresentationML.

Preserve ancestor transforms and per-layer z order. A semantic component may own
several separated native fragments; it need not be a single pre-existing group.
"""
import copy
import gc
import io
import shutil
from pathlib import Path
from zipfile import ZipFile
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from .mathml import NS, el, child
from .common import digest, identity, write_json, writer, safe_path, read_json

SHAPES={'sp','grpSp','pic','cxnSp','graphicFrame'}


def replace_native_components(pptx,donor_pptx,request,output):
    """Splice verified native subtrees while retaining external root identities.

    Donor geometry must already be authored in the same slide coordinate space.
    No imported picture may substitute for a native component in this route.
    """
    if request.get('source_pptx_hash')!=digest(pptx) or request.get('donor_pptx_hash')!=digest(donor_pptx):
        raise ValueError('Stale native replacement request')
    output=Path(output)
    if output.exists():raise FileExistsError(output)
    inventory(pptx);inventory(donor_pptx)
    targets=[c['target_id'] for c in request.get('replacements',[])]
    if len(set(targets))!=len(targets):raise ValueError('Duplicate native replacement target')
    deck=Presentation(pptx);donor=Presentation(donor_pptx)
    if (deck.slide_width,deck.slide_height)!=(donor.slide_width,donor.slide_height):raise ValueError('Native component coordinate canvases differ')
    slide=deck.slides[0];source_tree=slide.shapes._spTree;donor_tree=donor.slides[0].shapes._spTree
    records=[]
    for change in request.get('replacements',[]):
        old=next((e for e in source_tree.iter() if _tag(e) in SHAPES and _id(e)==change['target_id']),None)
        new=next((e for e in donor_tree.iter() if _tag(e) in SHAPES and _id(e)==change['donor_id']),None)
        if old is None or new is None or not change.get('reason'):raise ValueError('Missing native component/replacement reason')
        if _tag(new) in ('pic','graphicFrame') or new.xpath('.//p:pic | .//p:graphicFrame'):raise ValueError('Native replacement cannot conceal raster/external content')
        xf=_transform(old)
        parent=old.getparent()
        # Absolute coordinates in the donor cannot be injected into a scaled or
        # rotated parent without an explicit transform conversion.
        ancestor=parent
        while _tag(ancestor)=='grpSp':
            if not np.allclose(_group_matrix(_transform(ancestor)),np.eye(3),atol=1e-9):
                raise ValueError('Transform donor into parent coordinates before splicing')
            ancestor=ancestor.getparent()
        old_ids={_id(e) for e in old.iter() if _tag(e) in SHAPES}
        outside=[e for e in source_tree.iter() if _tag(e) in SHAPES and e not in set(old.iter())]
        for e in outside:
            for conn in e.xpath('./p:nvCxnSpPr/p:cNvCxnSpPr/a:stCxn | ./p:nvCxnSpPr/p:cNvCxnSpPr/a:endCxn'):
                if int(conn.get('id')) in old_ids-{change['target_id']}:
                    raise ValueError('External endpoint depends on replaced child; explicit remapping required')
        clone=copy.deepcopy(new);all_ids=[_id(e) for e in source_tree.iter() if _tag(e) in SHAPES]
        next_id=max(all_ids)+1;mapping={_id(clone):change['target_id']}
        for e in clone.iter():
            if _tag(e) not in SHAPES:continue
            sid=_id(e)
            if sid not in mapping:mapping[sid]=next_id;next_id+=1
            e.find('.//p:cNvPr',NS).set('id',str(mapping[sid]))
        for conn in clone.xpath('.//a:stCxn | .//a:endCxn'):
            if int(conn.get('id')) not in mapping:raise ValueError('Donor has an external endpoint')
            conn.set('id',str(mapping[int(conn.get('id'))]))
        before=identity(old.xml);parent.replace(old,clone)
        if change.get('z_before_id'):
            sibling=next((e for e in parent if _tag(e) in SHAPES and _id(e)==change['z_before_id']),None)
            if sibling is None or sibling is clone:raise ValueError('Z-order anchor must be another direct sibling')
            parent.remove(clone);parent.insert(parent.index(sibling),clone)
        records.append({'target_id':change['target_id'],'old_xml_hash':before,'new_xml_hash':identity(clone.xml),'reason':change['reason']})
    if not records:raise ValueError('No native replacements')
    output.parent.mkdir(parents=True,exist_ok=True);deck.save(output)
    receipt={'source_pptx_hash':digest(pptx),'donor_pptx_hash':digest(donor_pptx),'output_pptx_hash':digest(output),
             'replacements':records,'status':'SAVED_REQUIRES_CONTEXT_AND_NATIVE_CHECK'}
    write_json(output.with_suffix('.native-replacements.json'),receipt);return receipt


def replace_picture_layers(pptx,request,base,output):
    """Replace named picture payloads without changing their layout or siblings."""
    if request.get('source_pptx_hash')!=digest(pptx):raise ValueError('Stale picture-layer request')
    output=Path(output)
    if output.exists():raise FileExistsError(output)
    replacements=request.get('replacements',[])
    if not replacements:raise ValueError('No picture layers requested')
    keys=[(r.get('slide',1),r['shape_id']) for r in replacements]
    if len(set(keys))!=len(keys):raise ValueError('Duplicate picture replacement')
    deck=Presentation(pptx);records=[];sample_ids=set();sample_hashes=set()
    for item in replacements:
        slide=deck.slides[item.get('slide',1)-1];sid=item['shape_id']
        matches=[e for e in slide.shapes._spTree.iter() if _tag(e)=='pic' and _id(e)==sid]
        if len(matches)!=1:raise ValueError('Picture id is missing or ambiguous')
        element=matches[0];path=safe_path(base,item['asset']['path'])
        if digest(path)!=item['asset']['sha256']:raise ValueError('Replacement asset changed')
        if not item.get('reason'):raise ValueError('Layer replacement needs provenance')
        mode=item.get('replacement_mode')
        if mode=='new_illustrative_sample':
            # A full new sample is a redraw of this picture owner, not amodal
            # recovery of the historical hidden pixels or a protected local edit.
            if item.get('completion_receipt_hash') or item.get('local_edit_receipt_hash'):
                raise ValueError('Full illustrative replacement cannot claim local-edit or completion protection')
            category=item.get('category')
            instance=item.get('semantic_instance_id')
            source_refs=item.get('source_evidence_refs')
            review=item.get('whole_subject_review')
            if (item.get('kind')!='illustrative_photo'
                    or not isinstance(category,str) or not category.strip() or len(category)>80
                    or not isinstance(instance,str)
                    or not instance or instance in sample_ids or not isinstance(source_refs,list)
                    or not source_refs or any(not isinstance(ref,str) or not ref for ref in source_refs)
                    or not isinstance(review,dict) or review.get('status')!='PASS'
                    or review.get('asset_sha256')!=digest(path)
                    or review.get('category')!=category
                    or review.get('semantic_instance_id')!=instance
                    or review.get('whole_subject_continuity')!='PASS'
                    or review.get('single_subject')!='PASS'
                    or review.get('reviewer',{}).get('kind') not in ('human','host_model')
                    or not isinstance(review.get('evidence_refs'),list)
                    or not review['evidence_refs']
                    or any(not isinstance(ref,str) or not ref for ref in review['evidence_refs'])
                    or not item.get('generation_tool')):
                raise ValueError('New illustrative sample needs category, distinct instance, source and whole-subject review')
            if digest(path) in sample_hashes:
                raise ValueError('Distinct illustrative samples cannot reuse one asset')
            sample_ids.add(instance);sample_hashes.add(digest(path))
            with Image.open(path) as candidate:
                width,height=candidate.size
                if min(width,height)<128 or max(width,height)/min(width,height)>1.5:
                    raise ValueError('Full illustrative sample lacks usable resolution or aspect ratio')
        elif item.get('local_edit_receipt_hash'):
            completion=path.with_suffix('.local-edit.json')
            if not completion.is_file() or digest(completion)!=item['local_edit_receipt_hash']:
                raise ValueError('Local raster edit receipt is missing or changed')
            proof=read_json(completion)
            picture_name=element.find('p:nvPicPr/p:cNvPr',NS).get('name')
            if (proof.get('schema_version')!='local-raster-result-1'
                    or proof.get('status')!='PIXEL_PROTECTION_PASS_PPT_CONTEXT_UNVERIFIED'
                    or proof.get('output_hash')!=digest(path)
                    or proof.get('baseline_pptx_hash')!=digest(pptx)
                    or proof.get('outside_changed_pixels')!=0
                    or proof.get('layer_id')!=picture_name
                    or item.get('layer_id')!=picture_name
                    or not proof.get('authorization_hash')):
                raise ValueError('Local raster asset is not bound to this existing picture layer')
        elif mode in (None,'amodal_completion'):
            if not item.get('completion_receipt_hash'):raise ValueError('Layer completion needs provenance')
            completion=path.with_suffix('.completion.json')
            if not completion.is_file() or digest(completion)!=item['completion_receipt_hash']:
                raise ValueError('Layer completion receipt is missing or changed')
            if read_json(completion).get('output_hash')!=digest(path):
                raise ValueError('Completion receipt belongs to another asset')
        else:
            raise ValueError('Unknown picture-layer replacement mode')
        before=identity(element.xml);blip=element.find('p:blipFill/a:blip',NS)
        if blip is None or blip.get('{'+NS['r']+'}link'):raise ValueError('Unsupported linked/missing picture payload')
        _,rid=slide.part.get_or_add_image_part(str(path));blip.set('{'+NS['r']+'}embed',rid)
        if item.get('name'):element.find('p:nvPicPr/p:cNvPr',NS).set('name',item['name'])
        records.append({'shape_id':sid,'old_xml_hash':before,'new_xml_hash':identity(element.xml),
                        'asset_hash':digest(path),'completion_receipt_hash':item.get('completion_receipt_hash'),
                        'local_edit_receipt_hash':item.get('local_edit_receipt_hash'),
                        'replacement_mode':mode or 'amodal_completion',
                        'semantic_instance_id':item.get('semantic_instance_id'),
                        'whole_subject_review':item.get('whole_subject_review')})
    output.parent.mkdir(parents=True,exist_ok=True);deck.save(output)
    receipt={'source_pptx_hash':digest(pptx),'output_pptx_hash':digest(output),'replacements':records,
             'status':'SAVED_REQUIRES_RENDER_AND_EDIT_CHECK',
             'scope':'Full illustrative replacements authorize visible changes only inside named picture owners; no hidden-original recovery is claimed'}
    write_json(output.with_suffix('.layers.json'),receipt);return receipt


def _tag(element):return element.tag.rsplit('}',1)[-1]


def _id(element):
    props=element.find('.//p:cNvPr',NS)
    return int(props.get('id')) if props is not None else None


def _transform(element):
    tag=_tag(element)
    p=element.find('p:grpSpPr/a:xfrm',NS) if tag=='grpSp' else element.find('p:xfrm',NS) if tag=='graphicFrame' else element.find('p:spPr/a:xfrm',NS)
    return p


def _group_matrix(xfrm):
    off=xfrm.find('a:off',NS);ext=xfrm.find('a:ext',NS)
    co=xfrm.find('a:chOff',NS);ce=xfrm.find('a:chExt',NS)
    ox,oy=float(off.get('x')),float(off.get('y'));w,h=float(ext.get('cx')),float(ext.get('cy'))
    cw,ch=float(ce.get('cx')),float(ce.get('cy'))
    if min(cw,ch)<=0:raise ValueError('Degenerate saved group transform')
    m=np.array([[w/cw,0,ox-float(co.get('x'))*w/cw],[0,h/ch,oy-float(co.get('y'))*h/ch],[0,0,1.]])
    return _rotation_flip(xfrm,ox,oy,w,h)@m


def _rotation_flip(xfrm,x,y,w,h):
    angle=float(xfrm.get('rot','0'))/60000*np.pi/180;c,s=np.cos(angle),np.sin(angle)
    fx=-1 if xfrm.get('flipH')=='1' else 1;fy=-1 if xfrm.get('flipV')=='1' else 1
    cx,cy=x+w/2,y+h/2
    return np.array([[1,0,cx],[0,1,cy],[0,0,1.]])@np.array([[c*fx,-s*fy,0],[s*fx,c*fy,0],[0,0,1.]])@np.array([[1,0,-cx],[0,1,-cy],[0,0,1.]])


def inventory(pptx,slide_number=1,*,presentation=None):
    deck=presentation if presentation is not None else Presentation(pptx);slide=deck.slides[slide_number-1];rows=[];seen=set()
    def visit(tree,matrix,parent):
        from .native_markup import active_children,variants
        for e in active_children(tree):
            tag=_tag(e)
            if tag not in SHAPES:continue
            sid=_id(e)
            if sid in seen:raise ValueError('Duplicate saved shape id')
            seen.add(sid);xf=_transform(e)
            if xf is None:raise ValueError('Shape lacks explicit transform; resolve placeholders before extraction')
            off=xf.find('a:off',NS);ext=xf.find('a:ext',NS)
            x,y=float(off.get('x')),float(off.get('y'));w,h=float(ext.get('cx')),float(ext.get('cy'))
            corners=np.array([[x,y,1],[x+w,y,1],[x+w,y+h,1],[x,y+h,1]]).T
            pts=matrix@_rotation_flip(xf,x,y,w,h)@corners
            bbox=[float(pts[0].min()),float(pts[1].min()),float(pts[0].max()-pts[0].min()),float(pts[1].max()-pts[1].min())]
            props=e.find('.//p:cNvPr',NS)
            rows.append({'shape_id':sid,'parent_id':parent,'kind':tag,'name':props.get('name',''),
                         'bbox_emu':bbox,'text':'\n'.join(e.xpath('.//a:t/text()')) if tag!='grpSp' else '',
                         'xml_hash':identity(e.xml)})
            if variants(e):rows[-1]['compatibility']=variants(e)
            if tag=='pic':
                has_svg=bool(e.xpath('.//*[local-name()="svgBlip"]'))
                has_master=bool(e.xpath('.//*[local-name()="vectorMaster" and namespace-uri()="https://lotra.ai/ns/vector-master/1"]'))
                rows[-1]['representation']='svg' if has_svg else 'raster_preview_with_svg_master' if has_master else 'raster'
                rows[-1]['paper_vector_master']=has_svg or has_master
            if tag=='sp':
                if any(n.tag=='{http://schemas.microsoft.com/office/drawing/2010/main}m' for n in e.iter()):
                    rows[-1]['equation_editing']='office_math'
                elif props.get('descr')=='lotraAI native equation: editable formatted text runs':
                    rows[-1]['equation_editing']='native_formatted_text'
            if tag=='grpSp':visit(e,matrix@_group_matrix(xf),sid)
    visit(slide.shapes._spTree,np.eye(3),None)
    canvas=[deck.slide_width,deck.slide_height]
    if presentation is None:
        del slide,deck
        gc.collect()
    return {'pptx_hash':digest(pptx),'slide':slide_number,'canvas_emu':canvas,'objects':rows}


def export_owners(pptx,manifest,output,*,inventories=None):
    if manifest.get('version')!='semantic-owners-1' or manifest.get('pptx_hash')!=digest(pptx):
        raise ValueError('Owner manifest must bind this saved PPT')
    owners=manifest.get('components',[])
    if not owners or len({o['id'] for o in owners})!=len(owners):raise ValueError('Empty/duplicate component owners')
    output=Path(output)
    if output.exists():raise FileExistsError(output)
    inventories={} if inventories is None else inventories
    with ZipFile(pptx) as package:
        native_xml_bytes=sum(i.file_size for i in package.infolist()
                             if i.filename.startswith('ppt/slides/slide') and i.filename.endswith('.xml'))
    if native_xml_bytes>48*1024*1024:
        # Dense native colour paths must not be copied into a second enormous
        # in-memory deck. The actual saved named groups are already reusable.
        # Keep their authoritative PPTX and explicit per-owner pointers.
        entries=[]
        for owner in owners:
            num=owner.get('slide',1)
            if num not in inventories:inventories[num]=inventory(pptx,num)
            rows={r['shape_id']:r for r in inventories[num]['objects']}
            selected=set(owner.get('shape_ids',[]))
            if not selected or not selected<=rows.keys():raise ValueError('Owner contains absent/empty native parts')
            while True:
                expanded=selected|{r['shape_id'] for r in rows.values() if r['parent_id'] in selected}
                if expanded==selected:break
                selected=expanded
            leaves=[rows[k] for k in selected if rows[k]['kind']!='grpSp']
            if not leaves:raise ValueError('Owner has no actual component artwork')
            rasters=sum(r.get('representation')=='raster' for r in leaves)
            entries.append({'semantic_id':owner['id'],'slide':num,'source_shape_ids':sorted(selected),
                'root_names':[rows[k]['name'] for k in owner['shape_ids']],
                'representation':'mixed' if rasters else 'vector','raster_count':rasters,
                'complete_status':owner.get('complete_status','UNVERIFIED'),
                'source_artwork_hash':identity([r['xml_hash'] for r in leaves]),'native_behavior':'UNVERIFIED'})
        with writer(output):
            out=output/'Components_editable.pptx';shutil.copyfile(pptx,out)
            index={'source_pptx_hash':digest(pptx),'native_pptx_hash':digest(out),'owner_manifest_hash':identity(manifest),
                'components':entries,'status':'AUTHORITATIVE_GROUPED_NATIVE_SNAPSHOT',
                'reuse_mode':'Use the exact named native groups on their original slide; other semantic owners remain independent',
                'reason':'Dense native paths retained without materializing a duplicate full drawing tree'}
            write_json(output/'index.json',index)
        return index
    source=Presentation(pptx);deck=Presentation();deck.slide_width=source.slide_width;deck.slide_height=source.slide_height
    entries=[]
    with writer(output):
        for owner in owners:
            num=owner.get('slide',1);original=source.slides[num-1]
            if num not in inventories:inventories[num]=inventory(pptx,num,presentation=source)
            inv=inventories[num];rows={r['shape_id']:r for r in inv['objects']}
            selected=set(owner.get('shape_ids',[]))
            if not selected or not selected<=rows.keys():raise ValueError('Owner contains absent/empty native parts')
            # Selecting a group means all its parts; selection of fragments keeps
            # their ancestor transforms, not unrelated siblings.
            while True:
                expanded=selected|{r['shape_id'] for r in rows.values() if r['parent_id'] in selected}
                if expanded==selected:break
                selected=expanded
            leaves=[r for k,r in rows.items() if k in selected and r['kind']!='grpSp']
            if not leaves:raise ValueError('Owner has no actual component artwork')
            def prune(tree):
                from .native_markup import prune_owners,active_children
                clone=copy.deepcopy(tree);prune_owners(clone,selected)
                return [e for e in clone if _tag(e) in SHAPES or e.tag=='{'+NS['mc']+'}AlternateContent']
            forest=prune(original.shapes._spTree);slide=deck.slides.add_slide(deck.slide_layouts[6])
            group=slide.shapes.add_group_shape();group.name=owner['id'];gid=max(rows)+1
            group._element.find('p:nvGrpSpPr/p:cNvPr',NS).set('id',str(gid))
            rasters=0;detached=[]
            for element in forest:
                from .resource_transfer import transfer
                transfer(element,original.part,slide.part)
                rasters+=sum(_tag(pic)=='pic' and not pic.xpath('.//*[local-name()="svgBlip"]') for pic in element.iter())
                for connection in element.iter('{'+NS['a']+'}stCxn','{'+NS['a']+'}endCxn'):
                    if int(connection.get('id')) not in selected:
                        detached.append({'shape_id':_id(element),'external_id':int(connection.get('id'))})
                        connection.getparent().remove(connection)
                group._element.append(element)
            b=[r['bbox_emu'] for r in leaves];x=min(v[0] for v in b);y=min(v[1] for v in b)
            w=max(v[0]+v[2] for v in b)-x;h=max(v[1]+v[3] for v in b)-y
            xf=group._element.find('p:grpSpPr/a:xfrm',NS)
            for tag,attrs in [('off',{'x':int(Inches(.3)),'y':int(Inches(.3))}),('ext',{'cx':max(1,round(w)),'cy':max(1,round(h))}),
                              ('chOff',{'x':round(x),'y':round(y)}),('chExt',{'cx':max(1,round(w)),'cy':max(1,round(h))})]:
                item=xf.find('a:'+tag,NS)
                if item is None:item=child(xf,'a:'+tag)
                for key,value in attrs.items():item.set(key,str(value))
            entries.append({'semantic_id':owner['id'],'slide':len(deck.slides),'source_shape_ids':sorted(selected),
                            'representation':'mixed' if rasters else 'vector','raster_count':rasters,
                            'detached_external_attachments':detached,'complete_status':owner.get('complete_status','UNVERIFIED'),
                            'source_artwork_hash':identity([r['xml_hash'] for r in leaves]),'native_behavior':'UNVERIFIED'})
        out=output/'Components_editable.pptx';deck.save(out)
        index={'source_pptx_hash':digest(pptx),'native_pptx_hash':digest(out),'owner_manifest_hash':identity(manifest),
               'components':entries,'status':'EXTRACTED_REQUIRES_RENDER_AND_EDIT'}
        write_json(output/'index.json',index)
    return index
