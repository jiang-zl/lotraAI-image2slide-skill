"""Owned LibreOffice save/reopen checks; canonical PPTX is never roundtripped."""
import os
import json
from pathlib import Path
import subprocess
import uuid
from .common import digest,write_json


def save_reopen(pptx,output,executable,timeout=120,renderer=None):
    from pptx import Presentation
    from .render import render
    from .office_compatibility import validate
    pptx=Path(pptx).resolve(strict=True);output=Path(output).resolve()
    output.mkdir(parents=True,exist_ok=False)
    before=digest(pptx);saved=output/'saved'
    if renderer is not None and renderer.get('project'):
        from .editor_backend import profile_directory
        profile=profile_directory(renderer,output);temp=profile/'temp'
    else:profile=output/'profile';profile.mkdir();temp=output/'temp'
    for path in (temp,saved):path.mkdir()
    env=dict(os.environ)
    env.update({k:str(temp) for k in ('TEMP','TMP','TMPDIR','XDG_CACHE_HOME','XDG_CONFIG_HOME','XDG_RUNTIME_DIR')})
    env.update(SAL_USE_VCLPLUGIN='svp',GSETTINGS_BACKEND='memory')
    command=[str(executable),'-env:UserInstallation='+profile.as_uri(),'--headless','--norestore',
        '--convert-to','pptx:Impress MS PowerPoint 2007 XML','--outdir',str(saved),str(pptx)]
    try:
        with (output/'Save.log').open('wb') as log:
            process=subprocess.Popen(command,stdout=log,stderr=log,env=env,
                creationflags=subprocess.CREATE_NO_WINDOW if os.name=='nt' else 0,start_new_session=os.name!='nt')
            try:process.wait(timeout=timeout)
            finally:
                from .editor_backend import _stop
                _stop(process)
        target=saved/pptx.name
        if process.returncode or not target.is_file() or not target.stat().st_size:
            raise RuntimeError('LibreOffice PPTX save failed: '+(output/'Save.log').read_text(encoding='utf-8',errors='replace')[-2000:])
        validate(target)
        if len(Presentation(target).slides)!=len(Presentation(pptx).slides):raise ValueError('Saved slide count changed')
        # Open the application-saved PPTX again through the actual application.
        if renderer is not None:
            # Use the same fixed client and request-bound completion protocol
            # as the declared project renderer. A --convert-to wrapper can
            # remain alive after writing its PDF; file presence alone is not
            # an operation acknowledgement and cannot turn a timeout into PASS.
            from .portable_editor import render as editor_render
            from .transient_export import register
            receipt=editor_render(target,output/'reopen-render',renderer,timeout=timeout,native_math_import=True)
            register(output/'reopen-render',target,receipt)
        else:
            receipt=render(target,output/'reopen-render',executable=executable,timeout=timeout)
        if digest(pptx)!=before:raise ValueError('LibreOffice modified canonical input')
        result={'status':'PASS','backend':'LibreOffice','input_sha256':before,'saved_sha256':digest(target),
            'pdf_sha256':receipt['pdf_hash'],'scope':'LibreOffice save and actual reopen export of owned copy; PowerPoint/WPS editing UNVERIFIED',
            'canonical_input_unchanged':True,'target_powerpoint':'UNVERIFIED','temporary_directory':str(temp),
            'reopen_transport':'REQUEST_BOUND_EDITOR' if renderer is not None else 'BOUNDED_DECLARED_CLI'}
        if renderer is not None:
            proof=output/'reopen-render/math-editor.json'
            observations=json.loads(proof.read_bytes()).get('observations',[]) if proof.is_file() else []
            result['saved_math_style']={'status':'NOT_QUALIFIED' if any(r.get('source_style_retention')=='NOT_QUALIFIED' for r in observations) else 'NOT_COMPARED',
                'native_import_style_zones':sum(r.get('style_mode')=='application_saved_native_import' for r in observations),
                'scope':'Actual saved OMML reopened by native Math; source style retention is a separate obligation, not a PASS'}
            result['saved_math_geometry']={'status':'NOT_QUALIFIED' if observations else 'NOT_COMPARED',
                'scope':'Private renderer uses the imported Math intrinsic area; this is not raw application geometry or source-size retention'}
            plan=output/'reopen-render/editor-plan.json'
            metadata=json.loads(plan.read_bytes()).get('native_saved_math',{}) if plan.is_file() else {}
            result['native_math_compatibility']={'resolved_branches':metadata.get('compatibility_branches',[]),
                'expected_math_objects':metadata.get('expected_math_objects',0),
                'raw_application_branch_selection':'NOT_QUALIFIED_BY_ADAPTER_RENDER',
                'scope':'Source OMML primary branch resolved only for the declared native renderer; canonical PPTX and raw saved copy unchanged'}
        write_json(output/'Save-reopen.json',result);return result
    except Exception as error:
        write_json(output/'Backend-failure.json',{'status':'FAILED','backend':'LibreOffice','stage':'save-reopen',
            'input_sha256':before,'reason':str(error),'canonical_input_unchanged':digest(pptx)==before})
        raise
