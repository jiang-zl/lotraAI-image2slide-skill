"""Lossless, contract-bound local edits to an already-permitted picture layer.

The external editor may propose arbitrary pixels. Only an advance-frozen edit
mask reaches the new working asset; every RGBA channel outside stays identical.
This is not an exception to the native-icon or scientific-content requirements.
"""

from pathlib import Path

import numpy as np
from PIL import Image
from scipy import ndimage as ndi

from .common import digest, identity, safe_path, write_json

RENDER_SUPPORT_MARGIN_PX = 8


def _render_support(allowed, preview_size):
    with Image.fromarray(allowed) as image:
        resized = np.asarray(image.resize(tuple(preview_size), Image.Resampling.NEAREST)) > 0
    return (ndi.maximum_filter(resized.astype(np.uint8),
                               size=2 * RENDER_SUPPORT_MARGIN_PX + 1) * 255).astype(np.uint8)


def freeze_full_canvas_render_mask(spec, base, pptx, shape_id, preview_size, output):
    """Predeclare display influence for one unrotated full-canvas picture.

    Other picture transforms require their own measured transform protocol;
    this function refuses to guess a safe mask for them.
    """
    import hashlib
    from pptx import Presentation
    from .saved_components import inventory
    from .mathml import NS
    _, pixels, allowed, _ = validate_contract(spec, base)
    deck = Presentation(pptx)
    if spec["baseline_pptx_hash"] != digest(pptx) or len(deck.slides) != 1:
        raise ValueError("Full-canvas render mask needs one frozen saved slide")
    preview_size = list(preview_size)
    if (len(preview_size) != 2 or preview_size[0] < 16 or preview_size[1] < 16
            or abs(preview_size[1] - round(preview_size[0] * deck.slide_height / deck.slide_width)) > 1):
        raise ValueError("Declared preview canvas does not match the saved slide")
    rows = [row for row in inventory(pptx)["objects"] if row["shape_id"] == shape_id]
    if len(rows) != 1 or rows[0]["kind"] != "pic" or rows[0]["name"] != spec["layer_id"]:
        raise ValueError("Frozen picture identity is absent or ambiguous")
    def pictures(shapes):
        for shape in shapes:
            if shape.shape_type == 13:
                yield shape
            elif shape.shape_type == 6:
                yield from pictures(shape.shapes)
    matches = [shape for shape in pictures(deck.slides[0].shapes) if shape.shape_id == shape_id]
    if len(matches) != 1 or hashlib.sha256(matches[0].image.blob).hexdigest() != spec["source"]["sha256"]:
        raise ValueError("Frozen source pixels do not match the saved picture payload")
    picture = matches[0]
    if picture.rotation or any(getattr(picture, field) for field in
                               ("crop_left", "crop_right", "crop_top", "crop_bottom")):
        raise ValueError("Crop or rotation needs a separate render-mask transform protocol")
    ancestor = picture._element.getparent()
    while ancestor is not None:
        if ancestor.tag.rsplit("}", 1)[-1] == "grpSp":
            xf = ancestor.find("p:grpSpPr/a:xfrm", NS)
            if xf is None or xf.get("rot") or xf.get("flipH") or xf.get("flipV"):
                raise ValueError("Transformed picture parent needs a separate render-mask protocol")
            off, ext = xf.find("a:off", NS), xf.find("a:ext", NS)
            child_off, child_ext = xf.find("a:chOff", NS), xf.find("a:chExt", NS)
            if (off is None or ext is None or child_off is None or child_ext is None
                    or off.get("x") != child_off.get("x") or off.get("y") != child_off.get("y")
                    or ext.get("cx") != child_ext.get("cx") or ext.get("cy") != child_ext.get("cy")):
                raise ValueError("Scaled picture parent needs a separate render-mask protocol")
        ancestor = ancestor.getparent()
    x, y, width, height = rows[0]["bbox_emu"]
    if any(abs(a-b) > 2 for a, b in zip((x,y,width,height),
                                         (0,0,deck.slide_width,deck.slide_height))):
        raise ValueError("Only an exact full-canvas picture has this render-mask mapping")
    if pixels.shape[:2] != (preview_size[1], preview_size[0]):
        raise ValueError("Picture pixel canvas differs from exact preview mapping")
    output = Path(output)
    if output.exists():
        raise FileExistsError(output)
    Image.fromarray(_render_support(allowed, preview_size)).save(output, format="PNG")
    receipt = {"schema_version": "picture-render-mask-1", "status": "FROZEN_BEFORE_CANDIDATE",
               "base_contract_hash": identity(spec), "baseline_pptx_hash": digest(pptx),
               "layer_id": spec["layer_id"], "shape_id": shape_id,
               "preview_size": preview_size, "support_margin_px": RENDER_SUPPORT_MARGIN_PX,
               "asset_allowed_mask_hash": spec["allowed_mask"]["sha256"],
               "render_mask_hash": digest(output),
               "scope": "full-canvas fixed-preview resampling influence; actual render still must be checked"}
    write_json(output.with_suffix(".render-mask.json"), receipt)
    return receipt


def _resource(base, ref):
    if not isinstance(ref, dict) or not isinstance(ref.get("path"), str):
        raise ValueError("Raster edit resource path is required")
    path = safe_path(base, ref["path"])
    if digest(path) != ref.get("sha256"):
        raise ValueError("Raster edit resource hash changed")
    return path


def _png_rgba(path):
    with Image.open(path) as im:
        if im.format != "PNG":
            raise ValueError("Local raster working assets must be lossless PNG, not JPEG")
        if im.mode not in ("RGB", "RGBA"):
            raise ValueError("Local raster working asset needs explicit RGB/RGBA pixels")
        return np.asarray(im.convert("RGBA")).copy()


def _mask(path, shape, *, binary):
    with Image.open(path) as im:
        if im.format != "PNG" or im.mode != "L" or im.size != (shape[1], shape[0]):
            raise ValueError("Raster edit mask must be an aligned grayscale PNG")
        mask = np.asarray(im).copy()
    if binary and not np.isin(mask, [0, 255]).all():
        raise ValueError("Allowed edit mask must be binary and frozen in advance")
    return mask


def validate_contract(spec, base):
    if spec.get("schema_version") != "local-raster-edit-1":
        raise ValueError("Unknown local raster edit contract")
    if spec.get("layer_kind") not in ("illustrative_photo", "environment", "permitted_raster_layer"):
        raise ValueError("Local raster edit is restricted to a previously permitted picture layer")
    if not spec.get("instance_id") or not spec.get("layer_id") or not spec.get("intent"):
        raise ValueError("Instance, layer and edit intent are required")
    if not spec.get("baseline_pptx_hash") or not spec.get("authorization_hash"):
        raise ValueError("Frozen PPT baseline and separate authorization are required")
    source = _resource(base, spec["source"])
    pixels = _png_rgba(source)
    allowed = _mask(_resource(base, spec["allowed_mask"]), pixels.shape, binary=True)
    if not allowed.any():
        raise ValueError("Empty edit mask cannot authorize an image change")
    if spec.get("blend_alpha"):
        blend = _mask(_resource(base, spec["blend_alpha"]), pixels.shape, binary=False)
        if np.any(blend[~(allowed > 0)]):
            raise ValueError("Blend region cannot expand the authorized edit mask")
    else:
        blend = allowed
    if spec.get("render_plan"):
        plan = spec["render_plan"]
        mask_file = _resource(base, plan["mask"])
        receipt_file = _resource(base, plan["receipt"])
        from .common import read_json
        proof = read_json(receipt_file)
        unsigned = {key: value for key, value in spec.items() if key != "render_plan"}
        if (proof.get("schema_version") != "picture-render-mask-1"
                or proof.get("base_contract_hash") != identity(unsigned)
                or proof.get("baseline_pptx_hash") != spec["baseline_pptx_hash"]
                or proof.get("layer_id") != spec["layer_id"]
                or proof.get("support_margin_px") != RENDER_SUPPORT_MARGIN_PX
                or proof.get("render_mask_hash") != digest(mask_file)):
            raise ValueError("Frozen display-influence mask is not bound to this picture edit")
        observed = _mask(mask_file, (proof["preview_size"][1], proof["preview_size"][0], 4), binary=True)
        if not np.array_equal(observed, _render_support(allowed, proof["preview_size"])):
            raise ValueError("Render protection mask was widened or replaced")
    return source, pixels, allowed, blend


def compose(spec, base, candidate_ref, output):
    """Save a fresh PNG and a receipt; never overwrite the approved work asset."""
    source, original, allowed, blend = validate_contract(spec, base)
    if candidate_ref.get("contract_hash") != identity(spec):
        raise ValueError("Candidate was not generated for this frozen edit mask/contract")
    candidate = _png_rgba(_resource(base, candidate_ref))
    if candidate.shape != original.shape:
        raise ValueError("Candidate changed the working canvas")
    weight = blend.astype(np.uint16)[..., None]
    mixed = ((original.astype(np.uint16) * (255 - weight)
              + candidate.astype(np.uint16) * weight + 127) // 255).astype(np.uint8)
    mixed[allowed == 0] = original[allowed == 0]
    output = Path(output)
    if output.exists():
        raise FileExistsError(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    Image.fromarray(mixed).save(output, format="PNG")
    report = verify(spec, base, candidate_ref, output)
    receipt = {"schema_version": "local-raster-result-1", "status": "PIXEL_PROTECTION_PASS_PPT_CONTEXT_UNVERIFIED",
               "contract_hash": identity(spec), "source_hash": digest(source),
               "candidate_hash": candidate_ref["sha256"], "output_hash": digest(output),
               "baseline_pptx_hash": spec["baseline_pptx_hash"],
               "instance_id": spec["instance_id"], "layer_id": spec["layer_id"],
               "authorization_hash": spec["authorization_hash"],
               "allowed_pixels": int((allowed > 0).sum()),
               "changed_pixels": report["changed_pixels"],
               "outside_changed_pixels": report["outside_changed_pixels"],
               "scope": "lossless working asset only; saved PPT layer and parent rendering need separate checks"}
    write_json(output.with_suffix(".local-edit.json"), receipt)
    return receipt


def verify(spec, base, candidate_ref, output):
    """Recompute expected composite; reject altered masks, output and outer pixels."""
    _, original, allowed, blend = validate_contract(spec, base)
    if candidate_ref.get("contract_hash") != identity(spec):
        raise ValueError("Candidate/contract identity differs")
    candidate = _png_rgba(_resource(base, candidate_ref))
    result = _png_rgba(output)
    if candidate.shape != original.shape or result.shape != original.shape:
        raise ValueError("Local raster canvas mismatch")
    weight = blend.astype(np.uint16)[..., None]
    expected = ((original.astype(np.uint16) * (255 - weight)
                 + candidate.astype(np.uint16) * weight + 127) // 255).astype(np.uint8)
    expected[allowed == 0] = original[allowed == 0]
    if not np.array_equal(expected, result):
        raise ValueError("Local raster result differs from frozen-mask deterministic composite")
    changed = np.any(result != original, axis=2)
    outside = changed & (allowed == 0)
    if outside.any():
        raise ValueError("Protected RGBA pixels changed outside edit mask")
    return {"status": "PASS_PIXEL_SCOPE_ONLY", "changed_pixels": int(changed.sum()),
            "outside_changed_pixels": 0, "output_hash": digest(output),
            "contract_hash": identity(spec)}
