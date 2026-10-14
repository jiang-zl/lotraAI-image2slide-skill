"""Audit actual Office component exports and saved reuse operations.

This checks the component deck that was actually tested, not every object in its
source figure. Native equation and full-context claims require separate probes.
"""
from pathlib import Path
from io import BytesIO
from lxml import etree as E
from PIL import Image
import numpy as np
from pptx import Presentation
from zipfile import ZipFile
from .inspect import inventory as xml_inventory
from .common import digest,read_json,safe_path,write_json


def _walk(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type==6:yield from _walk(shape.shapes)


def verify(input_deck,office_directory,output):
    folder=Path(office_directory);raw=read_json(folder/'probe.json');saved=folder/'Native_reuse_test.pptx'
    if raw.get('input_hash')!=digest(input_deck) or raw.get('output_hash')!=digest(saved):
        raise ValueError('Component operation evidence does not bind actual PPT bytes')
    original=Presentation(input_deck);reopened=Presentation(saved);checks=[]
    if len(original.slides)!=len(raw.get('components',[])) or len(reopened.slides)!=len(original.slides):
        raise ValueError('Component coverage does not match tested slides')
    for record in raw['components']:
        i=record['slide']-1;slide=reopened.slides[i];source=original.slides[i]
        names={s.name:s for s in _walk(slide.shapes)}
        with ZipFile(saved) as package:
            xml_names,_=xml_inventory(E.fromstring(package.read('ppt/slides/slide'+str(i+1)+'.xml')))
        own=[]
        def add(name,ok,detail=None):own.append({'check':name,'status':'PASS' if ok else 'FAIL','detail':detail})
        add('office_saved_reopen_operations',record['status']=='PASS_LISTED_OPERATIONS' and not record.get('reason'),record.get('operations'))
        vector=not any(s.shape_type==13 for s in _walk(source.shapes))
        for kind in ('preview','svg'):
            resource=record.get(kind,{})
            if not resource.get('path'):add(kind+'_export',False,resource);continue
            file=safe_path(folder,resource['path']);add(kind+'_payload',digest(file)==resource.get('sha256'))
            if kind=='svg':
                svg=E.parse(str(file),E.XMLParser(resolve_entities=False,no_network=True))
                image_nodes=svg.xpath('//*[local-name()="image"]')
                external=[n.get('{http://www.w3.org/1999/xlink}href') or n.get('href','') for n in image_nodes]
                external=[u for u in external if u and not u.startswith('data:')]
                add('svg_self_contained',not external,{'external_images':len(external)})
                add('svg_declared_representation',not image_nodes if vector else bool(image_nodes),
                    {'declared':'native_vector' if vector else 'contains_declared_raster','embedded_images':len(image_nodes),
                     'native_geometry_nodes':len(svg.xpath('//*[local-name()="path" or local-name()="rect" or local-name()="ellipse" or local-name()="polygon"]'))})
        for op in record.get('operations',[]):
            kind=op['kind'];shape=names.get(op.get('name'))
            if kind=='ungroup':add('independent_saved_parts',all(n in xml_names for n in op['part_names']))
            elif kind in ('copy_move_resize','import_replacement_picture'):
                actual=[float(v)/12700 for v in (shape.left,shape.top,shape.width,shape.height)] if shape is not None else []
                add('saved_'+kind,len(actual)==4 and max(abs(a-b) for a,b in zip(actual,op['expected_saved']))<.025,
                    {'actual_pt':actual,'expected_pt':op['expected_saved']})
                if kind=='import_replacement_picture' and shape is not None:
                    asset=safe_path(folder,op['source_asset']);same=False
                    if shape.shape_type==13 and digest(asset)==op['source_sha256']:
                        with Image.open(asset) as a,Image.open(BytesIO(shape.image.blob)) as b:
                            same=a.size==b.size and np.array_equal(np.asarray(a.convert('RGBA')),np.asarray(b.convert('RGBA')))
                    add('replacement_pixels_in_saved_PPT',same)
            elif kind=='edit_text':add('saved_text_content',shape is not None and shape.text==op['expected_saved'])
            elif kind in ('edit_fill','edit_line'):
                rgb=None
                try:rgb=str((shape.line.color if kind=='edit_line' else shape.fill.fore_color).rgb)
                except (AttributeError,TypeError):pass
                n=op['expected_saved'];wanted=f'{n&255:02X}{(n>>8)&255:02X}{(n>>16)&255:02X}'
                add('saved_'+kind+'_color',rgb==wanted,{'actual':rgb,'expected':wanted})
            elif kind=='crop_picture':
                add('saved_crop_present',shape is not None and shape.shape_type==13 and shape.crop_left>0)
            elif kind=='edit_path_node':
                # Exact node coordinates are checked by the actual reopened
                # COM operation; XML independently proves a native freeform.
                add('saved_freeform_geometry',shape is not None and bool(shape._element.xpath('./p:spPr/a:custGeom')))
            else:add('known_operation',False,kind)
        checks.append({'slide':i+1,'owner':record['owner'],'checks':own,
                       'status':'PASS_LISTED_OPERATIONS' if own and all(c['status']=='PASS' for c in own) else 'FAIL'})
    result={'version':'component-roundtrip-1','source_pptx_hash':digest(input_deck),'saved_pptx_hash':digest(saved),
            'office_receipt_hash':digest(folder/'probe.json'),'checker_hash':digest(Path(__file__)),
            'components':checks,'status':'PASS_LISTED_OPERATIONS' if all(c['status']=='PASS_LISTED_OPERATIONS' for c in checks) else 'FAIL',
            'scope':'Listed component edits and actual SVG/PNG exports only; raster internals, full-page context, equation editor and hidden completion truth are separate'}
    write_json(output,result);return result
