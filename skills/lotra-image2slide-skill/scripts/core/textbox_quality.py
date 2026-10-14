"""Per-owner native text/math evidence, actual glyph metrics and scoped retention.

Font names/sizes cannot be recovered as metadata from a raster. Source-only
readings and visual domains precede candidates; program scores never certify
the accuracy of the host reading itself.
"""
import copy
import math
import unicodedata
from pathlib import Path
import numpy as np
from lxml import etree as E
from PIL import Image
from scipy import ndimage as ndi
from .common import digest, identity, read_json, write_json, safe_path
from .mathml import NS, canonical_ast
from .inspect import xml, inventory, parse_math_container

VERSION = "textbox-quality-1.6.0"
TARGETS = {"foreground_f1": .90, "edge_f1_1px": .95,
           "shape_f1": .90, "weight_agreement": .92, "size_agreement": .96,
           "position_error_px": 1.5, "line_anchor_error_px": 1.5,
           "line_count_match": 1., "content_match": 1., "math_match": 1., "format_match":1.}
LOWER = {"position_error_px", "line_anchor_error_px"}
# Fixed sampling allowances; they cannot be enlarged by a repair proposal.
ALLOWANCE = {"foreground_f1": .012, "edge_f1_1px": .012,
             "shape_f1": .015, "weight_agreement": .035, "size_agreement": .025,
             "position_error_px": .75, "line_anchor_error_px": .75,
             "line_count_match": 0., "content_match": 0., "math_match": 0., "format_match":0.}


def literal(value):
    # Never NFKC-fold mathematical symbols, fullwidth punctuation or scripts.
    return unicodedata.normalize("NFC", value.replace("\r\n", "\n").replace("\r", "\n"))


def native_rows(pptx):
    nodes, duplicates = inventory(xml(pptx))
    rows = []
    for name, item in nodes.items():
        if name in duplicates:
            raise ValueError("Ambiguous text owner: " + name)
        node = item["node"]
        if E.QName(node).localname=='grpSp':continue
        bodies = [(name, node.find("p:txBody", NS))]
        table = node.find(".//a:tbl", NS)
        if table is not None:
            bodies = [(name + "@cell:" + str(r) + ":" + str(c), cell.find("a:txBody", NS))
                      for r, row in enumerate(table.findall("a:tr", NS))
                      for c, cell in enumerate(row.findall("a:tc", NS))]
        for key, body in bodies:
            if body is None:
                continue
            text = []
            for p in body.findall("a:p", NS):
                text.append("".join(n.text or "" if n.tag == "{" + NS["a"] + "}t" else "\n"
                                    for n in p.iter() if n.tag in (
                                        "{" + NS["a"] + "}t", "{" + NS["a"] + "}br")))
            asts = [parse_math_container(n) for n in body.findall(".//m:oMath", NS)]
            if not any(text) and not asts:
                continue
            styles = []
            for n in body.xpath(".//a:rPr | .//a:defRPr", namespaces=NS):
                styles.append({"font": [x.get("typeface") for x in n if x.tag in (
                    "{" + NS["a"] + "}latin", "{" + NS["a"] + "}ea")],
                    "size_pt": float(n.get("sz")) / 100 if n.get("sz") else None,
                    "bold": n.get("b"), "italic": n.get("i"), "spacing_100pt": n.get("spc")})
            rows.append({"id": key, "object": name, "text": literal("\n".join(text)),
                         "math_spans": asts, "native_styles": styles,
                         "paragraph_count":len(body.findall("a:p",NS)),
                         "alignment": [p.get("algn") for p in body.findall("a:p/a:pPr", NS)]})
    return rows


def contracts(state):
    from .external_sync import native_contract
    return [native_contract(state, r) for r in state.get("textbox_quality", {}).get("contracts", {}).values()]


def math_roles(spans):
    roles=[]
    def walk(node,path):
        if isinstance(node,dict):
            if node.get("op")=="accent":roles.append(path+".mark")
            for key,value in node.items():
                if key in ("sub","sup","accent","numerator","denominator","degree"):
                    roles.append(path+"."+key)
                if isinstance(value,(dict,list)):walk(value,path+"."+key)
        elif isinstance(node,list):
            for i,value in enumerate(node):walk(value,path+"."+str(i))
    for i,ast in enumerate(spans):
        path="math_span_"+str(i);roles.append(path);walk(ast,path)
    return roles


def discover(task, state, candidate=None):
    candidate = candidate or state.get("active_candidate") or state.get("selected")
    if not candidate:
        return {"status": "AWAITING_SAVED_NATIVE_OWNERS", "items": []}
    rows = native_rows(safe_path(task, candidate["pptx"]["path"]))
    bound = {r["id"] for r in contracts(state)}
    return {"schema": VERSION, "source_hash": state["source"]["sha256"],
            "baseline_hash": candidate["pptx"]["sha256"],
            "items": [{"id": r["id"], "object": r["object"],
                       "kind": "mixed_math" if r["math_spans"] else "text",
                       "status": "bound" if r["id"] in bound else "requires_source_only_rereading"}
                      for r in rows], "limits": "Inventory is not semantic or visual acceptance."}


def _box(box, size):
    if (not isinstance(box, list) or len(box) != 4 or
            any(type(v) not in (int, float) or not math.isfinite(v) for v in box) or
            box[0] < 0 or box[1] < 0 or box[2] <= 0 or box[3] <= 0 or
            box[0] + box[2] > size[0] + .01 or box[1] + box[3] > size[1] + .01):
        raise ValueError("Freeze a nonempty original-canvas glyph domain")
    return box


def _bind(task, state, request):
    if request.get("source_hash") != state["source"]["sha256"]:
        raise ValueError("Bind the immutable original, not a candidate crop")
    if digest(safe_path(task,state["source"]["path"])) != state["source"]["sha256"]:
        raise ValueError("Original image changed; preserve its prior source identity")
    candidate = state.get("active_candidate") or state.get("selected")
    if not candidate or request.get("baseline_hash") != candidate["pptx"]["sha256"]:
        raise ValueError("Source contract must bind the current native owner inventory")
    rows = {r["id"]: r for r in native_rows(safe_path(task, candidate["pptx"]["path"]))}
    entries = request.get("entries", [])
    if not 1 <= len(entries) <= 8 or len({r["id"] for r in entries}) != len(entries):
        raise ValueError("Bind 1..8 distinct owners per source-only packet")
    registry = state.setdefault("textbox_quality", {"contracts": {}, "history": []})
    for entry in entries:
        row = copy.deepcopy(entry)
        if row["id"] not in rows or row.get("object") != rows[row["id"]]["object"]:
            raise ValueError("Bind one exact saved text owner/cell")
        _box(row["roi"], state["source_size"])
        for box in row.get("word_rois", []):
            _box(box, state["source_size"])
            x, y, w, h = row["roi"]
            if box[0] < x or box[1] < y or box[0] + box[2] > x+w or box[1] + box[3] > y+h:
                raise ValueError("Every source word domain must lie inside its frozen ROI")
        if not row.get("word_rois") or not row.get("ownership_evidence"):
            raise ValueError("Exclude arrows/borders/upper layers; whole rectangles are not glyph ownership")
        readings = row.get("readings", [])
        if (len(readings) != 2 or readings[0].get("id") == readings[1].get("id") or
                any(r.get("mode") != "original_only" or not r.get("evidence") for r in readings)):
            raise ValueError("Two distinct original-only readings are required, not native text reused as an answer")
        for reading in readings:
            ref = reading["source_view"]
            file = safe_path(task, ref["path"])
            if digest(file) != ref["sha256"]:
                raise ValueError("Source reading crop changed")
            # Verify exact original pixels instead of accepting a labelled candidate.
            original = Image.open(safe_path(task, state["source"]["path"])).convert("RGB")
            x, y, w, h = map(round, row["roi"])
            crop = Image.open(file).convert("RGB")
            if crop.size != (w,h) or crop.tobytes() != original.crop((x,y,x+w,y+h)).tobytes():
                raise ValueError("Reading evidence must be an exact original crop at the declared ROI")
        if (literal(readings[0].get("text", "")) != literal(readings[1].get("text", "")) or
                identity([canonical_ast(a) for a in readings[0].get("math_spans", [])]) !=
                identity([canonical_ast(a) for a in readings[1].get("math_spans", [])])):
            raise ValueError("Source readings disagree; reread/correct the original guide before drawing")
        row["text"] = literal(readings[1].get("text", ""))
        row["math_spans"] = readings[1].get("math_spans", [])
        row["source_line_count"]=readings[1].get("source_line_count",len(row["text"].split("\n")) if row["text"] else 1)
        required=set(math_roles(row["math_spans"]))
        if not required <= {z.get("role") for z in row.get("zones",[])}:
            raise ValueError("Observe each formula span and small script/accent role independently: "+str(sorted(required)))
        row["source_hash"] = state["source"]["sha256"]
        row["schema"] = VERSION
        row["source_font_metadata"] = row.get("source_font_metadata", {
            "font": None, "point_size": None, "bold_flag": None, "status": "UNAVAILABLE_FROM_RASTER"})
        if any(row["source_font_metadata"].get(k) is not None for k in ("font","point_size","bold_flag")):
            raise ValueError("Raster pixels do not establish font metadata; keep observable style or user-declared targets separate")
        for zone in row.get("zones", []):
            _box(zone["roi"], state["source_size"])
            zx,zy,zw,zh=zone["roi"];x,y,w,h=row["roi"]
            if zx<x or zy<y or zx+zw>x+w or zy+zh>y+h:
                raise ValueError("A math role belongs inside its original textbox observation ROI")
            if not zone.get("role") or not zone.get("evidence"):
                raise ValueError("Math/glyph zones need original role and observed evidence")
            zx,zy,zw,zh=map(round,zone["roi"])
            if not _mask(original.crop((zx,zy,zx+zw,zy+zh)),row.get("selector",{})).any():
                raise ValueError("Original role has no visible glyph; correct its domain or leave it unresolved: "+zone["role"])
        row["targets"] = {**TARGETS, **row.get("targets", {})}
        if set(row["targets"]) != set(TARGETS) or any(
                not isinstance(v, (int,float)) or not math.isfinite(v) or v < 0
                for v in row["targets"].values()):
            raise ValueError("Use fixed named typography targets")
        if any((v>TARGETS[k] if k in LOWER else v<TARGETS[k]) for k,v in row["targets"].items()):
            raise ValueError("Do not weaken mandatory typography targets to rescue a drawing")
        if any(v>1 for k,v in row["targets"].items() if k not in LOWER):
            raise ValueError("Agreement targets must be within0..1")
        old = registry["contracts"].get(row["id"])
        if old:
            if identity(old) == identity(row):
                continue
            if not request.get("source_correction") or old["targets"] != row["targets"]:
                raise ValueError("Keep source history and targets; explain original-reading/domain correction")
            registry["history"].append({"prior": old, "reason": request["source_correction"]})
        registry["contracts"][row["id"]] = row
    return {"status": "SOURCE_TEXTBOX_CONTRACTS_BOUND", "count": len(registry["contracts"]),
            "limits": "Separate readings prevent circular native matching; a model can still misread both. Original glyph metrics and discrepancy review remain necessary."}


def bind(task, request):
    from . import agent_runtime as a
    with a.writer(Path(task).resolve()):
        task, state = a._state(task)
        result = _bind(task, state, request)
        a._save(task, state)
        write_json(task/"source/textbox-contracts.json", state["textbox_quality"])
        return result


def _mask(rgb, selector):
    a = np.asarray(rgb).astype(float)
    mode = selector.get("mode", "neutral_dark")
    if mode == "red":
        return (a[...,0] > 100) & (a[...,0]-a[...,1] > 65) & (a[...,0]-a[...,2] > 65)
    if mode != "neutral_dark":
        raise ValueError("Unsupported glyph selector; supply a measured owned mask for coloured backgrounds")
    return (a.mean(2) < selector.get("max_mean_rgb_255", 180)) & (
            np.ptp(a, axis=2) < selector.get("max_chroma_rgb_255", 80))


def _f1(a,b):
    return float(2*np.count_nonzero(a & b)/(a.sum()+b.sum())) if a.any() and b.any() else 0.


def _extent(mask):
    y,x = np.nonzero(mask)
    return [int(x.min()),int(y.min()),int(x.max()-x.min()+1),int(y.max()-y.min()+1)] if len(x) else None


def _lines(mask):
    rows = np.flatnonzero(mask.any(1))
    if not len(rows):
        return []
    groups = np.split(rows, np.flatnonzero(np.diff(rows)>2)+1)
    return [[int(g[0]),int(g[-1])] for g in groups if len(g)>0]


def glyph_metrics(source, actual, selector=None):
    """Fixed canvas first; shape-normalized F1 is a separate diagnostic."""
    a,b = _mask(source, selector or {}), _mask(actual, selector or {})
    aa,bb = _extent(a),_extent(b)
    if aa is None:
        return {"status": "required_unmeasured", "reason": "empty original glyph ownership; correct the source domain"}
    if bb is None:
        return {"status":"observed_absence","reason":"No actual glyph in the expected fixed source domain; may be missing or displaced",
                "values":{"foreground_f1":0.,"edge_f1_1px":0.,"shape_f1":0.,"weight_agreement":0.,"size_agreement":0.,
                    "position_error_px":None,"line_anchor_error_px":None,"line_count_match":0.},
                "details":{"source_ink_bbox":aa,"actual_ink_bbox":None,"line_pitch_error_px":None},
                "limits":"Zero local presence is observed, not a invented location or absence of the native AST. Locate the actual glyph before repair."}
    ae=a & ~ndi.binary_erosion(a);be=b & ~ndi.binary_erosion(b)
    da,db=ndi.distance_transform_edt(~ae),ndi.distance_transform_edt(~be)
    precision=float((da[be]<=1).mean());recall=float((db[ae]<=1).mean())
    edge=2*precision*recall/(precision+recall) if precision+recall else 0.
    al,bl=_lines(a),_lines(b)
    a_crop=a[aa[1]:aa[1]+aa[3],aa[0]:aa[0]+aa[2]]
    b_crop=b[bb[1]:bb[1]+bb[3],bb[0]:bb[0]+bb[2]]
    normalized=np.asarray(Image.fromarray(b_crop.astype("uint8")*255).resize(
        (aa[2],aa[3]),Image.Resampling.NEAREST))>127
    ratios=[bb[k]/aa[k] for k in (2,3)]
    density_a=a_crop.mean();density_b=b_crop.mean()
    weight=float(min(density_a,density_b)/max(density_a,density_b))
    anchors=float(max(abs(x[1]-y[1]) for x,y in zip(al,bl))) if len(al)==len(bl) else None
    return {"status":"measured", "values":{
        "foreground_f1": _f1(a,b), "edge_f1_1px": edge,
        "shape_f1":_f1(a_crop,normalized), "weight_agreement":weight,
        "size_agreement":float(min(min(r,1/r) for r in ratios)),
        "position_error_px":float(max(abs(aa[k]-bb[k]) for k in (0,1))),
        "line_anchor_error_px": anchors, "line_count_match":float(len(al)==len(bl))},
        "details":{"source_ink_bbox":aa,"actual_ink_bbox":bb,"source_line_ink_rows":al,
          "actual_line_ink_rows":bl,"density_source":float(density_a),"density_actual":float(density_b),
          "width_ratio":ratios[0],"height_ratio":ratios[1],
          "edge_distance_p95_px":float(max(np.percentile(da[be],95),np.percentile(db[ae],95))),
          "line_pitch_error_px":float(max(abs((al[i][0]-al[i-1][0])-(bl[i][0]-bl[i-1][0]))
                   for i in range(1,len(al)))) if len(al)==len(bl) and len(al)>1 else None},
        "limits":"Density is a weight proxy including antialiasing. Ink-row anchors are not recovered typographic baselines. Normalized shape score cannot compensate absolute size/position; no per-glyph warping."}


def semantic(row, native):
    expected=row.get("native_expected",{})
    formats=all(native.get(key)==value for key,value in expected.items()
                if key in ("alignment",))
    for key in ("bold","italic"):
        if key in expected:
            formats=formats and bool(native["native_styles"]) and all(
                r.get(key)==("1" if expected[key] else "0") for r in native["native_styles"])
    return {"content_match": float(literal(row["text"]) == native["text"]),
            "math_match": float(identity([canonical_ast(a) for a in row["math_spans"]]) ==
                                identity([canonical_ast(a) for a in native["math_spans"]])),
            "format_match":float(formats)}


def _affected(rows, allowed):
    return [r for r in rows if allowed is None or r["object"] in allowed]


def _closure(state,pptx,allowed):
    if not allowed:return allowed
    from .source_quality import changed_closure
    names=changed_closure(state,allowed)|set(allowed)
    nodes,_=inventory(xml(pptx))
    # Exact native parent movement also affects descendants, including
    # imported groups absent from a stale recipe/scene.
    pending=[name for name,fields in allowed.items() if set(fields)&{
        "group_translate","bbox.left","bbox.top","bbox.width","bbox.height","style.rotation","replace_component"}]
    while pending:
        parent=pending.pop()
        children=[name for name,item in nodes.items() if item.get("parent")==parent]
        for name in children:
            if name not in names:names.add(name)
        pending.extend(children)
    return {name:allowed.get(name,[]) for name in names}


def guard_saved(task, state, pptx, baseline=None, allowed=None):
    allowed=_closure(state,pptx,allowed)
    rows=contracts(state);native={r["id"]:r for r in native_rows(pptx)}
    required=state["objectives"].get("source_textbox_quality",False)
    from .native_core import guarded_text_ids
    scoped=guarded_text_ids(task,pptx,allowed)
    rows=[r for r in rows if r['id'] not in scoped]
    missing=set(native)-{r["id"] for r in rows}-scoped
    # First saved draw is allowed to establish actual IDs. Its adoption/export
    # is blocked until the original-only source packets bind every owner.
    if missing and required and baseline:
        changed={r["id"] for r in native.values() if allowed is None or r["object"] in allowed}
        if missing & changed:
            raise ValueError("Changed native text/math needs original-only textbox contracts: "+str(sorted(missing & changed)))
    failures=[]
    for row in _affected(rows,allowed):
        if row["id"] not in native or not all(semantic(row,native[row["id"]]).values()):
            failures.append(row["id"])
    if failures:
        raise ValueError("Exact source text/math blocks this native candidate: "+str(failures))
    return {"schema":VERSION,"bound_count":len(rows),"unbound":sorted(missing),
            "scope":[r["id"] for r in _affected(rows,allowed)],"semantic_failures":failures}


def report(task, state, candidate, ids=None, force=False):
    from .native_core import typography
    return typography(task,state,candidate,ids,force,_legacy_report)


def _legacy_report(task, state, candidate, ids=None, force=False):
    pptx=safe_path(task,candidate["pptx"]["path"])
    if digest(pptx)!=candidate["pptx"]["sha256"]:
        raise ValueError("Saved native typography candidate changed")
    rows=[r for r in contracts(state) if ids is None or r["id"] in ids]
    preview=safe_path(task,candidate["preview"]["path"])
    receipt=read_json(safe_path(task,candidate["render"]["path"]))
    if receipt.get("pptx_hash")!=digest(pptx) or digest(preview)!=candidate["preview"]["sha256"]:
        raise ValueError("Actual typography requires this saved native host render")
    ownership_file=safe_path(task,candidate["render"]["path"]).parent/"text-ownership.json"
    ownership=read_json(ownership_file) if ownership_file.exists() else None
    if ownership:
        isolated=ownership_file.parent/ownership["preview_path"]
        if ownership.get("pptx_hash")!=digest(pptx) or ownership.get("preview_hash")!=digest(isolated):
            raise ValueError("Actual isolated native glyph ownership changed")
    basis={"pptx":digest(pptx),"preview":digest(preview),"source":state["source"]["sha256"],
           "contracts":identity(rows),"method":VERSION,"host":receipt.get("backend")}
    basis["glyph_ownership"]=identity(ownership) if ownership else "legacy_bounds_only"
    original_file=safe_path(task,state["source"]["path"])
    if digest(original_file)!=state["source"]["sha256"]:
        raise ValueError("Original image identity changed")
    bounds=safe_path(task,candidate["render"]["path"]).parent/"text-bounds.json"
    if not bounds.is_file() or read_json(bounds).get("pptx_hash")!=digest(pptx):
        raise ValueError("Actual saved typography owner bounds missing/stale")
    basis.update(render_receipt=digest(safe_path(task,candidate["render"]["path"])),
                 owner_bounds=digest(bounds),fonts=receipt.get("font_inventory_hash"))
    inputs={ref["path"]:ref["sha256"] for row in rows for ref in
            ([r["source_view"] for r in row["readings"]]+([row["reference_mask"]] if row.get("reference_mask") else []))}
    if any(digest(safe_path(task,path))!=sha for path,sha in inputs.items()):
        raise ValueError("Original crop/mask evidence changed; do not reuse stale cached scores")
    directory=pptx.parent/"textbox-quality";directory.mkdir(exist_ok=True)
    cache_file=directory/(identity(basis)+".json")
    if cache_file.exists() and not force:
        return read_json(cache_file)
    original=Image.open(safe_path(task,state["source"]["path"])).convert("RGB")
    actual=Image.open(isolated if ownership else preview).convert("RGB").resize(tuple(state["source_size"]),Image.Resampling.LANCZOS)
    native={r["id"]:r for r in native_rows(pptx)}
    # Office owner bounds are used only to isolate glyph ownership, never to
    # move/register rendered pixels into a convenient source window.
    from .component_report import _native_ink
    bounds=safe_path(task,candidate["render"]["path"]).parent/"text-bounds.json"
    if not bounds.is_file() or read_json(bounds).get("pptx_hash")!=digest(pptx):
        raise ValueError("Actual saved typography owner bounds missing/stale")
    ink=_native_ink(pptx,bounds,state["source_size"])
    results=[]
    for row in rows:
        x,y,w,h=map(round,row["roi"])
        s=np.asarray(original.crop((x,y,x+w,y+h))).copy()
        b=np.asarray(actual.crop((x,y,x+w,y+h))).copy()
        owned=np.zeros((h,w),bool)
        for box in row["word_rois"]:
            wx,wy,ww,hh=map(round,box);owned[wy-y:wy-y+hh,wx-x:wx-x+ww]=True
        s[~owned]=255
        if row.get("reference_mask"):
            ref=row["reference_mask"];file=safe_path(task,ref["path"])
            if digest(file)!=ref["sha256"]:raise ValueError("Original glyph ownership mask changed")
            mask=np.asarray(Image.open(file).convert("L"))>127
            if mask.shape!=(h,w):raise ValueError("Original glyph mask needs this exact ROI canvas")
            s[~mask]=255
        owner=ink.get(row["id"]) or ink.get(row["object"])
        record={"id":row["id"],"object":row["object"],"source_font_metadata":row["source_font_metadata"],
                "targets":row["targets"],"native_styles":native.get(row["id"],{}).get("native_styles")}
        record["native_format_evidence"] = "EXPLICIT_SOURCE_FORMAT" if row.get("native_expected") else "NO_METADATA_ASSERTION"
        if owner is None:
            record.update(status="required_unmeasured",reason="actual owner/cell glyph bounds unavailable")
            results.append(record);continue
        domain=np.zeros((h,w),bool);ox,oy,ow,oh=owner
        left=max(0,int(math.floor(ox-x-1)));top=max(0,int(math.floor(oy-y-1)))
        right=min(w,int(math.ceil(ox+ow-x+1)));bottom=min(h,int(math.ceil(oy+oh-y+1)))
        if right>left and bottom>top:domain[top:bottom,left:right]=True
        b[~domain]=255
        for peer in contracts(state):
            if peer["object"]==row["object"]:continue
            peer_box=ink.get(peer["id"]) or ink.get(peer["object"])
            if peer_box:
                px,py,pw,ph=peer_box
                l=max(0,int(math.floor(px-x)));t=max(0,int(math.floor(py-y)))
                r=min(w,int(math.ceil(px+pw-x)));bt=min(h,int(math.ceil(py+ph-y)))
                if r>l and bt>t:b[t:bt,l:r]=255
        measured=glyph_metrics(s,b,row.get("selector"));record.update(measured)
        if "values" in record:
            record["values"].update(semantic(row,native[row["id"]]))
            if row["math_spans"]:
                # Script/accent islands must not create fictitious text rows.
                record["values"]["line_count_match"]=float(
                    native[row["id"]]["paragraph_count"]==row["source_line_count"])
                if record["values"]["line_anchor_error_px"] is None and record["details"]["actual_ink_bbox"] is not None:
                    record["values"]["line_anchor_error_px"]=float(abs(
                        record["details"]["source_ink_bbox"][1]+record["details"]["source_ink_bbox"][3]-
                        record["details"]["actual_ink_bbox"][1]-record["details"]["actual_ink_bbox"][3]))
            below=[k for k,t in row["targets"].items() if record["values"].get(k) is None or
                   (record["values"][k]>t if k in LOWER else record["values"][k]<t)]
            record["below_targets"]=below;record["status"]="minimum_met" if not below else "below_fixed_targets"
            # Every originally observed math role has its own independent score.
            record["zones"]=[]
            for zone in row.get("zones",[]):
                zx,zy,zw,zh=map(round,zone["roi"]);local=[zx-x,zy-y,zx-x+zw,zy-y+zh]
                zs=Image.fromarray(s).crop(local);za=Image.fromarray(b).crop(local)
                record["zones"].append({"role":zone["role"],"roi":zone["roi"],
                        **glyph_metrics(zs,za,row.get("selector"))})
            for zone in record["zones"]:
                for key in ("foreground_f1","edge_f1_1px","shape_f1","position_error_px"):
                    value=zone.get("values",{}).get(key);target=row["targets"][key]
                    if value is None or (value>target if key in LOWER else value<target):
                        record["below_targets"].append("zone:"+zone["role"]+":"+key)
            if record["below_targets"]:record["status"]="below_fixed_targets"
            pitch=record["details"].get("line_pitch_error_px")
            record["line_spacing"]={"error_px":pitch,"target_px":1.5,
                "status":"measured" if pitch is not None else "not_applicable_or_unresolved",
                "limits":"Original ink-row pitch; never inferred from font point size. Single-line/script islands have no prose-pitch assertion."}
            if pitch is not None and pitch>1.5:
                record["below_targets"].append("line_pitch_error_px");record["status"]="below_fixed_targets"
        results.append(record)
    result={"schema":VERSION,"basis":basis,"items":results,
            "glyph_ownership_status":"BOUND_ACTUAL_NATIVE_TEXT_LAYER" if ownership else "LEGACY_BOUNDS_ONLY",
            "status":"minimum_met" if results and all(r["status"]=="minimum_met" for r in results) else "below_or_unmeasured",
            "limits":"Per-owner metrics, not a whole-page score. Requested fonts/PDF aliases are not font-use proof. Original raster metadata stays unknown; absolute targets/residuals are retained."}
    write_json(cache_file,result);return result


def packet_diagnostics(task,state,candidate,roi):
    """Attach only nearby text rows, keeping full metrics on disk."""
    x,y,w,h=roi
    ids=[r["id"] for r in contracts(state) if r["roi"][0]<x+w and r["roi"][1]<y+h
         and r["roi"][0]+r["roi"][2]>x and r["roi"][1]+r["roi"][3]>y]
    if not ids or not candidate.get("preview"):return None
    result=report(task,state,candidate,ids[:8])
    file=safe_path(task,candidate["pptx"]["path"]).parent/"textbox-packet-diagnostics.json"
    write_json(file,result)
    return {"path":file.relative_to(task).as_posix(),"sha256":digest(file),
            "owners":[{"id":r["id"],"status":r["status"],"below_targets":r.get("below_targets",[]),
                       "content_match":r.get("values",{}).get("content_match"),
                       "math_match":r.get("values",{}).get("math_match")} for r in result["items"]],
            "read":"Read the located owner/role only. Wrong content/script is a source or serialization cause; bad normalized shape is family/style; good shape with fixed-pose deficits is layout. Missing ownership is unmeasured, not success."}


def retention_guard(task, state, candidate, allowed):
    allowed=_closure(state,safe_path(task,candidate["pptx"]["path"]),allowed)
    required=state["objectives"].get("source_textbox_quality",False)
    guard=guard_saved(task,state,safe_path(task,candidate["pptx"]["path"]),state.get("selected"),allowed)
    if required and guard["unbound"]:
        raise ValueError("Read and bind every native text/math owner before adoption")
    ids=guard["scope"]
    if not ids:
        return {"status":"NO_TEXTBOX_PROPERTY_CHANGED"}
    current=report(task,state,candidate,ids)
    if current["glyph_ownership_status"] not in ("BOUND_ACTUAL_NATIVE_TEXT_LAYER","NATIVE_CORE_WITH_VALID_LEGACY_OWNERS"):
        raise ValueError("Use actual isolated native text ownership before typography retention")
    previous=state.get("selected")
    before=report(task,state,previous,ids) if previous and previous["pptx"]["sha256"]!=candidate["pptx"]["sha256"] else None
    old={r["id"]:r for r in before["items"]} if before else {}
    regressions=[];gains=[]
    for row in current["items"]:
        prior=old.get(row["id"]);v=row.get("values",{})
        if not v or v.get("content_match")!=1 or v.get("math_match")!=1 or v.get("format_match")!=1:
            regressions.append({"id":row["id"],"property":"required_content_or_glyph_measure"})
            continue
        if not prior:
            # Initial drawing establishes working progress. Its independent
            # deficits remain open; it is not a final acceptance certificate.
            continue
        for k,new in v.items():
            earlier=prior.get("values",{}).get(k)
            if new is None:
                regressions.append({"id":row["id"],"property":k,"reason":"required_unmeasured"});continue
            if earlier is None:
                gains.append({"id":row["id"],"property":k,"before":None,"after":new,"reason":"newly_measured_not_numeric_improvement"});continue
            delta=(earlier-new) if k in LOWER else (new-earlier)
            target=row["targets"][k]
            both_satisfactory=(earlier<=target and new<=target) if k in LOWER else (earlier>=target and new>=target)
            # Preserve the declared satisfactory interval; don't force a
            # meaningless monotonic race among already-met sampling values.
            if delta < -ALLOWANCE[k] and not both_satisfactory:
                regressions.append({"id":row["id"],"property":k,"before":earlier,"after":new})
            if delta > ALLOWANCE[k]:gains.append({"id":row["id"],"property":k,"before":earlier,"after":new})
        # Multiline pitch has its own protection; good shape or first-line
        # alignment cannot compensate drifting later lines.
        oldpitch=prior.get("details",{}).get("line_pitch_error_px")
        newpitch=row.get("details",{}).get("line_pitch_error_px")
        if oldpitch is not None and (newpitch is None or
                (newpitch-oldpitch>.75 and not(oldpitch<=1.5 and newpitch<=1.5))):
            regressions.append({"id":row["id"],"property":"line_pitch_error_px","before":oldpitch,"after":newpitch})
        # Math roles cannot be hidden by the enclosing textbox aggregate.
        oldzones={z["role"]:z for z in prior.get("zones",[])}
        for zone in row.get("zones",[]):
            a=oldzones.get(zone["role"],{}).get("values",{});b=zone.get("values",{})
            for k in ("foreground_f1","edge_f1_1px","shape_f1","weight_agreement","size_agreement","position_error_px"):
                if b.get(k) is None:
                    regressions.append({"id":row["id"],"property":"zone:"+zone["role"]+":"+k,"reason":"required_unmeasured"})
                elif a.get(k) is None:
                    gains.append({"id":row["id"],"property":"zone:"+zone["role"]+":"+k,"before":None,"after":b[k],"reason":"newly_measured_not_numeric_improvement"})
                else:
                    delta=a[k]-b[k] if k in LOWER else b[k]-a[k]
                    target=row["targets"][k]
                    satisfied=(a[k]<=target and b[k]<=target) if k in LOWER else (a[k]>=target and b[k]>=target)
                    if delta < -ALLOWANCE[k] and not satisfied:
                        regressions.append({"id":row["id"],"property":"zone:"+zone["role"]+":"+k})
    receipt={"schema":VERSION,"before":before,"after":current,"regressions":regressions,"gains":gains,
             "status":"SCOPED_PROGRESS_RETAINABLE" if not regressions else "REJECTED_TYPOGRAPHY_REGRESSION",
             "limits":"Retained progress does not turn below-target measures into passed. Missing/content/script failures cannot be waived."}
    write_json(safe_path(task,candidate["pptx"]["path"]).parent/"textbox-retention.json",receipt)
    if regressions:
        raise ValueError("Preserve independent text/math properties; typography regression: "+str(regressions[:4]))
    return receipt


def source_packet(task,state,candidate):
    """Return at most eight blind original crops; never print expected answers."""
    from .native_core import contract
    selected=set((contract(task) or {}).get('owners',[]))
    missing=[r for r in discover(task,state,candidate)["items"] if r["status"]!="bound" and r['object'] not in selected]
    if missing and state.get('native_import') and state.get('quality_intent'):
        from .retained_typography import report as retained_report
        retained=retained_report(task,state,candidate)
        verified={r['object'] for r in retained['items'] if r['passed']}
        missing=[r for r in missing if r['object'] not in verified]
    if not missing:return None
    from . import source_quality as q
    selected=missing[:8];views=[];items=[]
    directory=task/"source/textbox-reading";directory.mkdir(exist_ok=True)
    for row in selected:
        unit=next((u for u in state.get("quality",{}).get("units",{}).values()
                   if u.get("profile") in ("text","formula","table") and row["object"] in u.get("objects",[])),None)
        if unit:
            roi=unit['roi']
        else:
            owner=next((cid for cid,built in state.get('built',{}).items()
                        if any(o['id']==row['object'] for o in built.get('objects',[]))),None)
            comp=next((c for c in state['analysis']['components'] if c['id']==owner),None)
            roi=comp['roi'] if comp else [0,0,*state['source_size']]
            row={**row,'status':'requires_original_roi_localization',
                 'localization':'This is the independent source component search area, not the saved textbox box. Locate and freeze the exact original glyph ROI before submitting readings.'}
        x,y,w,h=map(round,roi)
        file=directory/(identity({"source":state["source"]["sha256"],"roi":roi})+".png")
        if not file.exists():
            Image.open(task/state["source"]["path"]).convert("RGB").crop((x,y,x+w,y+h)).save(file)
        view={"role":"source","path":file.relative_to(task).as_posix(),"sha256":digest(file),"roi":roi}
        views.append(view);items.append({**row,"roi":roi,"source_view":view,
             "read":"Read only these original pixels; native strings/prior expectations are deliberately absent. Record exact words, math AST, visible line structure and glyph domains. Reload references/textbox-quality.md only for this action."})
    return {"items":items,"views":views,"source_hash":state["source"]["sha256"],
            "baseline_hash":candidate["pptx"]["sha256"],
            "return":"entries[{id,object,roi,word_rois,ownership_evidence,readings[2 distinct original_only readings with source_view,text,math_spans,evidence],zones?,targets?}]"}


def mirror(scene,pptx,allowed):
    """Mirror saved native typography without recreating unrelated objects."""
    from .relative_layout import mirror_native_positions
    result=mirror_native_positions(scene,pptx,allowed)
    rows={r["object"]:r for r in native_rows(pptx) if "@cell:" not in r["id"]}
    for obj in result["objects"]:
        fields=allowed.get(obj["id"],[]);row=rows.get(obj["id"])
        if not fields or not row:continue
        run=next((r for r in row["native_styles"] if r.get("size_pt") is not None),None)
        if run:
            values={"style.font_size":run["size_pt"],"style.bold":run["bold"]=="1",
                    "style.italic":run["italic"]=="1","style.font":next(iter(run["font"]),None),
                    "style.char_spacing_pt":float(run.get("spacing_100pt") or 0)/100}
            for field in fields:
                if field in values and values[field] is not None:obj["style"][field[6:]]=values[field]
        node=inventory(xml(pptx))[0][obj["id"]]["node"]
        if "style.line_spacing_pt" in fields:
            spacing=node.find("p:txBody/a:p/a:pPr/a:lnSpc/a:spcPts",NS)
            if spacing is not None:obj["style"]["line_spacing_pt"]=float(spacing.get("val"))/100
        if "text_body" in fields and not row["math_spans"]:
            paragraphs=[]
            for p in node.findall("p:txBody/a:p",NS):
                words=[]
                for r in p.findall("a:r",NS):
                    props=r.find("a:rPr",NS)
                    if props is None:
                        from lxml import etree
                        props=etree.Element("{"+NS["a"]+"}rPr")
                    face=props.find("a:latin",NS)
                    words.append({"text":r.findtext("a:t",namespaces=NS) or "",
                        "font":face.get("typeface") if face is not None else obj["style"].get("font"),
                        "font_size":float(props.get("sz"))/100 if props.get("sz") else obj["style"].get("font_size"),
                        "bold":props.get("b")=="1","italic":props.get("i")=="1",
                        "char_spacing_pt":float(props.get("spc") or 0)/100})
                paragraphs.append(words)
            obj["content"]={"paragraphs":paragraphs}
        if "text_content" in fields:
            if obj["content"].get("paragraphs"):
                node=inventory(xml(pptx))[0][obj["id"]]["node"]
                native=[r.text or "" for r in node.findall(".//a:t",NS)]
                old=[r for p in obj["content"]["paragraphs"] for r in p if "text" in r]
                if len(old)!=len(native):raise ValueError("Saved native text topology cannot be mirrored")
                for target,text in zip(old,native):target["text"]=text
            else:obj["content"]["text"]=row["text"]
    return result


def prepare_atlas(task, request):
    """Cache native glyph-only samples; no stale scene text/style rebuilding."""
    from . import agent_runtime as a
    from .native_commit import _nodes
    from .agent_edit import _set
    from .inspect import active_children
    from pptx import Presentation
    from lxml import etree as E
    task,state=a._state(task);candidate=state.get("selected")
    if not candidate or request.get("baseline_hash")!=candidate["pptx"]["sha256"]:
        raise ValueError("Bind the current actual native font baseline")
    ids=request.get("ids",[])
    if not 1<=len(ids)<=8 or len(set(ids))!=len(ids):
        raise ValueError("Use1..8 representative original-bound native words")
    wanted=[r for r in contracts(state) if r["id"] in ids]
    native={r["id"]:r for r in native_rows(task/candidate["pptx"]["path"])}
    if len(wanted)!=len(ids) or any(native[r["id"]]["math_spans"] or "@cell:" in r["id"] for r in wanted):
        raise ValueError("Use source-confirmed plain text; mixed math follows source_math")
    variants=request.get("variants",[])
    if not 1<=len(variants)<=4 or any(set(v)-{"font","bold","italic","font_size"} for v in variants):
        raise ValueError("Bound at most4 explicit typeface/style/size alternatives")
    receipt=read_json(task/candidate["render"]["path"])
    basis={"pptx":candidate["pptx"]["sha256"],"owners":ids,"contracts":identity(wanted),
           "variants":variants,"host":receipt.get("backend"),"fonts":receipt.get("font_inventory_hash"),"method":VERSION}
    directory=task/"cache/text-font"/identity(basis);directory.mkdir(parents=True,exist_ok=True)
    manifest=directory/"atlas.json"
    if manifest.exists():
        result=read_json(manifest)
        if digest(directory/"Atlas.pptx")!=result["pptx_hash"]:
            raise ValueError("Cached native font sample changed")
        return result
    deck=Presentation(task/candidate["pptx"]["path"]);names={r["object"] for r in wanted}
    def prune(parent):
        for node in list(parent):
            tag=E.QName(node).localname
            if tag=="grpSp":
                prune(node)
                if not list(active_children(node)):parent.remove(node)
            elif tag in ("sp","cxnSp","pic","graphicFrame"):
                props=node.find(".//p:cNvPr",NS)
                if props is None or props.get("name") not in names:parent.remove(node)
    # Only the current one-page figure; preserve its native parent transforms.
    if len(deck.slides)!=1:raise ValueError("Bind one figure/slide before font sampling")
    prune(deck.slides[0]._element.find("p:cSld/p:spTree",NS))
    leaves=[copy.deepcopy(n) for n in deck.slides[0].shapes._spTree if E.QName(n).localname in ("sp","grpSp")]
    for variant in variants:
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        for n in leaves:slide.shapes._spTree.insert_element_before(copy.deepcopy(n),"p:extLst")
        for name,node in _nodes(slide._element).items():
            if name not in names:continue
            fresh=E.fromstring(E.tostring(node))
            for key,value in variant.items():_set(fresh,"style."+key,value)
            node.getparent().replace(node,fresh)
    file=directory/"Atlas.pptx";deck.save(file)
    result={"schema":"native-text-font-atlas-1","basis":basis,"pptx":str(file),"pptx_hash":digest(file),
            "styles":[{"inherited_actual_native":True}]+variants,
            "original_crops":[r["readings"][1]["source_view"] for r in wanted],
            "status":"SAVED_SAMPLE_REQUIRES_ACTUAL_SAME_HOST_RENDER",
            "limits":"Screen this actual native sample once; full-context fixed per-word metrics and protection decide adoption. No font installation or source raster metadata claim."}
    write_json(manifest,result);return result
