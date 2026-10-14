import sys,copy,tempfile,unittest,json,math
from pathlib import Path
from zipfile import ZipFile
import numpy as np
from PIL import Image,ImageDraw,ImageFilter
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree as E
D=Path(__file__).resolve().parent;sys.dont_write_bytecode=True;sys.path.insert(0,str(D/'skill/scripts'))
from core.common import write_json,read_json,digest
from core import recursive_containment as rc,office_math_contract as mc,task_completion as tc,icon_clarity as ic,target_actions as ta
from core.mathml import NS,math_zone,math_node,linear_math
from core.inspect import xml,inventory,parse_math_container
from core.boundary_geometry import saved_nodes
from core import neighbor_geometry as ng
from core.neighbor_actions import _payload,_same_zip
from core.source_reading import semantic_ast
from core.native_task import altered,canonical_math

RESULTS=[]
class Tests(unittest.TestCase):
 def setUp(self):
  self.temp=tempfile.TemporaryDirectory(dir=D,prefix='v121-check-');self.t=Path(self.temp.name)
  self.addCleanup(self.temp.cleanup)
  im=Image.new('RGB',(220,220),'white');im.save(self.t/'original.png')
  self.s={'source':{'path':'original.png','sha256':digest(self.t/'original.png')},'source_size':[220,220],'quality':{'units':{'part':{'id':'part'}}}}
  p=Presentation();p.slide_width=220*9525;p.slide_height=220*9525;slide=p.slides.add_slide(p.slide_layouts[6])
  for name,preset,bbox in [('outer',MSO_SHAPE.RECTANGLE,[5,5,210,210]),('middle',MSO_SHAPE.ROUNDED_RECTANGLE,[40,45,140,135]),('inner',MSO_SHAPE.OVAL,[70,65,65,60]),('a',MSO_SHAPE.RECTANGLE,[89,84,5,5]),('b',MSO_SHAPE.RECTANGLE,[100,86,6,6]),('protected',MSO_SHAPE.RECTANGLE,[210,210,3,3])]:
   shape=slide.shapes.add_shape(preset,*[round(v*9525) for v in bbox]);shape.name=name;shape.fill.solid();shape.line.fill.background()
  p.save(self.t/'Input.pptx');self.c={'pptx':{'path':'Input.pptx','sha256':digest(self.t/'Input.pptx')}}
  self.n=saved_nodes(self.t/'Input.pptx',self.s['source_size']);self.paint,self.errors=ng.read(self.t,self.s,self.c)
  self.contract={'schema':rc.VERSION,'source_hash':self.s['source']['sha256'],'canvas':self.s['source_size'],
   'source_units':[{'id':'part','role':'element','target':'pair','source_evidence':'Original paired elements'}],
   'elements':[{'id':'pair','objects':['a','b'],'source_bbox':[89,84,17,8],'source_evidence':'Two distinct original sibling parts'}],
   'containers':[{'id':i,'members':m,'source_bbox':self.n[body]['frame'],'source_evidence':'Original nested body and direct members',
     'domain':{'kind':'native_body','object':body,'padding_px':[0,0,0,0]},'target':{'axes':['x','y'],'center':'domain_bbox_center','measurement_basis':'direct_member_paint_union','authorization':'Explicit whole-content center target','source_uncertainty_px':0.},'compiled_tolerance_px':.75}
     for i,body,m in [('inner-c','inner',['pair']),('middle-c','middle',['inner-c']),('outer-c','outer',['middle-c'])]]}
  from core.centering_contract import decisions
  by={r['id']:r for r in self.contract['elements']};cs={r['id']:r for r in self.contract['containers']}
  for row in cs.values():row['member_alignment']=decisions(row,by,cs)
 def test_three_levels_each_axis_and_no_parent_shortcut(self):
  r=rc.evaluate(self.contract,self.s,self.paint,self.n)
  self.assertEqual(r['coverage']['required_measurements']['required'],6);self.assertEqual(len(r['items']),3)
  self.assertFalse(r['passed']);self.assertTrue(any(not v['passed'] for i in r['items'] for v in i['axes'].values()))
 def test_aggregate_center_preserves_sibling_gap(self):
  r=rc.evaluate(self.contract,self.s,self.paint,self.n);inner=r['items'][0];delta={a:-v['actual'] for a,v in inner['axes'].items()}
  data=_payload(self.t/'Input.pptx',self.s['source_size'],{n:delta for n in ('a','b')})
  from zipfile import ZIP_DEFLATED
  with ZipFile(self.t/'Input.pptx') as z,ZipFile(self.t/'After.pptx','w',ZIP_DEFLATED) as out:
   for info in z.infolist():out.writestr(info,data if info.filename=='ppt/slides/slide1.xml' else z.read(info.filename))
  new=saved_nodes(self.t/'After.pptx',self.s['source_size'])
  oldgap=np.subtract(self.n['b']['frame'][:2],self.n['a']['frame'][:2]);newgap=np.subtract(new['b']['frame'][:2],new['a']['frame'][:2])
  self.assertTrue(np.allclose(oldgap,newgap,atol=.001));self.assertEqual(E.tostring(new['protected']['node']),E.tostring(self.n['protected']['node']))
  c={'pptx':{'path':'After.pptx','sha256':digest(self.t/'After.pptx')}};paint,_=ng.read(self.t,self.s,c)
  check=rc.evaluate(self.contract,self.s,paint,new);self.assertTrue(check['items'][0]['passed']);self.assertFalse(check['passed'])
 def test_missing_original_member_measurement_blocks_both_axes(self):
  paint=dict(self.paint);paint.pop('b');r=rc.evaluate(self.contract,self.s,paint,self.n)
  self.assertEqual(r['items'][0]['status'],'required_unmeasured');self.assertFalse(any(v['passed'] for v in r['items'][0]['axes'].values()))
 def test_missing_axis_rejected(self):
  self.contract['containers'][0]['target']['axes']=['x']
  with self.assertRaises(ValueError):rc.validate(self.contract,self.s)
 def test_missing_source_unit_rejected(self):
  self.contract['source_units']=[]
  with self.assertRaises(ValueError):rc.validate(self.contract,self.s)
 def test_multiple_direct_parents_rejected(self):
  self.contract['containers'][1]['members'].append('pair')
  with self.assertRaises(ValueError):rc.validate(self.contract,self.s)
 def test_cycle_rejected(self):
  self.contract['containers'][0]['members'].append('outer-c')
  with self.assertRaises(ValueError):rc.validate(self.contract,self.s)
 def test_nonrectangular_domain_is_not_nominal_box(self):
  self.paint['a']['geometry']=ng.box(70,65,75,70);r=rc.evaluate(self.contract,self.s,self.paint,self.n)
  self.assertGreater(r['items'][0]['overflow_fraction'],0);self.assertFalse(r['items'][0]['passed'])
 def test_supported_roundrect_and_custom_polygon(self):
  c=copy.deepcopy(self.contract);c['containers'][0]['domain']={'kind':'logical_polygon','points':[[65,65],[140,65],[140,135],[65,135]],'padding_px':[1,2,3,4]}
  r=rc.evaluate(c,self.s,self.paint,self.n);self.assertNotEqual(r['items'][0]['status'],'required_unmeasured');self.assertNotEqual(r['items'][1]['status'],'required_unmeasured')
 def test_bad_custom_domain_is_unmeasured(self):
  self.contract['containers'][0]['domain']['object']='no-body';r=rc.evaluate(self.contract,self.s,self.paint,self.n)
  self.assertEqual(r['items'][0]['status'],'required_unmeasured')
 def test_no_depth_cutoff(self):
  c=copy.deepcopy(self.contract);c['containers']=[{**copy.deepcopy(c['containers'][0]),'id':'c'+str(i),'members':['pair'] if i==0 else ['c'+str(i-1)]} for i in range(150)]
  self.assertEqual(len(rc.validate(c,self.s)[2]),150)
 def test_original_only_crop_proof_rejects_candidate(self):
  Image.new('RGB',(30,30),'black').save(self.t/'fake.png');r={'path':'fake.png','sha256':digest(self.t/'fake.png'),'roi':[0,0,30,30],'source_hash':self.s['source']['sha256']}
  with self.assertRaises(ValueError):rc.source_view(self.t,self.s,r)
 def test_structured_math_simple_complex_and_products(self):
  token=lambda t:{'op':'token','text':t}
  cases=[token('Q'),{'op':'sub','base':token('t'),'sub':token('L')},{'op':'fraction','numerator':token('x'),'denominator':token('y')},
   {'op':'product','lower':token('i=1'),'upper':token('K'),'body':token('x')}]
  for ast in cases:
   m=math_zone(ast,18,'Cambria Math');actual=parse_math_container(m.find('.//m:oMath',NS))
   self.assertEqual(canonical_math(actual),canonical_math(ast));self.assertTrue(m.findall('.//m:oMath',NS))
  self.assertIn('∏',linear_math(cases[-1]))
 def test_native_runs_builder_rejected(self):
  from core.build import text_body
  p=Presentation();sl=p.slides.add_slide(p.slide_layouts[6]);sh=sl.shapes.add_textbox(0,0,1000000,1000000)
  with self.assertRaises(ValueError):text_body(sh,{'runs':[{'ast':{'op':'token','text':'Q'}}],'inline_math':'native_runs'},{})
 def test_implicit_math_uses_real_zone(self):
  from core.build import text_body
  p=Presentation();sl=p.slides.add_slide(p.slide_layouts[6]);sh=sl.shapes.add_textbox(0,0,1000000,1000000)
  text_body(sh,{'runs':[{'text':'At time '},{'ast':{'op':'token','text':'t'}}]},{'font':'Times New Roman'})
  self.assertEqual(len(sh._element.findall('.//m:oMath',NS)),1);self.assertEqual(mc._segments(sh._element),[[{'text':'At time '},{'ast':{'op':'token','text':'t'}}]])
 def test_plain_text_italic_scripts_are_not_math(self):
  from core.build import text_body
  p=Presentation();sl=p.slides.add_slide(p.slide_layouts[6]);sh=sl.shapes.add_textbox(0,0,1000000,1000000)
  text_body(sh,{'runs':[{'text':'t','italic':True},{'text':'L','baseline':-30000}]},{})
  self.assertEqual(len(sh._element.findall('.//m:oMath',NS)),0);self.assertNotEqual(semantic_ast(mc._segments(sh._element)),semantic_ast([[{'ast':{'op':'sub','base':{'op':'token','text':'t'},'sub':{'op':'token','text':'L'}}}]]))
 def test_operator_edit_probe_has_real_change(self):
  self.assertNotEqual(altered({'op':'token','text':'∑'}),{'op':'token','text':'∑'})
 def test_scope_missing_cannot_grant_relationship_write(self):
  # Real plan computes translations only after exact original binding; no
  # caller coordinates are interpreted by the solver.
  import unittest.mock as mock
  with mock.patch.object(rc,'load',return_value=self.contract),mock.patch.object(rc,'report',return_value=rc.evaluate(self.contract,self.s,self.paint,self.n)):
   report=rc.evaluate(self.contract,self.s,self.paint,self.n);report['capture_required']=[]
   with mock.patch.object(rc,'report',return_value=report):
    with self.assertRaises(ValueError):rc.repair_plan(self.t,self.s,self.c,{'max_translation_px':20,'write_set':{}})
 def test_branch_pruning_keeps_debt_and_independent_work(self):
  flow=tc.initialize(self.s);flow['backlog']={'a':{'basis':'b'},'independent':{'basis':'b'}}
  req={'obligations':['a'],'strategy':'same-strategy','attempt':{'outcome':'no_gain','evidence':[{'path':'report.json','sha256':'retained'}]}}
  tc.branch(self.s,req);r=tc.branch(self.s,req);self.assertTrue(r['pruned']);self.assertEqual(set(flow['backlog']),{'a','independent'});self.assertEqual(flow['status'],'ACTIVE')
 def test_no_draft_or_workload_stop(self):
  tc.initialize(self.s)
  for reason in ('draft','time','partial_pass','branch_exhausted'):
   with self.assertRaises(ValueError):tc.stop_report(self.t,self.s,self.c,{'reason':reason})
 def test_status_does_not_replace_goal_or_debt(self):
  flow=tc.initialize(self.s);flow['backlog']={'a':{'basis':'b'}};before=copy.deepcopy(flow);tc.amend(self.s,{'kind':'status'})
  self.assertEqual(flow,before)
 def test_supplement_keeps_original_and_debt_but_expires_acceptance(self):
  flow=tc.initialize(self.s);flow['backlog']={'a':{'basis':'b'}};self.s['acceptance']={'old':True};original=copy.deepcopy(self.s['source'])
  tc.amend(self.s,{'kind':'user_target','user_evidence':'actual human message','target':'also rebuild another diagram'})
  self.assertEqual(self.s['source'],original);self.assertEqual(set(flow['backlog']),{'a'});self.assertNotIn('acceptance',self.s)
 def test_true_block_cannot_hide_independent_work(self):
  flow=tc.initialize(self.s);flow['backlog']={'a':{},'independent':{}}
  write_json(self.t/'error.json',{'operation':'actual capability route','error':'external failure'})
  b={k:'real saved evidence' for k in ('operation','error','impact','attempted_routes','no_remaining_route','recovery')};b.update(obligations=['a'],evidence=[{'path':'error.json','sha256':digest(self.t/'error.json')}])
  with self.assertRaises(ValueError):tc.stop_report(self.t,self.s,self.c,{'reason':'external_block','blocks':[b]})
 def test_resume_roundtrip_preserves_target_debt_branch(self):
  f=tc.initialize(self.s);f.update(backlog={'a':{'basis':'old','next_action':'actual capture'}},branches={'x':{'attempts':[1,2],'pruned':True}})
  write_json(self.t/'state.json',self.s);again=read_json(self.t/'state.json');self.assertEqual(again,self.s);self.assertEqual(tc.target_basis(again),tc.target_basis(self.s))
 def test_clearer_edges_and_stochastic_type_are_separate(self):
  im=Image.new('RGB',(80,80),'white');draw=ImageDraw.Draw(im);draw.rectangle((15,20,65,60),fill='navy');draw.ellipse((20,48,32,62),fill='black');draw.ellipse((48,48,60,62),fill='black')
  sharp=ic.statistics(im);blur=ic.statistics(im.filter(ImageFilter.GaussianBlur(2)))
  self.assertGreater(sharp['edge_sharpness']/blur['edge_sharpness'],1.05)
 def test_enhancement_is_not_original_or_output(self):
  self.s['quality']['units']['part'].update(profile='illustration',roi=[0,0,70,70],objects=['a'])
  old=digest(self.t/'original.png');r=ic.enhance(self.t,self.s,'part');self.assertTrue(r['diagnostic_only']);self.assertEqual(digest(self.t/'original.png'),old);self.assertEqual(len(r['views']),3)
 def test_clarity_accepts_exact_native_canvas_pixel_quantization(self):
  self.s['source_size']=[1280,594];Image.new('RGB',(1600,742),'navy').save(self.t/'preview.png');self.c['preview']={'path':'preview.png','sha256':digest(self.t/'preview.png')}
  crop=ic._crop(self.t,self.s,[0,0,80,60],self.c);self.assertEqual(crop.size,(80,60))
 def test_clarity_rejects_actual_stretched_canvas(self):
  self.s['source_size']=[1280,594];Image.new('RGB',(1600,740),'navy').save(self.t/'preview.png');self.c['preview']={'path':'preview.png','sha256':digest(self.t/'preview.png')}
  with self.assertRaises(ValueError):ic._crop(self.t,self.s,[0,0,80,60],self.c)
 def test_failed_leaf_sets_not_counts(self):
  r=tc.failed_leaves({'g':{'passed':False,'items':[{'id':'x','passed':True},{'id':'y','passed':False,'status':'required_unmeasured'}]}})
  self.assertEqual({a['id'] for a in r},{'/g','/g/items/y'})
 def test_missing_nested_pair_decision_rejected(self):
  self.contract['hierarchy_decisions']=[]
  with self.assertRaises(ValueError):rc.validate(self.contract,self.s)
 def test_explicit_nested_decision_cannot_replace_direct_edge(self):
  pairs=rc.hierarchy_survey(self.contract)
  self.contract['hierarchy_decisions']=[{**r,'decision':'independent','source_evidence':'Fixture original'} for r in pairs]
  with self.assertRaises(ValueError):rc.validate(self.contract,self.s)
 def test_parent_pass_cannot_hide_inner_failure(self):
  # A deliberately large outer container makes its own x/y target pass while
  # the actual child is off-center; ancestry closure still fails.
  report=rc.evaluate(self.contract,self.s,self.paint,self.n)
  inner=report['items'][0];self.assertFalse(inner['passed'])
  self.contract['containers'][-1]['domain']={'kind':'logical_polygon','points':[[40,45],[180,45],[180,180],[40,180]],'padding_px':[0,0,0,0]}
  report=rc.evaluate(self.contract,self.s,self.paint,self.n)
  self.assertTrue(report['items'][-1]['own_geometry_pass']);self.assertFalse(report['items'][-1]['passed'])
 def test_bbox_internal_center_is_not_external_center(self):
  # The painted glyph geometry is independent of its text-frame layout flags.
  self.paint['a']['geometry']=ng.box(74,70,78,75);self.paint['b']['geometry']=ng.box(80,72,83,77)
  r=rc.evaluate(self.contract,self.s,self.paint,self.n);self.assertFalse(r['items'][0]['axes']['x']['passed']);self.assertFalse(r['items'][0]['axes']['y']['passed'])
 def test_original_bias_does_not_exempt_zero_target(self):
  self.contract['containers'][0]['original_bias']={'x':14,'y':9,'source_evidence':'Biased source fixture'}
  self.assertFalse(rc.evaluate(self.contract,self.s,self.paint,self.n)['items'][0]['passed'])
 def test_upgrade_criteria_cannot_reduce_to_file_exists(self):
  with self.assertRaises(ValueError):tc.initialize(self.s,'skill_upgrade',criteria=['file_exists'])
 def test_body_fit_requires_explicit_dimensions_grant(self):
  import unittest.mock as mock
  with mock.patch.object(rc,'load',return_value=self.contract):
   with self.assertRaises(ValueError):rc.repair_plan(self.t,self.s,self.c,{'operation':'fit_child_body','container':'inner-c','max_scale_loss':.05,'write_set':{'inner':['position.x','position.y']}})
 def test_missing_formula_behavior_cannot_be_content_pass(self):
  r=mc.behavior_report(self.t,self.s,self.c,{'owners':[{'object':'a','paragraphs':[[{'ast':{'op':'token','text':'Q'}}]]}]})
  self.assertFalse(r['passed']);self.assertEqual(r['status'],'required_unmeasured')
 def test_cell_formula_writer_keeps_tracks_and_other_cells(self):
  p=Presentation();sl=p.slides.add_slide(p.slide_layouts[6]);sh=sl.shapes.add_table(1,2,0,0,4000000,1000000);sh.name='T';sh.table.cell(0,0).text='Rate x';sh.table.cell(0,1).text='protected';file=self.t/'Table.pptx';p.save(file)
  before=inventory(xml(file))[0]['T']['node'];grid=E.tostring(before.find('.//a:tblGrid',NS));rows=before.findall('.//a:tc',NS)
  body={'T':{'paragraphs':[[{'text':'Rate '},{'ast':{'op':'token','text':'x'}}],[{'text':'protected'}]]}}
  data=ta._math_payload(file,body);after=inventory(E.fromstring(data))[0]['T']['node']
  self.assertEqual(E.tostring(after.find('.//a:tblGrid',NS)),grid);self.assertEqual(len(after.findall('.//m:oMath',NS)),1)
  self.assertEqual(mc._segments(after),[[{'text':'Rate '},{'ast':{'op':'token','text':'x'}}],[{'text':'protected'}]])
 def test_current_mathematical_structure_not_plain_metadata(self):
  p=Presentation();sl=p.slides.add_slide(p.slide_layouts[6]);sh=sl.shapes.add_textbox(0,0,1000000,1000000);sh.name='Q';sh.text='Q';file=self.t/'Q.pptx';p.save(file)
  n=inventory(xml(file))[0]['Q']['node'];self.assertFalse(any('ast' in s for row in mc._segments(n) for s in row));self.assertEqual(len(n.findall('.//m:oMath',NS)),0)
 def test_forged_accepted_result_cannot_avoid_fresh_fail(self):
  import unittest.mock as mock
  from core import quality_loop as ql,quality_intent as qi,execution_contract as ec,object_lifecycle as ol,global_layout_actions as gl
  Image.new('RGB',(220,220),'white').save(self.t/'preview.png');self.c['preview']={'path':'preview.png','sha256':digest(self.t/'preview.png')}
  result={'status':'ACCEPTED','basis':'fixed','implementation':'fixed','pptx_hash':self.c['pptx']['sha256'],'preview_hash':self.c['preview']['sha256'],
   'gates':{k:{'passed':True} for k in ('recursive_containment','structured_office_math','icon_clarity','all_neighbor_closure','coherent_paper_geometry')}}
  write_json(self.t/'claimed.json',result);self.s['quality_result']={'path':'claimed.json','sha256':digest(self.t/'claimed.json')}
  with mock.patch.object(ql,'_specifications',return_value={'passed':True}),mock.patch.object(gl,'obligations',return_value=[]),mock.patch.object(qi,'validate',return_value={}),mock.patch.object(ol,'verify_current'),mock.patch.object(ec,'_basis',return_value='fixed'),mock.patch.object(ec,'_implementation',return_value='fixed'),mock.patch.object(ql,'audit',return_value={'status':'WORK_REQUIRED','gates':{'critical':{'passed':False}},'failures':[{'goal':'critical'}]}):
   with self.assertRaises(ValueError):ql.final_gate(self.t,self.s,self.c,measured=result)
 def test_numeric_edit_preserves_other_digits(self):
  before={'op':'token','text':'0.99'};after=altered(before)
  self.assertEqual(after['text'],'1.99')
  from core.native_task import atom_change
  self.assertEqual(atom_change(before,after),('0','1',0))
 def test_bold_letter_style_does_not_change_base_semantic_letter(self):
  self.assertEqual(canonical_math({'op':'token','text':'𝐳'}),{'op':'token','text':'z'})
  self.assertEqual(canonical_math({'op':'token','text':'𝒢'}),{'op':'token','text':'𝒢'})
 def test_sharpness_alone_cannot_accept_an_icon(self):
  im=Image.new('RGB',(70,70),'white');draw=ImageDraw.Draw(im);draw.rectangle((10,10,60,60),fill='navy');im.filter(ImageFilter.GaussianBlur(2)).save(self.t/'original.png');im.save(self.t/'preview.png')
  self.s['source']['sha256']=digest(self.t/'original.png');self.s['source_size']=[70,70];self.s['quality']['units']['part'].update(profile='illustration',roi=[0,0,70,70],objects=list(self.n),features=[{'id':'body'}])
  from core.icon_hd import source_survey
  from core.quality_intent import source_view
  source=source_view(self.t,self.s,[0,0,70,70],1)
  target={'representation':'editable_vector','profile':'flat','min_sharpness_gain':1.05,'min_edge_support':.01,
   'edge_probes':[{'id':str(i),'points':p,'source_evidence':'Original rectangular boundary'} for i,p in enumerate([[[2,30],[18,30]],[[30,2],[30,18]],[[68,30],[53,30]]])]}
  survey=[dict(w,decision='covered',units=['part'],source_evidence='Full original reviewed') for w in source_survey(self.t,self.s)]
  self.c['preview']={'path':'preview.png','sha256':digest(self.t/'preview.png')}
  write_json(self.t/'render.json',{'backend':'independent_fixture','pptx_hash':self.c['pptx']['sha256'],'preview_hash':self.c['preview']['sha256']})
  self.c['render']={'path':'render.json','sha256':digest(self.t/'render.json')}
  ic.bind(self.t,self.s,self.c,{'schema':ic.VERSION,'source_hash':self.s['source']['sha256'],'source_survey':survey,'items':[{'unit':'part','category':'blurred_icon','source_evidence':'Blurred structured fixture','source_view':source,'structure_features':['body'],'relations':[],'relation_inventory_evidence':'Intentionally incomplete anatomy; numerical sharpness cannot waive it','target':target}]})
  self.c['preview']={'path':'preview.png','sha256':digest(self.t/'preview.png')};r=ic.report(self.t,self.s,self.c)
  gain=ic.statistics(Image.open(self.t/'preview.png'))['edge_sharpness']/ic.statistics(Image.open(self.t/'original.png'))['edge_sharpness']
  self.assertGreater(gain,1.05);self.assertFalse(r['passed']);self.assertFalse(r['items'][0].get('structure_pass',False))
 def test_native_output_encoding_cannot_hide_original_failure(self):
  from core.office import decode_native_output
  self.assertEqual(decode_native_output(None),'');self.assertEqual(decode_native_output('原错误'),'原错误')
  self.assertEqual(decode_native_output('原错误'.encode('utf-8')),'原错误')
  decoded=decode_native_output('保存错误'.encode('cp936'))
  self.assertTrue(decoded);self.assertNotIn('NoneType',decoded)
 def test_actual_host_job_paths_use_shared_wsl_bridge(self):
  import unittest.mock as mock
  from core import office_bridge as bridge
  with mock.patch.object(bridge.platform,'system',return_value='Linux'),mock.patch.object(bridge.subprocess,'run') as command:
   command.return_value.stdout='/mnt/c/Users/TestUser/AppData/Local/Temp/owned-job\n'
   result=bridge._local_path(r'C:\Users\TestUser\AppData\Local\Temp\owned-job')
   self.assertEqual(command.call_args.args[0][:2],['wslpath','-u']);self.assertEqual(result.name,'owned-job')
 def test_math_atom_pass_cannot_hide_failed_host_or_prose(self):
  import unittest.mock as mock
  from core import native_task as nt,execution_contract as ec
  contract={'owners':[{'object':'a','paragraphs':[[{'ast':{'op':'token','text':'Q'}}]]}]}
  ref={'path':'receipt.json','sha256':'fixture'};self.s['office_math_behavior']=ref
  record={'pptx_hash':self.c['pptx']['sha256'],'implementation':'current','plan':{'path':'Plan.json'},'office_directory':'Office','independent_check':{'path':'Independent.json'}}
  for failure in ({'status':'FAIL','actual_host_job_verified':False},{'status':'FAIL','actual_host_job_verified':True}):
   checked={**failure,'checks':[{'object_id':'a','action':'math','status':'PASS'}]}
   with mock.patch.object(ec,'_implementation',return_value='current'),mock.patch.object(ec,'resource',return_value=self.t/'receipt.json'),mock.patch.object(mc,'read_json',return_value=record),mock.patch.object(nt,'verify',return_value=checked):
    self.assertFalse(mc.behavior_report(self.t,self.s,self.c,contract)['passed'])
 def test_group_is_not_a_second_owner_of_descendant_math_prose(self):
  from core.native_task import plan
  from core.build import text_body
  p=Presentation();sl=p.slides.add_slide(p.slide_layouts[6]);group=sl.shapes.add_group_shape();group.name='semantic-group'
  child=group.shapes.add_textbox(100000,100000,2000000,500000);child.name='actual-owner'
  text_body(child,{'runs':[{'text':'At '},{'ast':{'op':'token','text':'t'}}]},{'font':'Times New Roman','font_size':18})
  file=self.t/'Grouped-math.pptx';p.save(file)
  scene={'canvas':{'width':960,'height':720,'target_width_inches':10},'objects':[
   {'id':'semantic-group','render_type':'group','parent':None,'content':{},'style':{},'bbox':[0,0,1,1]},
   {'id':'actual-owner','render_type':'text','parent':'semantic-group','content':{},'style':{},'bbox':[0,0,1,1]}],'relations':[],'prototypes':[]}
  result=plan(scene,file,self.t/'Group-plan.json')
  self.assertEqual(set(result['protected_prose']),{'actual-owner'});self.assertEqual(result['protected_prose']['actual-owner']['math_count'],1)
 def test_upgrade_cannot_claim_unexecuted_actual_behavior(self):
  from core.upgrade_validation import verify_behaviors
  with self.assertRaises(ValueError):verify_behaviors([],self.t)
 def test_repeated_formula_spans_cannot_collapse_to_one_passing_row(self):
  import unittest.mock as mock
  from core import native_task as nt
  from core.upgrade_validation import verify_behaviors
  plan=self.t/'Frozen-probe.json';write_json(plan,{'fixture':'two math spans in one owner'})
  rows=[{'object_id':'mixed','action':'math','status':'FAIL'},{'object_id':'mixed','action':'math','status':'PASS'},{'object_id':'line','action':'connector','status':'PASS'}]
  entry={'plan':str(plan),'plan_sha256':digest(plan),'office_directory':str(self.t),'expected_checks':[{**r,'status':'PASS'} for r in rows]}
  with mock.patch.object(nt,'verify',return_value={'actual_host_job_verified':True,'current_implementation_verified':True,'checks':rows,'status':'FAIL'}):
   with self.assertRaises(ValueError):verify_behaviors([entry],self.t)
 def test_actual_host_job_cannot_accept_old_executor_identity(self):
  import unittest.mock as mock
  from core import native_task as nt
  from core.upgrade_validation import verify_behaviors
  plan=self.t/'Old-executor-probe.json';write_json(plan,{'implementation':{'version':'old','files_sha256':'old'}})
  rows=[{'object_id':'formula','action':'math','status':'PASS'},{'object_id':'line','action':'connector','status':'FAIL'}]
  entry={'plan':str(plan),'plan_sha256':digest(plan),'office_directory':str(self.t),'expected_checks':rows}
  with mock.patch.object(nt,'verify',return_value={'actual_host_job_verified':True,'current_implementation_verified':False,'checks':rows,'status':'FAIL'}):
   with self.assertRaisesRegex(ValueError,'current implementation'):verify_behaviors([entry],self.t)
 def test_wrong_ppt_cannot_reuse_cached_before_render(self):
  from core.native_task import bind_before_render
  folder=self.t/'wrong-cached-render';folder.mkdir()
  write_json(folder/'render.json',{'backend':'PowerPoint','status':'PASS','scope':'target_render_only','pptx_hash':'another file'})
  write_json(folder/'job-receipt.json',{});write_json(folder/'office-session.json',{})
  with self.assertRaises(ValueError):bind_before_render(self.t/self.c['pptx']['path'],folder/'render.json')
 def test_incomplete_cached_render_cannot_supply_missing_evidence(self):
  from core.native_task import unpack_before_render
  with self.assertRaises(ValueError):unpack_before_render({'files':{'Final_preview.png':{'sha256':'fake','base64':'aGVsbG8='}}},self.t)
 def test_independent_missing_measurement_does_not_block_known_repair(self):
  import unittest.mock as mock
  self.s['recursive_containment']={'path':'fixture-contract.json','sha256':'fixture'}
  check=rc.evaluate(self.contract,self.s,self.paint,self.n);check['capture_required']=[{'object':'independent-table@cell:0:0','reason':'Actual capture missing'}]
  request={'container':'inner-c','max_translation_px':32.,'write_set':{n:['position.x','position.y'] for n in ('a','b')}}
  with mock.patch.object(rc,'load',return_value=self.contract):plan=rc._single_repair_plan(self.t,self.s,self.c,request,(check,self.paint,self.n))
  self.assertEqual(set(plan['translations']),{'a','b'});self.assertFalse(plan['before_check']['passed']);self.assertEqual(len(plan['check_set']),3)
  check['capture_required']=[{'object':'a','reason':'Actual moving member missing'}]
  with mock.patch.object(rc,'load',return_value=self.contract):
   with self.assertRaises(ValueError):rc._single_repair_plan(self.t,self.s,self.c,request,(check,self.paint,self.n))
 def test_frame_regression_remains_failed_without_aborting_whole_report(self):
  import unittest.mock as mock
  from core import frame_plan as f
  (self.t/'Scene.json').write_text('{}',encoding='utf-8');candidate={**self.c,'scene':{'path':'Scene.json'}}
  state={**self.s,'frame_plan':{'path':'contract.json'},'selected':{'pptx':{'path':'Earlier.pptx','sha256':'earlier'}}}
  failed={'status':'FAIL','passed':False,'items':[{'id':'frame-member','passed':False}]};prior={'status':'PASS','passed':True,'items':[{'id':'frame-member','passed':True}]}
  with mock.patch.object(f,'enabled',return_value=True),mock.patch.object(f,'_resource',side_effect=lambda task,ref:self.t/ref['path']),mock.patch.object(f,'load',return_value={}),mock.patch.object(f,'validate',return_value={}),mock.patch.object(f,'check',side_effect=[copy.deepcopy(failed),prior]):
   result=f.guard(self.t,state,candidate,report_only=True)
  self.assertFalse(result['passed']);self.assertFalse(result['eligible']);self.assertEqual(result['previously_satisfied_regressions'],['frame-member'])
  with mock.patch.object(f,'enabled',return_value=True),mock.patch.object(f,'_resource',side_effect=lambda task,ref:self.t/ref['path']),mock.patch.object(f,'load',return_value={}),mock.patch.object(f,'validate',return_value={}),mock.patch.object(f,'check',side_effect=[copy.deepcopy(failed),prior]):
   with self.assertRaisesRegex(ValueError,'regressed'):f.guard(self.t,state,candidate)
 def test_compressed_actual_render_cache_preserves_exact_bytes(self):
  import gzip,base64
  from core.native_task import unpack_before_render
  names=['Final.pdf','Final_preview.png','render.json','job-receipt.json','office-session.json'];files={}
  import hashlib
  for name in names:
   data=(name*30).encode();files[name]={'base64':base64.b64encode(gzip.compress(data,mtime=0)).decode(),'encoding':'gzip_base64','size_bytes':len(data),'sha256':hashlib.sha256(data).hexdigest()}
  unpack_before_render({'files':files},self.t)
  for name in names:self.assertEqual((self.t/name).read_bytes(),(name*30).encode())
  files[names[0]]['size_bytes']+=1
  with self.assertRaises(ValueError):unpack_before_render({'files':files},self.t)
 def test_actual_render_cache_resource_bound_rejects_before_office(self):
  import unittest.mock as mock
  from core import native_task as nt,native_capture_batch as captures
  folder=self.t/'bounded-cache';folder.mkdir()
  for name in ['Final.pdf','Final_preview.png','render.json','job-receipt.json','office-session.json']:(folder/name).write_text('1234')
  write_json(folder/'render.json',{'backend':'PowerPoint','pages':[{}]})
  with mock.patch.object(captures,'verify_actual_render',return_value=True),mock.patch.object(nt,'MAX_BEFORE_RENDER_BYTES',5):
   with self.assertRaisesRegex(ValueError,'before Office execution'):nt.bind_before_render(self.t/'Input.pptx',folder/'render.json')
 def test_native_emf_preview_requires_exact_source_hash(self):
  import shutil
  from core.native_task import verify_render
  directory=self.t/'emf-job';(directory/'before').mkdir(parents=True);stage=self.t/'emf-stage';(stage/'result/before').mkdir(parents=True)
  Image.new('RGB',(240,160),'white').save(directory/'before/Final_preview.png');shutil.copyfile(directory/'before/Final_preview.png',stage/'result/before/Final_preview.png')
  (directory/'After.emf').write_bytes(b'changed native metafile')
  raw={'before_png_sha256':digest(directory/'before/Final_preview.png'),'after_emf_sha256':'different original bytes'}
  self.assertFalse(verify_render({'preview_width':240,'canvas':{'width':120,'height':80}},raw,stage,directory))
 def test_owner_render_requires_every_repeated_math_span_and_text_action(self):
  from core.native_task import owner_render_coverage
  plan={'actions':[{'id':'mixed','action':'math'},{'id':'mixed','action':'math'},{'id':'prose','action':'text'}]}
  rows=[{'object':'mixed','operation_indices':[0,1]},{'object':'prose','operation_indices':[2]}]
  self.assertTrue(owner_render_coverage(plan,rows));self.assertFalse(owner_render_coverage(plan,rows[:1]))
  self.assertFalse(owner_render_coverage(plan,[rows[0],rows[0],rows[1]]))
  partial=copy.deepcopy(rows);partial[0]['operation_indices']=[0];self.assertFalse(owner_render_coverage(plan,partial))
  plan['actions'][2]['action']='connector';self.assertFalse(owner_render_coverage(plan,rows))
 def test_owner_render_changed_native_metafile_cannot_pass(self):
  import shutil
  from core.native_task import verify_render
  directory=self.t/'owner-job';(directory/'before').mkdir(parents=True);stage=self.t/'owner-stage';(stage/'result/before').mkdir(parents=True)
  Image.new('RGB',(240,160),'white').save(directory/'before/Final_preview.png');shutil.copyfile(directory/'before/Final_preview.png',stage/'result/before/Final_preview.png')
  (directory/'changed.emf').write_bytes(b'changed bytes')
  owner={'object':'math','operation_indices':[0],'emf':'changed.emf','preview':'unknown.png','canvas_px':[40,20],'emf_sha256':'old hash','preview_sha256':'old hash'}
  raw={'before_png_sha256':digest(directory/'before/Final_preview.png'),'after_owner_renders':[owner]}
  self.assertFalse(verify_render({'actions':[{'id':'math','action':'math'}],'preview_width':240,'canvas':{'width':120,'height':80}},raw,stage,directory))
 def test_pdf_preview_route_replays_pixels_and_rejects_tampering(self):
  import pymupdf
  from core.render import rasterize_pdf
  from core.native_task import verify_pdf_preview
  pdf=self.t/'Actual-layout.pdf';doc=pymupdf.open();page=doc.new_page(width=120,height=80);page.draw_rect((20,20,70,55),color=(0,0,1),fill=(0,0,1));doc.save(pdf);doc.close()
  rasterize_pdf(pdf,self.t/'view',240,1,160);preview=self.t/'view/Final_preview.png'
  self.assertTrue(verify_pdf_preview(pdf,preview,240,160))
  Image.new('RGB',(240,160),'white').save(preview)
  with self.assertRaises(ValueError):verify_pdf_preview(pdf,preview,240,160)
 def test_missing_actual_pdf_cannot_generate_native_preview(self):
  from core.native_task import complete_render
  folder=self.t/'missing-native-render';folder.mkdir();write_json(folder/'operations.json',{'input_sha256':'no actual export'})
  plan=self.t/'Plan.json';write_json(plan,{'canvas':{'width':120,'height':80},'preview_width':240})
  with self.assertRaises(ValueError):complete_render(plan,folder)
 def test_failed_partial_png_is_preserved_before_pdf_recovery(self):
  import pymupdf,shutil
  from core.native_task import complete_render
  folder=self.t/'partial-native-render';(folder/'before').mkdir(parents=True)
  pdf=folder/'before/Final.pdf';doc=pymupdf.open();doc.new_page(width=120,height=80);doc.save(pdf);doc.close();shutil.copyfile(pdf,folder/'After.pdf')
  partial=folder/'before/Final_preview.png';Image.new('RGB',(240,160),'red').save(partial);old=digest(partial)
  plan=self.t/'Plan.json';write_json(plan,{'canvas':{'width':120,'height':80},'preview_width':240})
  write_json(folder/'operations.json',{'before_pdf_sha256':digest(pdf),'after_pdf_sha256':digest(folder/'After.pdf')})
  result=complete_render(plan,folder);preserved=result['rows'][0]['failed_direct_PNG']
  self.assertEqual(digest(folder/preserved['path']),old);self.assertNotEqual(digest(partial),old)

class Recorder(unittest.TextTestResult):
 def addSuccess(self,test):super().addSuccess(test);RESULTS.append({'test':test.id(),'status':'PASS'})
 def addFailure(self,test,err):super().addFailure(test,err);RESULTS.append({'test':test.id(),'status':'FAIL','detail':self._exc_info_to_string(err,test)})
 def addError(self,test,err):super().addError(test,err);RESULTS.append({'test':test.id(),'status':'ERROR','detail':self._exc_info_to_string(err,test)})
def run(output):
 global D,RESULTS
 D=Path(output).resolve();D.mkdir(parents=True,exist_ok=True);RESULTS=[]
 result=unittest.TextTestRunner(verbosity=1,resultclass=Recorder).run(unittest.defaultTestLoader.loadTestsFromTestCase(Tests))
 from core.execution_contract import _implementation
 report={'status':'PASS' if result.wasSuccessful() else 'FAIL','tests':RESULTS,'implementation':_implementation(),
  'test_source_sha256':digest(Path(__file__)),'scope':'Current executable deterministic assertions; actual Office behavior and whole artwork are separately verified'}
 write_json(D/'Functional-regressions.json',report);return report
if __name__=='__main__':
 import argparse
 p=argparse.ArgumentParser();p.add_argument('--out',type=Path,required=True);args=p.parse_args()
 sys.exit(0 if run(args.out)['status']=='PASS' else 1)
