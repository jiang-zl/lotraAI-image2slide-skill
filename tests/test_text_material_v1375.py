"""Saved native text direction, explicit logical breaks and editable effects."""
from pathlib import Path
import sys
import tempfile
import unittest
from zipfile import ZipFile
from lxml import etree as E
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'skills/lotra-image2slide-skill/scripts'))
from core import text_layout,native_core_checks
from core.build import text_body,shape_style
from core.mathml import NS


class LogicalText(unittest.TestCase):
    def shape(self,content):
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(2),Inches(2))
        text_body(shape,content,{'font':'Arial','font_size':18})
        return prs,shape

    def test_vertical_native_word_keeps_one_logical_literal(self):
        content={'text':'组织架构','logical_text':'组织架构','text_layout':{
            'direction':'east_asian_vertical','source_evidence':'Four original Chinese characters in one vertical column'}}
        prs,shape=self.shape(content)
        with tempfile.TemporaryDirectory() as folder:
            file=Path(folder)/'vertical.pptx';prs.save(file)
            reopened=Presentation(file).slides[0].shapes[0]
            self.assertEqual(text_layout.saved_text(reopened._element,content),'组织架构')
            self.assertEqual(reopened._element.find('p:txBody/a:bodyPr',NS).get('vert'),'eaVert')
            self.assertEqual(len(reopened._element.findall('p:txBody/a:p',NS)),1)

    def test_only_declared_soft_boundaries_are_removed(self):
        content={'paragraphs':[[{'text':c}] for c in '组织架构'],'logical_text':'组织架构',
                 'text_layout':{'direction':'horizontal','soft_breaks':[1,2,3],'source_evidence':'Stacked visual glyphs; no semantic paragraph breaks'}}
        _,shape=self.shape(content)
        self.assertEqual(text_layout.saved_text(shape._element,content),'组织架构')
        self.assertEqual(text_layout.logical('AB\nCD',{}),'AB\nCD')
        with self.assertRaisesRegex(ValueError,'SOFT_BREAK_MAPPING_REQUIRED'):
            text_layout.logical('ABCD',{'soft_breaks':[2]})

    def test_literal_changes_cannot_hide_as_layout_normalization(self):
        content={'text':'组\n织\n架\n错','logical_text':'组织架构',
                 'text_layout':{'soft_breaks':[1,2,3],'source_evidence':'Original vertical word'}}
        with self.assertRaisesRegex(ValueError,'LOGICAL_TEXT_MISMATCH'):self.shape(content)

    def test_hard_paragraph_breaks_are_preserved(self):
        content={'text':'项目一\n项目二','logical_text':'项目一\n项目二',
                 'text_layout':{'direction':'horizontal','source_evidence':'Two actual original paragraphs'}}
        _,shape=self.shape(content)
        self.assertEqual(text_layout.saved_text(shape._element,content),'项目一\n项目二')


class NativeMaterials(unittest.TestCase):
    def test_saved_shadow_parameters_are_independently_checked(self):
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        shape=slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,Inches(1),Inches(1),Inches(3),Inches(.6))
        style={'fill':'EAE7E0','stroke':None,'shadow_effect':{'colour':'263342','opacity':.28,
                'blur_pt':1.5,'offset_pt':[2,3],'rotate_with_shape':False},'soft_edge_pt':.2}
        shape_style(shape,style,'shape')
        obj={'id':'bar','render_type':'shape','style':style,'content':{}}
        self.assertFalse([r for r in native_core_checks.style_checks(obj,shape._element) if r['status']=='FAIL'])
        with tempfile.TemporaryDirectory() as folder:
            file=Path(folder)/'material.pptx';prs.save(file)
            reopened=Presentation(file).slides[0].shapes[0]
            self.assertFalse([r for r in native_core_checks.style_checks(obj,reopened._element) if r['status']=='FAIL'])
            shadow=reopened._element.find('p:spPr/a:effectLst/a:outerShdw',NS);shadow.set('dist','0')
            self.assertIn('native_shadow_geometry',[r['check_id'] for r in native_core_checks.style_checks(obj,reopened._element) if r['status']=='FAIL'])

    def test_unbounded_effects_cannot_enter_the_writer(self):
        from core.material_effects import validate
        with self.assertRaises(ValueError):validate({'shadow_effect':{'colour':'000000','opacity':1,'blur_pt':999,'offset_pt':[0,0]}})
        with self.assertRaises(ValueError):validate({'soft_edge_pt':-1})


if __name__=='__main__':unittest.main()
