"""Source offsets, whole painted aggregates and corrupted native alignment."""
import copy
from pathlib import Path
import sys
import unittest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'skills/lotra-image2slide-skill/scripts'))
from core import native_core_checks as native, text_appearance as appearance, source_typography, source_preflight
from core.build import text_body
from core.mathml import NS


class Centering(unittest.TestCase):
    def text(self, content=None):
        content=content or {'text':'AB\nCD'}
        obj={'id':'caption','render_type':'text','content':content,
             'style':{'font':'Arial','font_size':18,'align':'center','valign':'middle','color':'000000'}}
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(3),Inches(1))
        text_body(shape,content,obj['style'])
        return obj,shape

    def failures(self,obj,shape):
        return {r['check_id'] for r in native.style_checks(obj,shape._element) if r['status']=='FAIL'}

    def test_both_native_axes_and_every_paragraph_are_checked_before_render(self):
        obj,shape=self.text();self.assertEqual(self.failures(obj,shape),set())
        shape._element.findall('p:txBody/a:p/a:pPr',NS)[1].set('algn','l')
        shape._element.find('p:txBody/a:bodyPr',NS).set('anchor','t')
        self.assertEqual(self.failures(obj,shape),{'p1.horizontal_alignment','text_vertical_alignment'})

    def test_asymmetric_native_padding_cannot_hide_under_center_flag(self):
        obj,shape=self.text();shape._element.find('p:txBody/a:bodyPr',NS).set('lIns','12700')
        self.assertIn('text_frame_margins',self.failures(obj,shape))
        obj,shape=self.text();shape._element.find('p:txBody/a:p/a:pPr',NS).set('marL','12700')
        self.assertIn('p0.horizontal_insets',self.failures(obj,shape))

    def test_source_owned_paragraph_offset_is_preserved(self):
        obj,shape=self.text({'text':'AB','paragraph_layout':[{'left_indent_pt':2,'right_indent_pt':0}]})
        self.assertEqual(self.failures(obj,shape),set())

    def fixture(self, kind='text'):
        image=Image.new('RGB',(200,130),'white');draw=ImageDraw.Draw(image)
        # An asymmetric complete aggregate, including a high small mark. Its
        # intentional source offset must survive without snapping to the ROI.
        draw.rectangle((70,50,99,64),fill='black');draw.rectangle((101,42,105,46),fill='black')
        feature={'id':'caption','kind':kind,'roi':[50,30,100,60],
                 'text_appearance':{'source_evidence':'Independent source aggregate including high mark',
                    'isolation':'All black content in the fixed white source window','colours':['000000'],
                    'colour_tolerance':.1,'limits':{'width':1,'height':1,'baseline':1,'spacing':1}}}
        return image,feature

    def shifted(self,source,dx,dy):
        image=Image.new('RGB',source.size,'white');image.paste(source,(dx,dy));return image

    def test_horizontal_ink_shift_fails_even_when_dimensions_and_baseline_pass(self):
        source,feature=self.fixture();result=appearance.compare(source,self.shifted(source,3,0),feature)
        self.assertFalse(result['passed']);checks={c['field']:c for c in result['checks']}
        self.assertTrue(all(checks[k]['passed'] for k in ('width','height','baseline','spacing')))
        self.assertFalse(checks['ink_center_x']['passed']);self.assertEqual(checks['ink_center_x']['value'],3)
        self.assertFalse(result['candidate_alignment'])

    def test_both_painted_axes_are_measured_for_complete_formula(self):
        source,feature=self.fixture('math');result=appearance.compare(source,self.shifted(source,3,2),feature)
        checks={c['field']:c for c in result['checks']}
        self.assertFalse(checks['ink_center_x']['passed']);self.assertFalse(checks['ink_center_y']['passed'])
        self.assertEqual(result['source']['ink_center_y'],53.5)

    def test_original_offset_is_accepted_but_geometric_center_snap_is_rejected(self):
        source,feature=self.fixture();self.assertTrue(appearance.compare(source,source,feature)['passed'])
        self.assertNotEqual(appearance.ink(source,feature)['ink_center_x'],feature['roi'][0]+feature['roi'][2]/2)
        self.assertFalse(appearance.compare(source,self.shifted(source,5,0),feature)['passed'])

    def test_alignment_intent_change_requires_source_rebind_before_native_write(self):
        obj,_=self.text();obj['bbox']=[0,0,100,30]
        frozen={'objects':[obj]};scene=copy.deepcopy(frozen);scene['objects'][0]['style']['valign']='top'
        rows,errors=source_typography._alignment(scene,frozen,{'caption'})
        self.assertIn('ALIGNMENT_REBIND_REQUIRED',errors[0]['reason'])
        self.assertEqual(rows[0]['status'],'SOURCE_INTENT_MISMATCH')

    def test_new_math_source_contract_cannot_omit_aggregate_ink_observation(self):
        source,feature=self.fixture('math');feature.pop('text_appearance')
        observation={'source_hash':'a'*64,'features':[feature],'units':[],'layout_review':{}}
        result=source_preflight.evaluate(source,observation,include_math=True)
        self.assertIn('TEXT_APPEARANCE_SOURCE_REQUIRED',[row['reason'].split(':')[0] for row in result['issues']])


if __name__=='__main__':unittest.main()
