"""Distinct typographic/compatible names and actual embedded style slots."""
import copy
import importlib
from pathlib import Path
import sys
import unittest
from unittest.mock import patch
from zipfile import ZipFile
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'skills/lotra-image2slide-skill/scripts'))
from core import font_resources as resources,font_embedding,font_retention,font_contract as fonts
from core import font_collection,native_core_checks,font_resolution
from core.build import text_body
from core.common import digest,identity


class CompatibleNames(unittest.TestCase):
    def setUp(self):
        self.fixture=importlib.import_module('test_font_workflow_v1373').Fonts()
        self.fixture.setUp();self.addCleanup(self.fixture.doCleanups)

    def face(self,name,weight,bold):
        value=self.fixture.face(name);path=Path(value['path'])
        with TTFont(path) as font:
            for platform,encoding,language in ((3,1,0x409),(1,0,0)):
                font['name'].setName('Shared Typographic Family',16,platform,encoding,language)
                font['name'].setName('Narrow Bold' if bold else 'Black',17,platform,encoding,language)
                font['name'].setName('Bold' if bold else 'Regular',2,platform,encoding,language)
            font['OS/2'].usWeightClass=weight
            font['OS/2'].fsSelection=32 if bold else 64
            font['head'].macStyle=1 if bold else 0
            font.save(path)
        current=resources.describe(path,0,digest(path));current['id']=identity({
            key:current[key] for key in ('file_sha256','face_index','family')})
        return current

    def test_shared_typographic_family_does_not_merge_real_typefaces(self):
        narrow=self.face('Fixture Narrow',700,True);black=self.face('Fixture Black',900,False)
        self.assertEqual(narrow['typographic_family'],black['typographic_family'])
        self.assertNotEqual(narrow['family'],black['family'])
        self.assertFalse(resources.matches(narrow,'Shared Typographic Family'))
        self.assertEqual(resources.embedding_slot(narrow),'bold')
        self.assertEqual(resources.embedding_slot(black),'regular')

    def test_distinct_faces_embed_and_reopen_without_slot_overwrite(self):
        narrow=self.face('Fixture Narrow',700,True);black=self.face('Fixture Black',900,False)
        bindings=[{'family':face['family'],'resource':face} for face in (narrow,black)]
        file=self.fixture.package('separate-fonts.pptx',families=[f['family'] for f in (narrow,black)])
        result=font_embedding.embed(file,bindings)
        with ZipFile(file) as archive:
            checked=font_retention.verify_package(archive,bindings)
        self.assertEqual(len(checked),2)
        self.assertEqual({row['slot'] for row in checked},{'bold','regular'})

    def test_slot_conflict_fails_before_touching_presentation_bytes(self):
        first=self.face('Fixture Narrow',700,True);second=self.face('Fixture Alternate',700,True)
        second['family']=first['family']
        bindings=[{'resource':face,'family':face['family']} for face in (first,second)]
        file=self.fixture.package();before=digest(file)
        with self.assertRaisesRegex(ValueError,'FONT_SLOT_CONFLICT'):
            font_embedding.embed(file,bindings)
        self.assertEqual(digest(file),before)

    def test_heavy_regular_face_preserves_emphasis_without_second_bold(self):
        black=self.face('Fixture Black',900,False)
        scene={'revision':0,'source':{'sha256':'original'},'font_workflow_policy':font_resolution.POLICY,
               'objects':[{'id':'lead','render_type':'text','semantic_type':'label','bbox':[0,0,100,30],
                           'style':{'font':black['family'],'bold':True,'font_size':18},
                           'content':{'runs':[{'text':'AB','color':'FF0000'}]}}]}
        request=fonts.requests(scene)[0]
        binding={'request':request,'request_hash':font_resolution.request_key(request),
                 'family':black['family'],'resource':black,'origin':'raster_inference','status':'approximate'}
        with patch.object(font_collection,'assert_language'):
            compiled=fonts.compile_scene(copy.deepcopy(scene),{'font_contract':{'bindings':{request['id']:binding}}})
        obj=compiled['objects'][0]
        self.assertTrue(obj['style']['bold'])
        self.assertFalse(obj['content']['runs'][0]['native_bold'])
        self.assertTrue(compiled['font_plan']['source_spans']['lead']['paragraphs'][0]['spans'][0]['style']['bold'])
        self.assertFalse(compiled['font_plan']['native_spans']['lead']['paragraphs'][0]['spans'][0]['style']['bold'])
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(3),Inches(1));shape.name='lead'
        text_body(shape,obj['content'],obj['style']);file=self.fixture.root/'black.pptx';prs.save(file)
        self.assertEqual([row for row in native_core_checks.style_checks(obj,shape._element) if row['status']=='FAIL'],[])
        fonts.guard_saved(compiled,file)

    def test_mixed_script_physical_emphasis_keeps_logical_intervals(self):
        black=self.face('Fixture Black',900,False);bold=self.face('Fixture Chinese',700,True)
        scene={'revision':0,'source':{'sha256':'original'},'objects':[
            {'id':'mixed','render_type':'text','semantic_type':'label','bbox':[0,0,100,30],
             'style':{'bold':True,'font_size':18},'content':{'runs':[
                 {'text':'AB中文C','color':'FF0000'},{'text':'D','color':'0000FF'}]}}]}
        bindings={}
        for request in fonts.requests(scene):
            face=black if request['role']=='latin' else bold
            bindings[request['id']]={'request':request,'request_hash':font_resolution.request_key(request),
                 'family':face['family'],'resource':face,'origin':'raster_inference','status':'approximate'}
        before=font_resolution.source_spans(scene)
        with patch.object(font_collection,'assert_language'):
            compiled=fonts.compile_scene(copy.deepcopy(scene),{'font_contract':{'bindings':bindings}})
        self.assertEqual(before,compiled['font_plan']['source_spans'])
        obj=compiled['objects'][0];prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(3),Inches(1));shape.name='mixed'
        text_body(shape,obj['content'],obj['style']);file=self.fixture.root/'mixed.pptx';prs.save(file)
        self.assertEqual([r.text for r in shape.text_frame.paragraphs[0].runs],['AB','中文','C','D'])
        self.assertEqual([r.font.bold for r in shape.text_frame.paragraphs[0].runs],[False,True,False,False])
        self.assertEqual([r for r in native_core_checks.style_checks(obj,shape._element) if r['status']=='FAIL'],[])
        fonts.guard_saved(compiled,file)
        shape.text_frame.paragraphs[0].runs[1].font.bold=False;prs.save(file)
        with self.assertRaisesRegex(ValueError,'FONT_TEXT_STYLE_CHANGE'):fonts.guard_saved(compiled,file)

    def test_compiled_mixed_table_retains_source_emphasis_and_valid_schema(self):
        from core import tables
        black=self.face('Fixture Black',900,False);bold=self.face('Fixture Chinese',700,True)
        scene={'revision':0,'source':{'sha256':'original'},'objects':[
            {'id':'grid','render_type':'table','semantic_type':'table','bbox':[0,0,100,30],
             'style':{'bold':True,'font_size':18},'content':{'rows':[['AB中文C','AB']]}}]}
        bindings={}
        for request in fonts.requests(scene):
            face=black if request['role']=='latin' else bold
            bindings[request['id']]={'request':request,'request_hash':font_resolution.request_key(request),
                'family':face['family'],'resource':face,'origin':'raster_inference','status':'approximate'}
        with patch.object(font_collection,'assert_language'):
            compiled=fonts.compile_scene(copy.deepcopy(scene),{'font_contract':{'bindings':bindings}})
            # The builder revalidates normalized content after compiling fonts.
            tables.normalize_scene(compiled)
        cells=compiled['objects'][0]['content']['cell_styles'][0]
        self.assertEqual(cells[0]['native_script_styles'],{'latin':{'bold':False,'italic':False},'ea':{'bold':True,'italic':False}})
        self.assertFalse(cells[1]['native_bold'])
        self.assertTrue(compiled['objects'][0]['style']['bold'])

    def test_table_physical_style_fields_reject_invalid_values(self):
        from core.tables import validate
        for bad in ({'native_bold':'false'},{'native_script_styles':{'latin':{'bold':False}}},
                    {'native_script_styles':{'other':{'bold':False,'italic':False}}}):
            with self.assertRaises(ValueError):validate({'rows':[['A']],'cell_styles':[[bad]]})


if __name__=='__main__':unittest.main()
