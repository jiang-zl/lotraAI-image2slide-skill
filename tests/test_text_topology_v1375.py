"""Real saved native text-body changes keep leaf authority and unrelated bytes."""
import copy,sys,tempfile,unittest
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from pptx import Presentation
from pptx.util import Pt,Inches
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'skills/lotra-image2slide-skill/scripts'))
from core.native_commit import merge_native
from core.mathml import NS

class TextTopology(unittest.TestCase):
    def setUp(self):
        base=ROOT/'.lotraAI-build/topology-tests';base.mkdir(parents=True,exist_ok=True)
        self.tmp=tempfile.TemporaryDirectory(dir=base);self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name)
        for name,lines,size in (('before',['组织架构'],25),('donor',list('组织架构'),14)):
            deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(1),Inches(3));shape.name='label'
            shape.text_frame.clear()
            for i,line in enumerate(lines):
                p=shape.text_frame.paragraphs[0] if i==0 else shape.text_frame.add_paragraph();p.line_spacing=Pt(13.29)
                r=p.add_run();r.text=line;r.font.name='Arial';r.font.size=Pt(size)
            if name=='before':
                neighbour=slide.shapes.add_textbox(Inches(4),Inches(1),Inches(2),Inches(1));neighbour.name='neighbour';neighbour.text='Preserve exact neighbour'
            deck.save(self.root/(name+'.pptx'))

    def native(self,file,name):
        with ZipFile(file) as archive:tree=E.fromstring(archive.read('ppt/slides/slide1.xml'))
        return next(node for node in tree.findall('p:cSld/p:spTree/p:sp',NS) if node.find('p:nvSpPr/p:cNvPr',NS).get('name')==name)

    def test_explicit_body_before_style_leaves_changes_topology_without_owner_rebuild(self):
        before=self.root/'before.pptx';output=self.root/'output.pptx'
        original=self.native(before,'label');neighbour=E.tostring(self.native(before,'neighbour'))
        merge_native(before,self.root/'donor.pptx',{'label':['style.font_size','style.line_spacing_pt','text_body','text_content']},output)
        saved=self.native(output,'label')
        self.assertEqual([n.text for n in saved.findall('.//a:t',NS)],list('组织架构'))
        self.assertEqual(saved.find('p:nvSpPr/p:cNvPr',NS).attrib,original.find('p:nvSpPr/p:cNvPr',NS).attrib)
        self.assertEqual(E.tostring(saved.find('p:spPr',NS)),E.tostring(original.find('p:spPr',NS)))
        self.assertTrue(all(n.get('sz')=='1400' for n in saved.findall('.//a:rPr',NS)))
        self.assertEqual(E.tostring(self.native(output,'neighbour')),neighbour)

    def test_font_leaf_without_explicit_body_still_rejects_topology_change(self):
        before=self.root/'before.pptx';unchanged=before.read_bytes();output=self.root/'rejected.pptx'
        with self.assertRaisesRegex(ValueError,'unchanged native run topology'):
            merge_native(before,self.root/'donor.pptx',{'label':['style.font_size']},output)
        self.assertEqual(before.read_bytes(),unchanged);self.assertFalse(output.exists())

if __name__=='__main__':unittest.main()
