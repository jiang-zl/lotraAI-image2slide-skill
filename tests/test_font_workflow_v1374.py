"""Cross-layer font contracts: hints, slots, colour intervals and finite steps."""
import copy
import importlib
from pathlib import Path
import sys
import unittest
from unittest.mock import patch
from lxml import etree as E
from pptx import Presentation
from pptx.util import Inches

REPO=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(REPO/'skills/lotra-image2slide-skill/scripts'))
from core import font_contract as fc,font_collection as collection,font_resolution as resolution
from core import font_selection as selection,native_core_checks,component_actions,packet_compaction
from core.common import identity,digest,write_json
from core.build import text_body


class FontContracts(unittest.TestCase):
    def setUp(self):
        fixture=importlib.import_module('test_font_workflow_v1373').Fonts()
        fixture.setUp();self.addCleanup(fixture.doCleanups)
        self.fixture=fixture;self.root=fixture.root;self.task=fixture.task

    def scene(self,runs,modern=False):
        scene={'revision':1,'source':{'sha256':identity('independent fixture original')},'objects':[{
            'id':'lead','semantic_type':'label','render_type':'text','bbox':[0,0,100,30],
            'style':{'font':self.fixture.old['family'],'font_size':20,'color':'333333'},'content':{'runs':runs}}]}
        if modern:scene['font_workflow_policy']=resolution.POLICY
        return scene

    def plan(self,scene,choose=None):
        bindings={}
        for request in fc.requests(scene):
            face=choose(request) if choose else self.fixture.new
            bindings[request['id']]={'request':request,'request_hash':resolution.request_key(request),
                'family':face['family'],'resource':face,'status':'approximate','origin':'raster_inference'}
        with patch.object(collection,'assert_language'):
            return fc.compile_scene(scene,{'font_contract':{'bindings':bindings}})

    def emit(self,scene):
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1));shape.name='lead'
        obj=scene['objects'][0];text_body(shape,obj['content'],obj['style'])
        file=self.root/'output.pptx';prs.save(file)
        failures=[row for row in native_core_checks.style_checks(obj,shape._element) if row['status']=='FAIL']
        return prs,shape,file,failures

    def test_ea_binding_and_unrelated_hint_have_consistent_three_slots(self):
        planned=self.plan(self.scene([{'text':'中文'}]));_,_,file,failures=self.emit(planned)
        self.assertEqual(failures,[])
        self.assertEqual(fc.guard_saved(planned,file)['status'],'SOURCE_FONT_BINDINGS_SERIALIZED')
        self.assertEqual(set(planned['objects'][0]['content']['runs'][0]['font_faces']),set(resolution.SLOTS))

    def test_two_colours_and_true_two_families_remain_separate(self):
        planned=self.plan(self.scene([{'text':'中文','color':'FF0000'},{'text':'甲乙','color':'333333'}]),
                          lambda row:self.fixture.new if row['location']=='p0.r0' else self.fixture.old)
        _,_,file,failures=self.emit(planned)
        self.assertEqual(failures,[]);fc.guard_saved(planned,file)
        self.assertEqual([r['color'] for r in planned['objects'][0]['content']['runs']],['FF0000','333333'])

    def test_mixed_scripts_in_one_run_have_complete_plan(self):
        planned=self.plan(self.scene([{'text':'中文AB'}]),lambda row:self.fixture.old if row['role']=='latin' else self.fixture.new)
        _,_,file,failures=self.emit(planned);self.assertEqual(failures,[]);fc.guard_saved(planned,file)

    def test_inline_real_office_math_retains_ordinary_colour_intervals(self):
        planned=self.plan(self.scene([{'text':'AB','color':'FF0000'},{'ast':{'op':'token','text':'a'}}]))
        _,shape,file,failures=self.emit(planned)
        self.assertEqual(failures,[]);fc.guard_saved(planned,file)
        self.assertTrue(shape._element.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'))

    def build_table(self,font_size=18):
        from PIL import Image
        from core.build import build
        from core.workflow_guard import developer_writer
        Image.new('RGB',(200,120),'white').save(self.task/'source/original.png')
        write_json(self.task/'source/evidence.json',{});write_json(self.task/'source/profile.json',{})
        scene={'schema_version':'0.1','revision':0,'canvas':{'width':200,'height':120,'unit':'px','target_width_inches':200/96},
            'relations':[],'prototypes':[],'objects':[{'id':'lead','render_type':'table','semantic_type':'data_table','bbox':[20,20,100,30],
                'style':{'font':self.fixture.old['family'],'font_size':font_size,'color':'333333'},
                'content':{'rows':[['中文','AB']],'cell_styles':[[{'color':'FF0000'},{}]]},
                'versions':dict.fromkeys(('content','structure','style','placement'),1),
                'locks':dict.fromkeys(('content','structure','style','placement'),False),'evidence_status':'observed'}]}
        for field,name in [('source','original.png'),('evidence','evidence.json'),('profile','profile.json')]:
            scene[field]={'path':'source/'+name,'sha256':digest(self.task/'source'/name)}
        planned=self.plan(scene,lambda row:self.fixture.new if row['role']=='ea' else self.fixture.old)
        folder=self.root/'.lotraAI-build/developer-drafts/table';folder.mkdir(parents=True)
        file=folder/'Draft_editable.pptx'
        with patch.object(collection,'assert_language'),developer_writer(self.root):build(planned,self.task,file)
        return planned,file

    def test_table_cells_share_complete_slot_resolution_and_saved_checks(self):
        planned,file=self.build_table();prs=Presentation(file);shape=prs.slides[0].shapes[0]
        failures=[row for row in native_core_checks.style_checks(planned['objects'][0],shape._element) if row['status']=='FAIL']
        self.assertEqual(failures,[]);fc.guard_saved(planned,file)

    def test_fractional_text_and_math_paragraph_sizes_survive_saved_native_units(self):
        from zipfile import ZipFile
        from core.mathml import NS
        for size,centipoints in ((18.15,1815),(20.15,2015),(12.34,1234),(18.155,1816)):
            with self.subTest(size=size):
                scene=self.scene([{'text':'中文','color':'FF0000'},{'ast':{'op':'token','text':'a'}}])
                scene['objects'][0]['style']['font_size']=size
                planned=self.plan(scene);_,_,file,failures=self.emit(planned)
                self.assertEqual(failures,[])
                with ZipFile(file) as z:root=E.fromstring(z.read('ppt/slides/slide1.xml'))
                sizes=[int(n.get('sz')) for n in root.findall('.//a:rPr',NS)+root.findall('.//a:defRPr',NS)+root.findall('.//a:endParaRPr',NS) if n.get('sz')]
                self.assertGreaterEqual(len(sizes),3)
                self.assertEqual(set(sizes),{centipoints})

    def test_fractional_table_size_survives_saved_native_units(self):
        from zipfile import ZipFile
        from core.mathml import NS
        planned,file=self.build_table(18.15)
        with ZipFile(file) as z:root=E.fromstring(z.read('ppt/slides/slide1.xml'))
        sizes=[int(n.get('sz')) for n in root.findall('.//a:defRPr',NS) if n.get('sz')]
        self.assertEqual(sizes,[1815,1815])
        self.assertEqual([row for row in native_core_checks.style_checks(planned['objects'][0],Presentation(file).slides[0].shapes[0]._element) if row['status']=='FAIL'],[])

    def test_table_cell_colour_tampering_is_rejected(self):
        from pptx.dml.color import RGBColor
        planned,file=self.build_table();prs=Presentation(file)
        prs.slides[0].shapes[0].table.cell(0,0).text_frame.paragraphs[0].font.color.rgb=RGBColor.from_string('333333')
        prs.save(file)
        with self.assertRaisesRegex(ValueError,'FONT_TEXT_STYLE_CHANGE'):fc.guard_saved(planned,file)

    def test_separate_script_runs_do_not_shallow_replace_inherited_faces(self):
        planned=self.plan(self.scene([{'text':'中文','color':'FF0000'},{'text':'AB'}]),
                          lambda row:self.fixture.old if row['role']=='latin' else self.fixture.new)
        _,_,file,failures=self.emit(planned);self.assertEqual(failures,[]);fc.guard_saved(planned,file)

    def test_declared_slot_hints_are_retained_without_mutating_input(self):
        original=self.scene([{'text':'中文','font_faces':dict.fromkeys(resolution.SLOTS,self.fixture.old['family'])}])
        before=copy.deepcopy(original);planned=self.plan(original)
        self.assertEqual(original,before)
        self.assertIn(dict.fromkeys(resolution.SLOTS,self.fixture.old['family']),planned['font_plan']['declared_faces'].values())
        with patch.object(collection,'assert_language'):
            self.assertEqual(planned,fc.compile_scene(planned))

    def test_hint_only_change_preserves_modern_binding(self):
        planned=self.plan(self.scene([{'text':'中文AB'}],True))
        changed=copy.deepcopy(planned);changed['objects'][0]['style']['font']='Another hint'
        for request in fc.requests(changed):
            self.assertTrue(fc._request_match(planned['font_plan']['bindings'][request['id']],request))

    def test_exact_family_intent_change_requires_a_new_binding(self):
        original=self.scene([{'text':'中文'}],True);planned=self.plan(original)
        original['objects'][0]['style']['font_intent']={'required_families':{'ea':'Exact source family'},'revision':'observed-2'}
        request=fc.requests(original)[0]
        self.assertFalse(fc._request_match(planned['font_plan']['bindings'][request['id']],request))

    def test_raster_intent_revision_requires_repreparation_without_claiming_exact_identity(self):
        original=self.scene([{'text':'中文'}],True);planned=self.plan(original)
        original['objects'][0]['style']['font_intent']={'revision':'source-observation-2'}
        request=fc.requests(original)[0]
        self.assertIsNone(request['required_family'])
        self.assertFalse(fc._request_match(planned['font_plan']['bindings'][request['id']],request))

    def test_invalid_intent_revision_is_rejected(self):
        scene=self.scene([{'text':'中文'}],True)
        for revision in ('',{},'x'*241):
            scene['objects'][0]['style']['font_intent']={'revision':revision}
            with self.assertRaisesRegex(ValueError,'bounded nonempty'):fc.requests(scene)

    def test_raster_revision_accepts_its_scoped_semantic_grant(self):
        scene=self.scene([{'text':'中文'}],True);planned=self.plan(scene)
        scene['objects'][0]['style']['font_intent']={'revision':'observed-2'}
        row=fc.requests(scene)[0]
        state={'font_contract':planned['font_plan'],'font_pending_candidate':{'allowed':{},
            'reason':{'font_revision_grants':[{'owner':'lead','field':['style','font_intent','revision']}]}}}
        result=fc.revision_details(self.task,state,scene,row)
        self.assertTrue(result['authorized']);self.assertEqual(result['prior_binding_hash'],identity(planned['font_plan']['bindings'][row['id']]))

    def test_real_submission_controls_and_four_batch_ids_fit_minimum_context(self):
        from core import agent_runtime as runtime
        fixture=importlib.import_module('test_packet_submission_v1371').PacketSubmission()
        fixture.setUp();self.addCleanup(fixture.doCleanups)
        state=fixture.fixture.state();state['objectives']['context_chars']=2500
        rows=[{'id':identity(i),'large_evidence':'x'*4000} for i in range(4)]
        extra=selection.batch([dict(row,owner='lead',location=str(i),role='ea',characters='中',literal_hash=identity(i),
            bold=False,italic=False,proposed_family='hint') for i,row in enumerate(rows)],'source_style',state)
        extra['requests']=rows
        with patch.object(runtime,'_preview',return_value=[]):
            packet=runtime._packet(fixture.task,state,'font_binding','source_fonts','Prepare source-supported fonts',extra=extra)
        self.assertEqual(packet['request_ids'],[row['id'] for row in rows])
        self.assertEqual(packet['action'],'submit');self.assertEqual(packet['required_step'],'source_style')

    def test_legacy_hint_identity_is_retained_until_explicit_migration(self):
        planned=self.plan(self.scene([{'text':'中文'}]));changed=copy.deepcopy(planned)
        changed['objects'][0]['style']['font']='Another hint'
        request=fc.requests(changed)[0]
        self.assertFalse(fc._request_match(planned['font_plan']['bindings'][request['id']],request))

    def test_scene_policy_flag_cannot_migrate_a_registered_legacy_task(self):
        scene=self.scene([{'text':'中文'}],True)
        with self.assertRaisesRegex(ValueError,'FONT_POLICY_MIGRATION_REQUIRED'):
            fc.requests(scene,{'task_id':'retained-old-task'})

    def test_saved_colour_loss_cannot_pass_font_binding_guard(self):
        planned=self.plan(self.scene([{'text':'中文','color':'FF0000'},{'text':'甲乙'}]))
        prs,shape,file,_=self.emit(planned)
        shape._element.find('.//{'+fc.A+'}srgbClr').set('val','333333');prs.save(file)
        with self.assertRaisesRegex(ValueError,'FONT_TEXT_STYLE_CHANGE'):fc.guard_saved(planned,file)

    def test_unused_slot_tampering_is_detected_independently(self):
        planned=self.plan(self.scene([{'text':'中文'}]));prs,shape,file,_=self.emit(planned)
        shape._element.find('.//{'+fc.A+'}cs').set('typeface','Wrong fallback');prs.save(file)
        with self.assertRaisesRegex(ValueError,'complete font plan'):fc.guard_saved(planned,file)

    def test_logical_colour_intervals_survive_equivalent_run_subdivision(self):
        before=self.scene([{'text':'中文','color':'FF0000'},{'text':'甲乙'}])
        after=self.scene([{'text':'中','color':'FF0000'},{'text':'文','color':'FF0000'},{'text':'甲乙'}])
        self.assertEqual(resolution.source_spans(before),resolution.source_spans(after))

    def test_regular_face_is_ineligible_for_unapproved_bold(self):
        request=fc.requests(self.scene([{'text':'中文','bold':True}],True))[0]
        face={**self.fixture.old,'language':'zh'}
        state={'source':{'sha256':identity('source')},'project':str(self.root),'capabilities':{}}
        with patch.object(collection,'prepare',return_value={'faces':{face['id']:face}}):
            self.assertEqual(collection.eligible(state,request),[])
        self.assertFalse(resolution.style_feasible(face,request))

    def test_synthetic_style_is_an_explicit_separate_plan(self):
        request=fc.requests(self.scene([{'text':'中文','bold':True}],True))[0];face=self.fixture.old
        self.assertFalse(resolution.style_feasible(face,request))
        self.assertTrue(resolution.style_feasible(face,request,synthetic=True))

    def test_required_style_step_cannot_be_skipped_by_a_driver_flag(self):
        request=fc.requests(self.scene([{'text':'中文'}],True))[0]
        state={'source':{'sha256':identity('source')},'project':str(self.root),'objectives':{},'sampled':True,'constrained':True}
        self.assertEqual(selection.required_step(state,request),'source_style')
        with self.assertRaisesRegex(ValueError,'FONT_STEP_REQUIRED: source_style'):
            selection.require_binding(state,request,self.fixture.old,'raster_inference',False)

    def test_failed_batch_has_no_partial_preparation_state(self):
        scene=self.scene([{'text':'中文'}],True);request=fc.requests(scene)[0]
        state={'source':scene['source'],'objectives':{},'font_workflow_policy':resolution.POLICY}
        packet={**selection.batch([request],'source_style',state),'packet_id':'current'}
        before=copy.deepcopy(state)
        with self.assertRaisesRegex(ValueError,'FONT_BATCH_SCOPE'):
            selection.receive(self.task,state,scene,packet,{'font_constraints':{'another-request':{}}})
        self.assertEqual(state,before)

    def test_batch_ids_and_required_step_survive_minimum_context(self):
        ids=[identity(i) for i in range(4)]
        packet={'packet_id':'current','kind':'font_binding','target':'source_fonts','source':identity('source'),
            'batch_id':identity('batch'),'request_ids':ids,'required_step':'source_style','font_batch_scope':{'requests':4},
            'requests':[{'id':key,'large_evidence':'x'*4000} for key in ids]}
        compact=packet_compaction.compact(self.task,packet,2500)
        self.assertEqual(compact['request_ids'],ids);self.assertEqual(compact['required_step'],'source_style')
        self.assertEqual(packet_compaction.restore(self.task,compact),packet)

    def test_same_source_group_cannot_adopt_different_families(self):
        scene=self.scene([{'text':'中文'},{'text':'甲乙'}],True);rows=fc.requests(scene)
        state={'source':scene['source'],'font_constraints':{row['id']:{'source_hash':scene['source']['sha256'],
            'request_hash':resolution.request_key(row),'group_id':'shared'} for row in rows}}
        bindings={row['id']:{'request':row,'family':family} for row,family in zip(rows,['One','Two'])}
        with self.assertRaisesRegex(ValueError,'FONT_GROUP_CONFLICT'):selection.check_group_bindings(state,bindings)

    def test_source_line_failure_has_a_legal_decision_action(self):
        row={'id':'lead','route':'fixed_parameters','owners':['lead']}
        contract={'components':[row],'policy':{'stage_order':['source','measurement','structure','fit','instance','assembly']}}
        ledger={'components':{'lead':{'status':'SOURCE_LINE_CONSTRAINT_REQUIRED'}}}
        action=component_actions.available(contract,ledger,{})[0]
        self.assertEqual(action['operation'],'source_or_measurement_decision');self.assertFalse(action['executable'])

    def test_unknown_component_status_is_an_explicit_diagnostic_stop(self):
        row={'id':'lead','route':'fixed_parameters','owners':['lead']}
        contract={'components':[row],'policy':{'stage_order':['source','measurement','structure','fit','instance','assembly']}}
        action=component_actions.available(contract,{'components':{'lead':{'status':'UNKNOWN_NEW_STATE'}}},{})[0]
        self.assertEqual(action['operation'],'diagnostic_stop');self.assertFalse(action['executable'])

    def test_all_known_status_route_pairs_have_finite_outcomes(self):
        for route in component_actions.ROUTE_STATES:
            for status in component_actions.KNOWN_STATUSES:
                with self.subTest(route=route,status=status):
                    row={'id':'lead','route':route,'owners':['lead'],'parameters':{},'fit':{'bindings':{}}}
                    record={'status':status,'repair_grant':[{'owner':'lead','field':['style','font']}]}
                    contract={'components':[row],'policy':{'stage_order':['source','measurement','structure','fit','instance','assembly']}}
                    actions=component_actions.available(contract,{'components':{'lead':record}},{})
                    if status in ('QUALIFIED','SPATIAL_REPAIR_WAITING'):
                        self.assertEqual(actions,[]);continue
                    self.assertEqual(len(actions),1)
                    if status not in component_actions.EVALUATION_STATES:
                        self.assertNotEqual(actions[0]['operation'],'evaluate')

    def test_saved_instance_confirmation_retains_its_qualification_route(self):
        row={'id':'lead','route':'instance','owners':['lead']}
        contract={'components':[row],'policy':{'stage_order':['source','measurement','structure','fit','instance','assembly']}}
        action=component_actions.available(contract,{'components':{'lead':{'status':'INSTANCE_NATIVE_CONFIRMATION_REQUIRED'}}},{})[0]
        self.assertEqual(action['operation'],'qualify');self.assertTrue(action['legal'])


class ScopedCompilation(unittest.TestCase):
    setUp=FontContracts.setUp
    scene=FontContracts.scene
    plan=FontContracts.plan
    emit=FontContracts.emit

    def test_unused_slot_expansion_does_not_rewrite_ungranted_legacy_nodes(self):
        before=self.plan(self.scene([{'text':'中文','color':'FF0000'}]))
        node=before['objects'][0]
        node['style']['font_faces']={'ea':self.fixture.new['family']}
        node['content']['runs'][0]['font_faces']={'ea':self.fixture.new['family']}
        with patch.object(collection,'assert_language'):compiled=fc.compile_scene(before)
        self.assertNotEqual(compiled['objects'],before['objects'])
        retained=fc.preserve_unwritten_nodes(compiled,before,{'unrelated.shape':['geometry']})
        self.assertEqual(retained['objects'],before['objects'])
        _,_,file,failures=self.emit(retained);self.assertEqual(failures,[]);fc.guard_saved(retained,file)

    def test_actual_ungranted_slot_or_resource_change_requires_explicit_revision(self):
        before=self.plan(self.scene([{'text':'中文'}]));changed=copy.deepcopy(before)
        changed['objects'][0]['content']['runs'][0]['font_faces']['ea']=self.fixture.old['family']
        with self.assertRaisesRegex(ValueError,'FONT_SCOPE_REVISION_REQUIRED'):fc.preserve_unwritten_nodes(changed,before,{'other':[]})
        changed=copy.deepcopy(before);next(iter(changed['font_plan']['bindings'].values()))['resource']['file_sha256']='changed'
        with self.assertRaisesRegex(ValueError,'bound resource changed'):fc.preserve_unwritten_nodes(changed,before,{'other':[]})

    def test_ungranted_colour_change_cannot_hide_as_compiler_metadata(self):
        before=self.plan(self.scene([{'text':'中文','color':'FF0000'}]));changed=copy.deepcopy(before)
        changed['objects'][0]['content']['runs'][0]['color']='333333'
        with self.assertRaisesRegex(ValueError,'native owner scope'):fc.preserve_unwritten_nodes(changed,before,{'other':[]})

    def test_resumed_font_plan_preserves_unwritten_nodes_before_qualified_guard(self):
        from core.component_constructor import construct
        before=self.plan(self.scene([{'text':'中文'}]))
        before['objects'][0]['style']['font_faces']={'ea':self.fixture.new['family']}
        before['objects'][0]['content']['runs'][0]['font_faces']={'ea':self.fixture.new['family']}
        with patch.object(collection,'assert_language'):resumed=fc.compile_scene(before)
        file=self.task/'source/legacy-native-scene.json';write_json(file,before)
        class GuardReached(Exception):pass
        def qualified_guard(task,state,scene,allowed,revisions):
            self.assertEqual(scene['objects'],before['objects'])
            self.assertEqual(allowed,{'edited.shape':['geometry']})
            raise GuardReached()
        with patch('core.workflow_guard.require'),patch('core.component_workflow.guard_changes',side_effect=qualified_guard),self.assertRaises(GuardReached):
            construct(self.task,{},resumed,{'scene':{'path':'source/legacy-native-scene.json','sha256':digest(file)}},{'edited.shape':['geometry']})


class CatalogueIsolation(unittest.TestCase):
    def setUp(self):
        import shutil,platform
        fixture=importlib.import_module('test_font_workflow_v1373').Fonts()
        fixture.setUp();self.addCleanup(fixture.doCleanups);self.root=fixture.root
        self.fonts=self.root/'fonts';self.fonts.mkdir();shutil.copyfile(fixture.old['path'],self.fonts/'real.ttf')
        self.roots=patch('core.font_resources.roots',return_value=[{'path':str(self.fonts),'host':'fixture'}])
        self.roots.start();self.addCleanup(self.roots.stop)
        self.caps={'platform':platform.system(),'render_preference':'auto','configuration_mode':'explicit',
            'windows_office':{'available':True}}

    def catalogue(self,preference='auto'):
        from core import font_resources
        caps=copy.deepcopy(self.caps);caps['render_preference']=preference
        caps['fonts']=font_resources.catalogue(self.root,caps,{})
        return caps

    def test_render_route_does_not_replace_an_existing_task_catalogue(self):
        from core import font_resources,tool_router
        task=self.catalogue();path=Path(task['fonts']['catalogue']);before=(path.read_bytes(),path.stat().st_mtime_ns)
        rendered=self.catalogue('powerpoint')
        self.assertNotEqual(task['fonts']['catalogue'],rendered['fonts']['catalogue'])
        self.assertEqual((path.read_bytes(),path.stat().st_mtime_ns),before)
        self.assertEqual(font_resources.available(task),font_resources.available(rendered))
        self.assertFalse(tool_router.environment_changed(task,self.root))

    def test_overwritten_legacy_shared_catalogue_requires_normal_capability_refresh(self):
        from core import tool_router,font_resources
        task=self.catalogue();task['fonts'].pop('catalogue_sha256')
        task['fonts']['catalogue']=str(self.root/'.lotraAI-build/toolchain/font-resources.json')
        self.catalogue('powerpoint');self.assertTrue(tool_router.environment_changed(task,self.root))
        with self.assertRaisesRegex(ValueError,'another environment'):font_resources.available(task)
        refreshed=self.catalogue();self.assertFalse(tool_router.environment_changed(refreshed,self.root))

    def test_real_font_file_change_remains_an_environment_change(self):
        import os
        from core import tool_router
        task=self.catalogue();file=self.fonts/'real.ttf';stamp=file.stat()
        os.utime(file,ns=(stamp.st_atime_ns,stamp.st_mtime_ns+1000000000))
        self.assertTrue(tool_router.environment_changed(task,self.root))

    def test_frozen_catalogue_tamper_is_rejected_before_returning_faces(self):
        from core import font_resources,tool_router
        task=self.catalogue();Path(task['fonts']['catalogue']).write_text('{}',encoding='utf-8')
        with self.assertRaisesRegex(ValueError,'Frozen font resource catalogue changed'):font_resources.available(task)
        self.assertTrue(tool_router.environment_changed(task,self.root))


class CurrentCandidatePacket(unittest.TestCase):
    def setUp(self):
        import test_packet_submission_v1371 as previous
        self.fixture=previous.PacketSubmission();self.fixture.setUp();self.addCleanup(self.fixture.doCleanups)
        self.task=self.fixture.task

    def test_first_layout_candidate_identity_is_managed_without_selected_baseline(self):
        from core import agent_runtime as runtime,context_control
        from core.common import write_json
        scene=self.task/'scratch/current-scene.json';write_json(scene,{'objects':[]})
        state=self.fixture.fixture.state();candidate={'revision':1,'pptx':self.fixture.fixture.source,
            'scene':runtime._file(self.task,scene)}
        state['active_candidate']=candidate;runtime._save(self.task,state)
        packet=self.fixture.generated('component_layout_work','final_layout')
        packet=__import__('core.packet_compaction',fromlist=['restore']).restore(self.task,packet)
        self.assertIsNone(packet['baseline']);self.assertEqual(packet['current_candidate'],candidate)
        self.assertEqual(context_control.authorize(self.task,packet['current_candidate']['scene']),scene)
        self.assertNotIn('preview',packet['current_candidate'])

    def test_compaction_retains_current_candidate_dependency_and_exact_scene_hash(self):
        from core import agent_runtime as runtime,context_control
        from core.common import write_json
        from core.packet_compaction import compact,restore
        scene=self.task/'scratch/current-scene.json';write_json(scene,{'objects':[]})
        state=self.fixture.fixture.state();state['active_candidate']={'pptx':self.fixture.fixture.source,
            'scene':runtime._file(self.task,scene)};runtime._save(self.task,state)
        packet=restore(self.task,self.fixture.generated('component_layout_work','final_layout'))
        packet['diagnostics']='retained original observations '*2000
        compacted=compact(self.task,packet,2500);state=self.fixture.fixture.state();state['packet']=compacted
        runtime._save(self.task,state);restored=restore(self.task,compacted)
        self.assertEqual(restored['current_candidate'],packet['current_candidate'])
        self.assertEqual(context_control.authorize(self.task,restored['current_candidate']['scene']),scene)
        scene.write_text('{}',encoding='utf-8')
        with self.assertRaisesRegex(ValueError,'dependency changed'):context_control.authorize(self.task,restored['current_candidate']['scene'])

if __name__=='__main__':unittest.main()
