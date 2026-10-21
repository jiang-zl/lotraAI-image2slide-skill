"""Compatibility and failure counterexamples; real host checks have separate receipts."""
from pathlib import Path
import json
import sys
import tempfile
import unittest
from unittest.mock import patch
from lxml import etree as E
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR,MSO_SHAPE
from pptx.util import Inches

REPO=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(REPO/'skills/lotra-image2slide-skill/scripts'))
from core import presentation_backend as backend
from core.geometry import set_native_route,native_route_spec,freeform_route_exact,custom_geometry
from core.mathml import NS,child
from core.office_compatibility import validate
from core.export_destination import reserve,_rename_new
from core.common import safe_path


class Contracts(unittest.TestCase):
    def setUp(self):
        root=REPO/'.lotraAI-build/presentation-tests';root.mkdir(parents=True,exist_ok=True)
        self.tmp=tempfile.TemporaryDirectory(dir=root);self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)

    def test_arbitrary_route_is_editable_unfilled_freeform_and_exact(self):
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,0,0,Inches(2),Inches(2))
        old_id=shape.shape_id;shape.name='source-road'
        points=[[1,1],[2,1.2],[3,1],[4,2]]
        spec=native_route_spec(points,914400)
        set_native_route(shape,points,914400)
        self.assertTrue(freeform_route_exact(shape._element,spec))
        file=self.root/'Freeform.pptx';prs.save(file);validate(file)
        saved=Presentation(file).slides[0].shapes[0]
        self.assertEqual(saved.shape_id,old_id);self.assertEqual(saved.name,'source-road')
        self.assertTrue(freeform_route_exact(saved._element,spec))
        saved._element.find('.//a:moveTo/a:pt',NS).set('x','1')
        self.assertFalse(freeform_route_exact(saved._element,spec))

    def test_connected_arbitrary_route_cannot_drop_ports_or_mutate(self):
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        node=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(1),Inches(1))
        route=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,0,0,Inches(2),Inches(2))
        child(route._element.find('p:nvCxnSpPr/p:cNvCxnSpPr',NS),'a:stCxn',id=node.shape_id,idx=1)
        before=E.tostring(route._element)
        with self.assertRaisesRegex(ValueError,'SOURCE_DECISION_REQUIRED'):
            set_native_route(route,[[1,1],[2,1.2],[3,1]],914400)
        self.assertEqual(E.tostring(route._element),before)

    def test_schema_valid_custom_connector_is_rejected_before_adoption(self):
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,0,0,Inches(1),Inches(1))
        custom_geometry(shape,[{'points':[[0,0],[1,.2],[.5,1]],'closed':False}],fill=False)
        path=self.root/'Invalid-office.pptx';prs.save(path)
        with self.assertRaisesRegex(ValueError,'Office connectors require preset'):validate(path)

    def test_production_builder_and_both_inspectors_agree_on_arbitrary_route(self):
        from core.fixtures import create_probe
        from core.recipes import node
        from core.build import build
        from core.inspect import check
        from core.native_core_inspect import check as native_check
        base=self.root/'.lotraAI-build/developer-drafts/scene'
        base.parent.mkdir(parents=True)
        scene=create_probe(base)
        scene['objects']=[node('source-route','connector',[10,10,250,50],
            {'points':[[10,10],[120,30],[180,15],[260,60]],'arrow':'triangle'},stroke='202020',line_width=2)]
        scene['relations']=[];scene['prototypes']=[]
        from core.workflow_guard import developer_writer
        file=base/'Generated.pptx'
        with developer_writer(self.root):build(scene,base,file)
        for inspector in (check,native_check):
            result=inspector(scene,base,file)
            failures=[row for row in result['checks'] if row['status']=='FAIL']
            self.assertEqual(failures,[],str(failures))

    def test_explicit_missing_backend_does_not_silently_use_alternative(self):
        with patch.object(backend,'discover',return_value=[]):
            value=backend.select(self.root,'wps',libreoffice={'executable':'declared-soffice'})
        self.assertEqual(value['status'],'CAPABILITY_BLOCKED');self.assertEqual(value['backend'],'WPS')

    def test_auto_can_choose_another_qualified_native_adapter(self):
        caps=[{'backend':'PowerPoint'},{'backend':'WPS'}]
        def qualified(project,cap,refresh):return {**cap,'status':'READY' if cap['backend']=='WPS' else 'CAPABILITY_BLOCKED'}
        with patch.object(backend,'discover',return_value=caps),patch.object(backend,'probe',side_effect=qualified):
            value=backend.select(self.root)
        self.assertEqual(value['backend'],'WPS');self.assertEqual(value['status'],'READY')
        self.assertEqual(len(value['selection']['failed_candidates']),1)

    def test_resume_reuses_destination_but_new_export_reserves_another(self):
        first=reserve(self.root,'same-input')
        resumed=reserve(self.root,'same-input')
        self.assertEqual(first['final'],resumed['final'])
        self.assertNotEqual(first['stage'],resumed['stage'])
        another=reserve(self.root,'same-input',new=True)
        self.assertNotEqual(first['final'],another['final'])
        self.assertTrue(Path(first['final']).name.startswith('output-slide-'))
        self.assertEqual(Path(first['final']).parent,self.root)

    def test_destination_default_does_not_expand_private_path_guard(self):
        external=self.root/'chosen-output'
        plan=reserve(self.root,'input',output_root=external)
        self.assertEqual(Path(plan['final']).parent,external)
        with self.assertRaises(ValueError):safe_path(self.root,'../escape')
        with self.assertRaises(ValueError):reserve(self.root,'input',output_root=self.root/'.lotraAI-build/private')

    def test_atomic_publish_never_replaces_even_an_empty_existing_directory(self):
        stage=self.root/'stage';stage.mkdir();(stage/'sentinel').write_bytes(b'own bytes')
        final=self.root/'existing';final.mkdir()
        with self.assertRaises(FileExistsError):_rename_new(stage,final)
        self.assertTrue(stage.exists());self.assertEqual(list(final.iterdir()),[])
        fresh=self.root/'new';_rename_new(stage,fresh)
        self.assertEqual((fresh/'sentinel').read_bytes(),b'own bytes')

    def test_moved_native_pdf_evidence_keeps_exact_disposed_ownership(self):
        from core import libreoffice_backend
        from core.common import digest,write_json,read_json
        from core.transient_export import register,release_task_pdfs
        import shutil,pymupdf
        destination=self.root/'chosen-volume';destination.mkdir()
        pptx=destination/'Final_editable.pptx';backend.fixture(pptx)
        def application_copy(source,output,executable):
            output=Path(output);(output/'saved').mkdir(parents=True)
            saved=output/'saved/Copy.pptx';shutil.copyfile(source,saved)
            render=output/'reopen-render';render.mkdir()
            pdf=render/'Final.pdf'
            doc=pymupdf.open();doc.new_page().insert_text((30,30),'Ownership fixture');doc.save(pdf);doc.close()
            receipt={'status':'PASS','backend':'LibreOffice','pptx_hash':digest(saved),
                'pdf':{'status':'EXPORTED'},'pdf_hash':digest(pdf)}
            write_json(render/'render.json',receipt);register(render,saved,receipt)
            return {'status':'MOCK_APPLICATION_FOR_OWNERSHIP_CONTRACT_ONLY'}
        audit=self.root/'private-audit'
        with patch.object(libreoffice_backend,'save_reopen',side_effect=application_copy):
            result=backend.validate_delivery(pptx,audit,{'presentation_backend':{'backend':'LibreOffice','executable':'fixture'}},destination=destination)
        marker=audit/'libreoffice-delivery/reopen-render/Pdf-lifecycle.json'
        proof=read_json(marker)
        self.assertEqual(proof['root'],str(marker.parent))
        self.assertEqual(proof['status'],'DISPOSED')
        self.assertTrue(Path(proof['pptx']).is_file())
        self.assertEqual(release_task_pdfs(audit)['status'],'DEPENDENTS_COMPLETE')
        self.assertEqual([p.name for p in destination.iterdir()],['Final_editable.pptx'])


if __name__=='__main__':unittest.main(verbosity=2)
