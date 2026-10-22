"""Font revision, exact saved-run and resumable package-commit regressions."""
import copy
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from contextlib import ExitStack
from unittest.mock import patch
from zipfile import ZipFile
from fontTools.fontBuilder import FontBuilder
from fontTools.pens.ttGlyphPen import TTGlyphPen
from lxml import etree as E
from pptx import Presentation
from pptx.util import Inches

REPO=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(REPO/'skills/lotra-image2slide-skill/scripts'))
from core import font_contract as fc,font_embedding as fe,font_collection as collection,font_resources as resources
from core.common import digest,identity,read_json,write_json


class Fonts(unittest.TestCase):
    def setUp(self):
        parent=REPO/'.lotraAI-build/font-tests';parent.mkdir(parents=True,exist_ok=True)
        self.tmp=tempfile.TemporaryDirectory(dir=parent);self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.task=self.root/'task';(self.task/'source').mkdir(parents=True)
        self.old=self.face('Review Old');self.new=self.face('Review New')
        self.bindings=[{'family':self.old['family'],'resource':self.old}]

    def face(self,family):
        file=self.root/(family.replace(' ','-')+'.ttf')
        chars='ABab0123甲乙中文';names=['.notdef']+['g'+str(i) for i in range(len(chars))]
        builder=FontBuilder(1000,isTTF=True);builder.setupGlyphOrder(names)
        builder.setupCharacterMap({ord(c):names[i+1] for i,c in enumerate(chars)})
        glyphs={}
        for name in names:
            pen=TTGlyphPen(None);pen.moveTo((100,0));pen.lineTo((500,0));pen.lineTo((300,700));pen.closePath();glyphs[name]=pen.glyph()
        builder.setupGlyf(glyphs);builder.setupHorizontalMetrics({n:(600,0) for n in names})
        builder.setupHorizontalHeader(ascent=800,descent=-200)
        builder.setupNameTable({'familyName':family,'styleName':'Regular','uniqueFontIdentifier':family,'fullName':family,'psName':family.replace(' ','')})
        builder.setupOS2(sTypoAscender=800,sTypoDescender=-200,usWinAscent=800,usWinDescent=200,fsType=0)
        builder.setupPost();builder.save(file)
        face=resources.describe(file,0,digest(file));face.update(id=identity(family),hosts=['fixture'])
        return face

    def package(self,name='input.pptx',families=None,texts=('甲','乙')):
        file=self.root/name;prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(1));shape.name='lead'
        for text,family in zip(texts,families or [self.old['family']]*len(texts)):
            run=shape.text_frame.paragraphs[0].add_run();run.text=text;run.font.name=family
            E.SubElement(run._r.get_or_add_rPr(),'{'+fc.A+'}ea',typeface=family)
        prs.save(file);return file

    def scene(self,families=None,texts=('甲','乙')):
        families=families or [self.old['family']]*len(texts)
        scene={'revision':1,'source':{'sha256':identity('source')},'objects':[{'id':'lead','render_type':'text',
            'semantic_type':'label','bbox':[0,0,100,30],'style':{'font':self.old['family']},
            'content':{'runs':[{'text':t,'font':f} for t,f in zip(texts,families)]}}]}
        return self.plan(scene)

    def plan(self,scene):
        bindings={}
        for request in fc.requests(scene):
            face=self.old if request['proposed_family']==self.old['family'] else self.new
            bindings[request['id']]={'family':face['family'],'resource':face,'status':'approximate','request':request,
                'request_hash':identity({k:request[k] for k in ('owner','location','role','characters','literal_hash','bold','italic','proposed_family')})}
        scene['font_plan']={'bindings':bindings,'version':fc.VERSION}
        return scene

    def test_real_direct_embedding_preserves_slide_bytes(self):
        file=self.package()
        with ZipFile(file) as z:before=z.read('ppt/slides/slide1.xml')
        result=fe.embed(file,self.bindings)
        self.assertEqual(result['status'],'EMBEDDED_FULL_EDITABLE_FONTS');self.assertEqual(result['pptx_sha256'],digest(file))
        with ZipFile(file) as z:
            self.assertEqual(before,z.read('ppt/slides/slide1.xml'))
            from core.font_retention import verify_package
            self.assertTrue(verify_package(z,self.bindings))

    def blocked_stage(self,file,provider=None):
        with patch.object(Path,'replace',side_effect=PermissionError('held Windows handle')):
            with self.assertRaises(fe.FontCommitPending):
                (fe.transfer_provider(file,provider,self.bindings) if provider else fe.embed(file,self.bindings))

    def test_commit_retry_reuses_verified_stage_without_rewriting_zip(self):
        file=self.package();before=digest(file);self.blocked_stage(file)
        self.assertEqual(before,digest(file))
        with patch.object(fe,'ZipFile',wraps=ZipFile) as archive:
            self.assertEqual(fe.embed(file,self.bindings)['status'],'EMBEDDED_FULL_EDITABLE_FONTS')
        self.assertFalse(any('x' in str(call.args[1:]) or 'w' in str(call.args[1:]) for call in archive.call_args_list))

    @unittest.skipUnless(os.name=='nt','real Windows sharing semantics')
    def test_windows_external_zip_lock_releases_and_commit_resumes(self):
        file=self.package()
        with ZipFile(file):
            with self.assertRaises(fe.FontCommitPending):fe.embed(file,self.bindings)
        self.assertEqual(fe.embed(file,self.bindings)['status'],'EMBEDDED_FULL_EDITABLE_FONTS')

    def test_completed_commit_is_repeat_safe(self):
        file=self.package();first=fe.embed(file,self.bindings);stamp=file.stat().st_mtime_ns
        self.assertEqual(fe.embed(file,self.bindings),first);self.assertEqual(stamp,file.stat().st_mtime_ns)

    def test_unknown_pending_file_is_preserved(self):
        file=self.package();pending=file.with_name(file.stem+'.fonts-pending.pptx');pending.write_bytes(b'unowned')
        with self.assertRaises(FileExistsError):fe.embed(file,self.bindings)
        self.assertEqual(pending.read_bytes(),b'unowned')

    def test_changed_source_blocks_commit(self):
        file=self.package();self.blocked_stage(file);file.write_bytes(file.read_bytes()+b'changed')
        with self.assertRaisesRegex(ValueError,'FONT_COMMIT_CONFLICT'):fe.embed(file,self.bindings)

    def test_changed_stage_blocks_commit(self):
        file=self.package();self.blocked_stage(file)
        pending=file.with_name(file.stem+'.fonts-pending.pptx');pending.write_bytes(pending.read_bytes()+b'changed')
        with self.assertRaisesRegex(ValueError,'FONT_COMMIT_CONFLICT'):fe.embed(file,self.bindings)

    def test_changed_binding_cannot_adopt_another_stage(self):
        file=self.package();self.blocked_stage(file)
        with self.assertRaisesRegex(ValueError,'FONT_COMMIT_CONFLICT'):fe.embed(file,[{'family':self.new['family'],'resource':self.new}])

    def test_current_font_budget_is_rechecked_on_resume(self):
        file=self.package();self.blocked_stage(file)
        with self.assertRaisesRegex(ValueError,'FONT_BUDGET_PENDING'):fe.embed(file,self.bindings,max_font_bytes=1)

    def test_provider_commit_recovers_without_copying_slide_content(self):
        provider=self.package('provider.pptx');fe.embed(provider,self.bindings)
        original=self.package('original.pptx');self.blocked_stage(original,provider)
        self.assertEqual(fe.transfer_provider(original,provider,self.bindings)['status'],'EMBEDDED_FULL_EDITABLE_FONTS')

    def test_each_run_binding_is_checked(self):
        scene=self.scene([self.old['family'],self.new['family']])
        file=self.package(families=[self.old['family'],self.new['family']])
        self.assertEqual(len(fc.guard_saved(scene,file)['checked_locations']),2)

    def test_swapped_run_families_are_rejected(self):
        scene=self.scene([self.old['family'],self.new['family']])
        file=self.package(families=[self.new['family'],self.old['family']])
        with self.assertRaisesRegex(ValueError,'lead/p0.r0/ea'):fc.guard_saved(scene,file)

    def test_changed_literal_is_rejected(self):
        with self.assertRaisesRegex(ValueError,'literal'):fc.guard_saved(self.scene(),self.package(texts=('乙','甲')))

    def test_extra_unbound_run_is_rejected(self):
        with self.assertRaisesRegex(ValueError,'location-bound'):fc.guard_saved(self.scene(texts=('甲',)),self.package())

    def binding_context(self):
        scene=self.scene(texts=('中文',));state={'project':str(self.root),'source':scene['source'],'objectives':{},
            'capabilities':{'fonts':{'environment_key':'fixture'}},'font_collection':{},
            'font_contract':{'version':fc.VERSION,'bindings':copy.deepcopy(scene['font_plan']['bindings']),'history':[]}}
        changed=copy.deepcopy(scene);changed.pop('font_plan');changed['objects'][0]['content']['runs'][0]['font']=self.new['family']
        request=fc.requests(changed)[0];prior=state['font_contract']['bindings'][request['id']]
        comparison=self.task/'source/ranking.json';write_json(comparison,{'source_hash':state['source']['sha256'],'selected_style':{'font':self.new['family']}})
        value={'request_id':request['id'],'family':self.new['family'],'resource_id':self.new['id'],'origin':'raster_inference',
            'status':'approximate','comparison_ref':{'path':'source/ranking.json','sha256':digest(comparison)},'alternatives':[self.old['family']],
            'prior_binding_hash':identity(prior),'revision_reason':'Located original glyph correction',
            'evidence':{'source_hash':state['source']['sha256'],'roi':[0,0,100,30],'description':'Original glyph observation'}}
        return scene,state,changed,request,value

    def collection_context(self):
        stack=ExitStack();faces={f['id']:f for f in (self.old,self.new)}
        for module,name,value in ((resources,'available',lambda caps:list(faces.values())),
            (collection,'prepare',lambda state:None),(collection,'load',lambda project:{'faces':faces}),
            (collection,'assert_language',lambda face,role:None),(collection,'verify_choice',lambda *args:None),
            (collection,'archive',lambda *args:None),(collection,'eligible',lambda *args,**kwargs:list(faces.values())),
            (collection,'ranking',lambda *args:None)):
            stack.enter_context(patch.object(module,name,value))
        return stack

    def test_whole_owner_authority_permits_its_run_font_revision(self):
        _,state,scene,request,value=self.binding_context();state['font_pending_candidate']={'allowed':{'lead':['replace_component']}}
        with self.collection_context():fc.bind(self.task,state,scene,[value],persist=False)
        self.assertEqual(state['font_contract']['bindings'][request['id']]['family'],self.new['family'])

    def test_text_body_is_not_implicit_font_authority(self):
        _,state,scene,_,value=self.binding_context();state['font_pending_candidate']={'allowed':{'lead':['text_body']}}
        with self.collection_context(),self.assertRaisesRegex(ValueError,'FONT_REVISION_GRANT_REQUIRED'):fc.bind(self.task,state,scene,[value],persist=False)

    def test_semantic_run_font_grant_survives_text_body_projection(self):
        _,state,scene,_,value=self.binding_context();state['font_pending_candidate']={'allowed':{'lead':['text_body']},
            'reason':{'font_revision_grants':[{'owner':'lead','field':['content','runs',0,'font']}]}}
        with self.collection_context():fc.bind(self.task,state,scene,[value],persist=False)

    def test_text_content_grant_cannot_authorize_a_font(self):
        _,state,scene,_,value=self.binding_context();state['font_pending_candidate']={'allowed':{'lead':['text_body']},
            'reason':{'font_revision_grants':[{'owner':'lead','field':['content','runs',0,'text']}]}}
        with self.collection_context(),self.assertRaisesRegex(ValueError,'FONT_REVISION_GRANT_REQUIRED'):fc.bind(self.task,state,scene,[value],persist=False)

    def test_stale_prior_hash_is_rejected_without_mutating_active_bindings(self):
        _,state,scene,_,value=self.binding_context();state['font_pending_candidate']={'allowed':{'lead':['style.font']}}
        value['prior_binding_hash']='stale';before=copy.deepcopy(state['font_contract'])
        with self.collection_context(),self.assertRaisesRegex(ValueError,'FONT_REVISION_STALE_BINDING'):fc.bind(self.task,state,scene,[value],persist=False)
        self.assertEqual(before,state['font_contract'])

    def test_packet_supplies_prior_hash_and_current_revision_authority(self):
        _,state,scene,request,_=self.binding_context();state['font_pending_candidate']={'allowed':{'lead':['replace_component']}}
        with self.collection_context():packet=fc.packet(self.task,state,scene)
        detail=packet['requests'][0]['binding_revision'];self.assertTrue(detail['authorized'])
        self.assertEqual(detail['prior_binding_hash'],identity(state['font_contract']['bindings'][request['id']]))

    def test_pending_binding_view_keeps_active_binding_unchanged(self):
        _,state,scene,request,value=self.binding_context();state['font_pending_candidate']={'allowed':{'lead':['style.font']}}
        working=copy.deepcopy(state)
        with self.collection_context():fc.bind(self.task,working,scene,[value],persist=False)
        state['font_pending_candidate']['staged_font_contract']=working['font_contract']
        self.assertEqual(state['font_contract']['bindings'][request['id']]['family'],self.old['family'])
        self.assertEqual(fc.pending_state(state)['font_contract']['bindings'][request['id']]['family'],self.new['family'])

    def test_whole_owner_grant_preserves_existing_user_font_lock(self):
        _,state,scene,_,value=self.binding_context();state['component_workflow']={"bound":True}
        state['font_pending_candidate']={'allowed':{'lead':['replace_component']}}
        ledger={'components':{'lead':{'user_locks':[{'owner':'lead','field':['style','font']}]}},'orphan_user_locks':[]}
        with self.collection_context(),patch('core.component_workflow._load',return_value=(None,ledger)),self.assertRaisesRegex(ValueError,'USER_LOCK_CONFLICT'):
            fc.bind(self.task,state,scene,[value],persist=False)

    def test_failed_batch_does_not_commit_earlier_valid_revision(self):
        _,state,scene,_,value=self.binding_context();state['font_pending_candidate']={'allowed':{'lead':['style.font']}}
        before=copy.deepcopy(state['font_contract']);invalid={**value,'request_id':'foreign'}
        with self.collection_context(),self.assertRaises(ValueError):fc.bind(self.task,state,scene,[value,invalid],persist=False)
        self.assertEqual(before,state['font_contract'])

    def test_mixed_scripts_use_their_own_font_slots(self):
        scene=self.scene(texts=('A甲',));ea=next(b for b in scene['font_plan']['bindings'].values() if b['request']['role']=='ea')
        ea.update(family=self.new['family'],resource=self.new)
        file=self.package(texts=('A甲',));prs=Presentation(file)
        props=prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]._r.rPr
        props.find('{'+fc.A+'}ea').set('typeface',self.new['family']);prs.save(file)
        self.assertEqual(len(fc.guard_saved(scene,file)['checked_locations']),2)

    def test_plain_multiline_text_keeps_its_literal_and_font(self):
        scene=self.scene();scene['objects'][0]['content']={'text':'A\nB'};self.plan(scene)
        file=self.root/'multiline.pptx';prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        shape=slide.shapes.add_textbox(0,0,Inches(3),Inches(1));shape.name='lead'
        from core.build import text_body
        text_body(shape,scene['objects'][0]['content'],{'font':self.old['family']});prs.save(file)
        self.assertEqual(len(fc.guard_saved(scene,file)['checked_locations']),1)

    def test_table_cells_are_checked_separately(self):
        scene=self.scene();obj=scene['objects'][0];obj['render_type']='table'
        obj['content']={'rows':[['甲','乙']],'cell_styles':[[{'font':self.old['family']},{'font':self.new['family']}]]};self.plan(scene)
        file=self.root/'table.pptx';prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        frame=slide.shapes.add_table(1,2,0,0,Inches(4),Inches(1));frame.name='lead'
        for index,family in enumerate((self.old['family'],self.new['family'])):
            cell=frame.table.cell(0,index);cell.text=('甲','乙')[index]
            run=cell.text_frame.paragraphs[0].runs[0];run.font.name=family;E.SubElement(run._r.get_or_add_rPr(),'{'+fc.A+'}ea',typeface=family)
        prs.save(file);self.assertEqual(len(fc.guard_saved(scene,file)['checked_locations']),2)
        frame.table.cell(0,0).text_frame.paragraphs[0].runs[0]._r.rPr.find('{'+fc.A+'}ea').set('typeface',self.new['family']);prs.save(file)
        with self.assertRaisesRegex(ValueError,'cell.0.0'):fc.guard_saved(scene,file)

    def test_inline_office_math_keeps_its_own_location(self):
        scene=self.scene();obj=scene['objects'][0]
        obj['content']={'runs':[{'text':'A'},{'ast':{'op':'token','text':'B'},'font':self.new['family']}]};self.plan(scene)
        file=self.root/'math.pptx';prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        shape=slide.shapes.add_textbox(0,0,Inches(3),Inches(1));shape.name='lead'
        from core.build import text_body
        text_body(shape,obj['content'],obj['style']);prs.save(file)
        self.assertEqual(len(fc.guard_saved(scene,file)['checked_locations']),2)

    def test_saved_font_build_resumes_without_reconstructing_native_shapes(self):
        from core.build import _finish_font_build
        from core.session import implementation_identity
        scene=self.scene();output=self.root/'drawing.pptx';temporary=self.package('drawing.pending.pptx')
        record={'implementation_version':'1.37.3','implementation':implementation_identity(),'scene_hash':identity(scene),
            'scene_revision':1,'objects':{},'construction_storage':{},'raw_pptx_sha256':digest(temporary)}
        write_json(output.with_suffix('.font-build.json'),record)
        with patch.object(Path,'replace',side_effect=PermissionError('Windows sharing lock')),self.assertRaises(fe.FontCommitPending):
            _finish_font_build(scene,output,record)
        with patch('core.build.Presentation',side_effect=AssertionError('must not reconstruct')):
            result=_finish_font_build(scene,output,record)
        self.assertTrue(output.is_file());self.assertEqual(result['pptx_hash'],digest(output))
        self.assertEqual(_finish_font_build(scene,output,record),result)

    def test_font_commit_retry_reuses_reserved_native_attempt(self):
        from core.component_transactions import attempt_folder
        folder=self.task/'native-attempt-001';folder.mkdir();marker=folder/'owned.font-commit.json';write_json(marker,{'phase':'READY'})
        tx={'native_attempts':['native-attempt-001'],'font_commit_pending':{'attempt':'native-attempt-001',
            'record':{'path':'native-attempt-001/owned.font-commit.json','sha256':digest(marker)}}}
        with patch('core.component_transactions.phase',side_effect=lambda task,tx,phase,**values:tx.update(values)):
            self.assertEqual(attempt_folder(self.task,tx),folder)
        self.assertEqual(len(tx['native_attempts']),1);self.assertEqual(tx['font_commit_retries'],1)

    def test_completed_commit_phase_is_recoverable_after_interruption(self):
        from core.component_transactions import attempt_folder
        folder=self.task/'native-attempt-001';folder.mkdir();marker=folder/'owned.font-commit.json'
        record={'schema':'lotra-font-commit-1','phase':'READY','staged_sha256':'retained'};write_json(marker,record)
        tx={'native_attempts':['native-attempt-001'],'font_commit_pending':{'attempt':'native-attempt-001',
            'record':{'path':'native-attempt-001/owned.font-commit.json','sha256':digest(marker)},
            'commit_identity':identity({k:v for k,v in record.items() if k!='phase'})}}
        record['phase']='COMMITTED';write_json(marker,record)
        with patch('core.component_transactions.phase',side_effect=lambda task,tx,phase,**values:tx.update(values)):
            self.assertEqual(attempt_folder(self.task,tx),folder)

    def test_commit_recovery_has_a_finite_retry_budget(self):
        from core.component_transactions import attempt_folder
        tx={'native_attempts':['native-attempt-001'],'font_commit_retries':3,
            'font_commit_pending':{'attempt':'native-attempt-001'}}
        with self.assertRaisesRegex(ValueError,'retry budget'):attempt_folder(self.task,tx)

    def test_unauthorized_revision_returns_to_evidenced_source_work(self):
        _,state,scene,request,_=self.binding_context();scene_file=self.task/'source/pending.json';write_json(scene_file,scene)
        state['font_pending_candidate']={'allowed':{'lead':['text_body']},'scene':{'path':'source/pending.json','sha256':digest(scene_file)},'baseline':None}
        evidence=self.task/'source/observation.json';write_json(evidence,{'source_hash':state['source']['sha256'],
            'pptx_hash':None,'observer':'host_model','observations':['Located run needs a font-specific source decision']})
        contract={'components':[{'id':'lead','owners':['lead']}]};ledger={'components':{'lead':{'status':'RECHECK_REQUIRED'}}}
        before=copy.deepcopy(state['font_contract'])
        with self.collection_context(),patch('core.component_workflow._load',return_value=(contract,ledger)),patch('core.component_workflow._save_ledger'):
            result=fc.retain_revision_decision(self.task,state,scene,{'requests':[request]},
                {'request_ids':[request['id']],'reason':'Font field is not authorized',
                 'evidence':{'path':'source/observation.json','sha256':digest(evidence)}})
        self.assertEqual(result['status'],'FONT_SOURCE_DECISION_REQUIRED');self.assertFalse(result['acceptance'])
        self.assertNotIn('font_pending_candidate',state);self.assertEqual(state['font_contract'],before)
        self.assertEqual(ledger['components']['lead']['status'],'SOURCE_DECISION_REQUIRED')
        self.assertTrue((self.task/ledger['font_redecision_history'][0]['path']).is_file())


if __name__=='__main__':unittest.main()
